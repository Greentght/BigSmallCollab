#!/usr/bin/env python3
"""Analyze already-exported fusion and distillation predictions.

This is a read-only post-hoc analysis. It never imports source model classes,
loads model weights, or trains either source models or fusion heads.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
from collections import defaultdict

import numpy as np

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from collab import artifacts
from experiments.storage import (PROJECT_DATA_ROOT, REPORTS_ROOT, RESULTS_ROOT,
                                 external_path, require_report_output)


DATASETS = [
    ("004", "BNCI2014004"),
    ("001", "BNCI2014001"),
    ("001-4", "BNCI2014001-4"),
    ("alexmi", "AlexMI"),
    ("15001", "BNCI2015001"),
]
DISPLAY_BY_CANONICAL = {canonical: display for display, canonical in DATASETS}
CANONICAL_BY_ALIAS = {}
for _display, _canonical in DATASETS:
    CANONICAL_BY_ALIAS[_display.casefold()] = _canonical
    CANONICAL_BY_ALIAS[_canonical.casefold()] = _canonical
CANONICAL_BY_ALIAS.update({"bnci2014001_4": "BNCI2014001-4"})
TEACHERS = ("mirepnet", "cbramod")
STUDENTS = ("ifnet", "eegnet", "adfcnn")
PAIRS = [(teacher, student) for teacher in TEACHERS for student in STUDENTS]
FEWSHOT_SEEDS = (666, 667, 668)
EXPECTED_SUBJECTS = {
    "BNCI2014004": 9,
    "BNCI2014001": 9,
    "BNCI2014001-4": 9,
    "AlexMI": 8,
    "BNCI2015001": 12,
}

SUMMARY_FIELDS = [
    "analysis", "protocol", "dataset", "teacher", "student", "method",
    "metric", "seed_set", "n_seeds", "n_subjects", "n_units",
    "value_mean", "value_std", "delta_vs_gate_mean", "delta_vs_gate_std",
    "pooled_count", "pooled_denominator", "detail", "source",
]
MISSING_FIELDS = [
    "scope", "protocol", "dataset", "teacher", "student", "method",
    "seed_set", "missing_item", "detail", "source_expected",
]


def canonical_dataset(value: str) -> str | None:
    return CANONICAL_BY_ALIAS.get(str(value or "").strip().casefold())


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def load_light_artifact(dataset: str, model: str, subject_key: int,
                       seed: int, split: str) -> dict:
    """Read only logits, labels, and alignment metadata from a saved artifact."""
    path = Path(artifacts.artifact_path(dataset, model, subject_key, seed, split))
    if not path.is_file():
        raise FileNotFoundError(path)
    with np.load(path, allow_pickle=False) as archive:
        required = {"logits", "y", "sample_uid", "split_policy"}
        missing = required - set(archive.files)
        if missing:
            raise ValueError(f"{path}: missing required fields {sorted(missing)}")
        return {
            "logits": np.asarray(archive["logits"]),
            "y": np.asarray(archive["y"], dtype=np.int64),
            "sample_uid": np.asarray(archive["sample_uid"]),
            "split_policy": str(archive["split_policy"].item()),
        }


def npz_dict(path: Path, names=None) -> dict[str, np.ndarray]:
    with np.load(path, allow_pickle=True) as archive:
        selected = archive.files if names is None else [name for name in names if name in archive.files]
        return {key: np.asarray(archive[key]) for key in selected}


LOSO_ARRAY_FIELDS = ("logits", "test_logits", "y", "labels", "test_y",
                     "trial_uid", "trial_uids", "sample_uid", "local_sample_uid")
KD_TEST_ARRAY_FIELDS = ("test_logits", "logits", "test_y", "y", "labels",
                        "test_uid", "test_sample_uid", "trial_uid", "sample_uid")


def sha256_array(value: np.ndarray) -> str:
    array = np.asarray(value)
    digest = hashlib.sha256()
    digest.update(str(array.dtype).encode("utf-8"))
    digest.update(str(tuple(array.shape)).encode("utf-8"))
    digest.update(np.ascontiguousarray(array).tobytes())
    return digest.hexdigest()


def softmax(logits: np.ndarray) -> np.ndarray:
    z = np.asarray(logits, dtype=np.float64)
    z = z - z.max(axis=1, keepdims=True)
    ez = np.exp(z)
    return ez / ez.sum(axis=1, keepdims=True)


def percent_accuracy(logits: np.ndarray, labels: np.ndarray) -> float:
    return float(np.mean(np.argmax(logits, axis=1) == np.asarray(labels))) * 100.0


def csv_float(value: object) -> float | None:
    try:
        if value is None or str(value).strip() in {"", "NA", "nan", "None"}:
            return None
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def int_field(row: dict[str, str], *names: str) -> int | None:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip() != "":
            try:
                return int(float(value))
            except (TypeError, ValueError):
                pass
    return None


def field(row: dict[str, str], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip() not in {"", "NA", "None"}:
            return str(value).strip()
    return ""


def stat_over_seed_means(values: list[tuple[int, int, float]]) -> tuple[float | None, float | None, int, int, int]:
    """Equal-weight subjects within each seed, then mean/sample SD over seeds."""
    if not values:
        return None, None, 0, 0, 0
    by_seed: dict[int, list[float]] = defaultdict(list)
    subjects: set[int] = set()
    for seed, subject, value in values:
        by_seed[int(seed)].append(float(value))
        subjects.add(int(subject))
    seed_means = np.asarray([np.mean(v) for _, v in sorted(by_seed.items())], dtype=float)
    sd = float(seed_means.std(ddof=1)) if len(seed_means) > 1 else 0.0
    return float(seed_means.mean()), sd, len(by_seed), len(subjects), len(values)


def seed_string(seeds) -> str:
    return ",".join(str(s) for s in sorted(set(int(x) for x in seeds)))


def add_summary(rows: list[dict], *, analysis: str, protocol: str,
                dataset: str, teacher: str = "", student: str = "",
                method: str = "", metric: str = "", values=None,
                delta_values=None, pooled_count=None, pooled_denominator=None,
                detail: str = "", source: str = "") -> None:
    values = values or []
    mean, sd, n_seeds, n_subjects, n_units = stat_over_seed_means(values)
    delta_mean = delta_sd = None
    if delta_values:
        delta_mean, delta_sd, _, _, _ = stat_over_seed_means(delta_values)
    rows.append({
        "analysis": analysis, "protocol": protocol,
        "dataset": DISPLAY_BY_CANONICAL.get(dataset, dataset),
        "teacher": teacher, "student": student, "method": method,
        "metric": metric,
        "seed_set": seed_string(seed for seed, _, _ in values),
        "n_seeds": n_seeds, "n_subjects": n_subjects, "n_units": n_units,
        "value_mean": mean, "value_std": sd,
        "delta_vs_gate_mean": delta_mean,
        "delta_vs_gate_std": delta_sd,
        "pooled_count": pooled_count,
        "pooled_denominator": pooled_denominator,
        "detail": detail, "source": source,
    })


def add_missing(missing: list[dict], *, scope: str, protocol: str,
                dataset: str = "", teacher: str = "", student: str = "",
                method: str = "", seed_set: str = "", missing_item: str,
                detail: str = "", source_expected: str = "") -> None:
    missing.append({
        "scope": scope, "protocol": protocol,
        "dataset": DISPLAY_BY_CANONICAL.get(dataset, dataset),
        "teacher": teacher, "student": student, "method": method,
        "seed_set": seed_set, "missing_item": missing_item,
        "detail": detail, "source_expected": source_expected,
    })


def load_fewshot_pair(dataset: str, teacher: str, student: str,
                      subject_key: int, seed: int) -> dict:
    big_tr = load_light_artifact(dataset, teacher, subject_key, seed, "train")
    small_tr = load_light_artifact(dataset, student, subject_key, seed, "train")
    big_te = load_light_artifact(dataset, teacher, subject_key, seed, "test")
    small_te = load_light_artifact(dataset, student, subject_key, seed, "test")
    if any(item["split_policy"] != "fewshot_stratified_random"
           for item in (big_tr, small_tr, big_te, small_te)):
        raise ValueError("unexpected few-shot split policy")
    for split_name, big, small in (("train", big_tr, small_tr), ("test", big_te, small_te)):
        if not np.array_equal(big["sample_uid"], small["sample_uid"]):
            raise ValueError(f"few-shot {split_name} UID mismatch")
        if not np.array_equal(big["y"], small["y"]):
            raise ValueError(f"few-shot {split_name} label mismatch")
    return {"big_tr": big_tr, "small_tr": small_tr, "y_tr": big_tr["y"],
            "big_te": big_te, "small_te": small_te, "y_te": big_te["y"]}


def gate_arrays(big_train: np.ndarray, small_train: np.ndarray,
                y_train: np.ndarray, big_test: np.ndarray,
                small_test: np.ndarray) -> tuple[np.ndarray, np.ndarray, dict]:
    """Match collab.fusion.gate_conf_acc alpha=beta=temperature=1."""
    p_big, p_small = softmax(big_test), softmax(small_test)
    acc_big = max(float(np.mean(np.argmax(big_train, axis=1) == y_train)), 1e-8)
    acc_small = max(float(np.mean(np.argmax(small_train, axis=1) == y_train)), 1e-8)
    conf_big = np.maximum(p_big.max(axis=1), 1e-8)
    conf_small = np.maximum(p_small.max(axis=1), 1e-8)
    score_big, score_small = conf_big * acc_big, conf_small * acc_small
    w_big = score_big / np.maximum(score_big + score_small, 1e-8)
    probs = w_big[:, None] * p_big + (1.0 - w_big[:, None]) * p_small
    return probs, w_big, {"train_acc_teacher_pct": acc_big * 100.0,
                          "train_acc_student_pct": acc_small * 100.0}


def read_fusion_csv(dataset: str, teacher: str, student: str) -> dict:
    path = REPORTS_ROOT / f"{dataset}_fewshot_fusion_{teacher}_{student}.csv"
    if not path.is_file():
        return {"path": path, "rows": {}, "error": "file missing"}
    out: dict[tuple[int, int, str], float] = {}
    duplicates = []
    for row in load_csv(path):
        try:
            key = (int(row["subject"]) - 1, int(row["seed"]), row["method"])
            if key in out:
                duplicates.append(key)
            out[key] = float(row["acc"])
        except (KeyError, ValueError):
            continue
    return {"path": path, "rows": out, "error": "duplicate key" if duplicates else ""}


def fewshot_analysis(summary: list[dict], missing: list[dict]) -> tuple[dict, list[dict]]:
    fusion_cells: dict[tuple, dict] = {}
    gate_diag: list[dict] = []
    for display, dataset in DATASETS:
        for teacher, student in PAIRS:
            saved = read_fusion_csv(dataset, teacher, student)
            if saved["error"]:
                add_missing(missing, scope="fusion", protocol="fewshot",
                            dataset=dataset, teacher=teacher, student=student,
                            method="concat_mlp / gate_conf_acc",
                            seed_set="666,667,668", missing_item=saved["error"],
                            source_expected=str(saved["path"]))
            saved_rows = saved["rows"]
            possible = sorted({(key[0], key[1]) for key in saved_rows})
            if not possible:
                continue
            per_method: dict[str, list[tuple[int, int, float]]] = defaultdict(list)
            method_units: dict[str, dict[tuple[int, int], float]] = defaultdict(dict)
            gate_units: dict[tuple[int, int], float] = {}
            complementarity: dict[str, list[tuple[int, int, float]]] = defaultdict(list)
            complementarity_counts: dict[str, int] = defaultdict(int)
            confusion_denominator = 0
            for subject_key, seed in possible:
                try:
                    cell = load_fewshot_pair(dataset, teacher, student,
                                             subject_key, seed)
                except (FileNotFoundError, ValueError, KeyError) as exc:
                    add_missing(missing, scope="fusion", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="teacher/student logits, features, y, sample_uid",
                                seed_set=str(seed), missing_item="unaligned or absent artifact",
                                detail=str(exc))
                    continue
                y = cell["y_te"]
                tlog = np.asarray(cell["big_te"]["logits"])
                slog = np.asarray(cell["small_te"]["logits"])
                tpred, spred = tlog.argmax(axis=1), slog.argmax(axis=1)
                probs_prob = (softmax(tlog) + softmax(slog)) / 2.0
                probs_logit = softmax((tlog + slog) / 2.0)
                probs_gate, w_big, gate_info = gate_arrays(
                    cell["big_tr"]["logits"], cell["small_tr"]["logits"],
                    cell["y_tr"], tlog, slog)
                metrics = {
                    "teacher_only": percent_accuracy(tlog, y),
                    "student_only": percent_accuracy(slog, y),
                    "mean_prob": float(np.mean(probs_prob.argmax(axis=1) == y) * 100.0),
                    "mean_logit": float(np.mean(probs_logit.argmax(axis=1) == y) * 100.0),
                    "gate_conf_acc": float(np.mean(probs_gate.argmax(axis=1) == y) * 100.0),
                }
                concat = saved_rows.get((subject_key, seed, "concat_mlp"))
                saved_gate = saved_rows.get((subject_key, seed, "gate_conf_acc"))
                if concat is None or saved_gate is None:
                    add_missing(missing, scope="fusion", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="concat_mlp / gate_conf_acc",
                                seed_set=str(seed), missing_item="fusion output cell",
                                detail=f"subject_key={subject_key}")
                    continue
                metrics["concat_mlp"] = float(concat)
                metrics["gate_conf_acc"] = float(saved_gate)
                recomputed_gate = float(np.mean(probs_gate.argmax(axis=1) == y) * 100.0)
                if abs(recomputed_gate - float(saved_gate)) > 0.011:
                    add_missing(missing, scope="fusion", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="gate_conf_acc", seed_set=str(seed),
                                missing_item="recomputed gate differs from saved output",
                                detail=f"subject_key={subject_key}; computed={recomputed_gate:.4f}; saved={saved_gate:.4f}",
                                source_expected=str(saved["path"]))
                unit_key = (seed, subject_key)
                for method, value in metrics.items():
                    method_units[method][unit_key] = value
                gate_units[unit_key] = float(saved_gate)
                y_ok_t, y_ok_s = tpred == y, spred == y
                states = {
                    "both_correct": y_ok_t & y_ok_s,
                    "teacher_only_correct": y_ok_t & ~y_ok_s,
                    "student_only_correct": ~y_ok_t & y_ok_s,
                    "both_wrong": ~y_ok_t & ~y_ok_s,
                }
                confusion_denominator += len(y)
                for state, mask in states.items():
                    rate = float(mask.mean() * 100.0)
                    complementarity[state].append((seed, subject_key, rate))
                    complementarity_counts[state] += int(mask.sum())
                oracle = float((y_ok_t | y_ok_s).mean() * 100.0)
                complementarity["oracle_accuracy"].append((seed, subject_key, oracle))
                gate_diag.append({
                    "dataset": dataset, "teacher": teacher, "student": student,
                    "seed": seed, "subject": subject_key + 1,
                    "teacher_train_acc_pct": gate_info["train_acc_teacher_pct"],
                    "student_train_acc_pct": gate_info["train_acc_student_pct"],
                    "w_teacher": w_big,
                })
            if not method_units.get("gate_conf_acc"):
                continue
            for method in ("teacher_only", "student_only", "mean_prob", "mean_logit",
                           "concat_mlp", "gate_conf_acc"):
                unit_map = method_units.get(method, {})
                values = [(seed, subject, value)
                          for (seed, subject), value in unit_map.items()]
                gate_common = [(seed, subject, value - gate_units[(seed, subject)])
                               for (seed, subject), value in unit_map.items()
                               if (seed, subject) in gate_units]
                add_summary(
                    summary, analysis="A", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student, method=method,
                    metric="accuracy_pct",
                    values=values, delta_values=gate_common,
                    detail="mean across subjects within each seed, then mean±sample SD across seeds; paired delta uses same subject×seed cells",
                    source=f"artifacts + {saved['path'].name}")
                fusion_cells[(dataset, teacher, student, method)] = values
                definition = {
                    "teacher_only": "teacher argmax(logits)",
                    "student_only": "student argmax(logits)",
                    "mean_prob": "argmax(0.5×softmax(teacher logits)+0.5×softmax(student logits))",
                    "mean_logit": "argmax(0.5×teacher logits+0.5×student logits)",
                    "concat_mlp": "saved concat MLP output; no head retraining in this analysis",
                    "gate_conf_acc": "saved gate output; gate recomputed from artifacts for validation",
                }[method]
                add_summary(
                    summary, analysis="B", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student, method=method,
                    metric="accuracy_pct", values=values,
                    detail=f"accuracy context for complementarity/oracle comparison; {definition}",
                    source=f"artifacts + {saved['path'].name}")
            for state in ("both_correct", "teacher_only_correct",
                          "student_only_correct", "both_wrong"):
                values = complementarity.get(state, [])
                add_summary(
                    summary, analysis="B", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student, method="teacher_vs_student",
                    metric=f"{state}_pct_of_test_trials", values=values,
                    pooled_count=complementarity_counts[state],
                    pooled_denominator=confusion_denominator,
                    detail="2×2 correctness cell; value aggregates subject rates within seed; pooled counts include every aligned test trial",
                    source="few-shot teacher/student test logits and labels")
            add_summary(
                summary, analysis="B", protocol="fewshot", dataset=dataset,
                teacher=teacher, student=student, method="oracle_teacher_or_student",
                metric="accuracy_pct", values=complementarity.get("oracle_accuracy", []),
                detail="per-trial oracle is correct when either teacher or student is correct; not a deployable predictor",
                source="few-shot teacher/student test logits and labels")
    return fusion_cells, gate_diag


def uid_keys(values: np.ndarray) -> list:
    array = np.asarray(values)
    if array.ndim == 2:
        return [tuple(item.tolist()) for item in array]
    if array.ndim != 1:
        raise ValueError(f"UID must have 1D or 2D shape, got {array.shape}")
    if np.issubdtype(array.dtype, np.number):
        return [(item.item(),) for item in array]
    return [str(item) for item in array.tolist()]


def align_to_reference(logits: np.ndarray, labels: np.ndarray | None,
                       uid: np.ndarray | None, uid_hash: str,
                       ref_uid: np.ndarray, ref_y: np.ndarray) -> np.ndarray:
    """Return logits in reference UID order, rejecting positional assumptions."""
    logits = np.asarray(logits)
    ref_keys = uid_keys(ref_uid)
    if uid is not None:
        keys = uid_keys(uid)
        if len(set(keys)) != len(keys) or set(keys) != set(ref_keys):
            raise ValueError("test UID set differs from reference artifact")
        row_for_uid = {key: index for index, key in enumerate(keys)}
        order = np.asarray([row_for_uid[key] for key in ref_keys], dtype=np.int64)
        logits = logits[order]
        if labels is not None and not np.array_equal(np.asarray(labels)[order], ref_y):
            raise ValueError("test labels disagree after UID alignment")
        return logits
    if not uid_hash or sha256_array(ref_uid) != uid_hash:
        raise ValueError("no explicit UID and test_uid_hash does not match reference UID order")
    if len(logits) != len(ref_y):
        raise ValueError("test logits length differs from reference labels")
    if labels is not None and not np.array_equal(np.asarray(labels), ref_y):
        raise ValueError("test labels disagree with reference artifact")
    return logits


def row_subject_key(row: dict[str, str]) -> int | None:
    key = int_field(row, "subject_index", "key")
    if key is not None:
        return key
    subject = int_field(row, "subject")
    return subject - 1 if subject is not None else None


def source_series(relative: Path) -> str:
    text = relative.as_posix()
    if text == "seed666_three_ablation" or re.fullmatch(r"three_seed_ablation_seed\d+", text):
        return "three_seed_ablation"
    if re.fullmatch(r"warmup_mi_combo_seed\d+", text):
        return "warmup_mi_combo"
    return re.sub(r"seed(?:ed)?_?\d+", "seed*", text)


def config_pair_for_dir(directory: Path) -> tuple[str, str]:
    candidates = [directory / "execution_provenance.json",
                  directory / "provenance.json"]
    config_path = ""
    for candidate in candidates:
        if candidate.is_file():
            try:
                data = json.loads(candidate.read_text(encoding="utf-8"))
                config_path = str(data.get("config_path", ""))
                break
            except (ValueError, OSError):
                pass
    if not config_path:
        return "", ""
    resolved = external_path(config_path)
    if not resolved.is_file():
        resolved = Path(config_path)
    if not resolved.is_file():
        return "", ""
    try:
        import yaml
        config = yaml.safe_load(resolved.read_text(encoding="utf-8")) or {}
    except Exception:
        return "", ""
    teacher = str(config.get("ft_teacher") or config.get("teacher") or "").lower()
    student = str(config.get("student") or "").lower()
    return teacher, student


def initial_hash(row: dict[str, str]) -> str:
    return field(row, "initial_state_hash", "initial_student_state_hash")


def condition_value(row: dict[str, str]) -> str:
    return field(row, "condition", "method", "condition_label")


def method_class(condition: str) -> str | None:
    key = condition.upper()
    if key in {"KD_ALL", "KD"}:
        return "KL"
    if key in {"KD_MASKED", "KD_MASK"}:
        return "KL_mask"
    if key == "MMD":
        return "MMD"
    if key == "KD_MMD":
        return "KD+MMD"
    if key in {"TASK_FEAT_LOGIT_KD", "PREALIGN_THEN_TASK_FEAT_LOGIT_KD"}:
        return "feature_alignment"
    if key == "PREALIGN_ONLY":
        return "feature_prealign_only"
    if key in {"WARMUP10", "WARMUP10_KD_MI", "DELAYED_KD", "DELAYED_KD_ALL",
               "DELAYED_KD_TCORRECT", "DELAYED_LOGIT_KD", "PREALIGN_THEN_KD"}:
        return "warmup_or_delayed_KD"
    if key in {"KD_MI_ALL", "KD_PROTO", "KD_MI_PROTO", "KD_MI_TCORRECT_MASK"}:
        return "specialized_KD_variant"
    return None


def is_ce_control(condition: str) -> bool:
    return condition.strip().casefold() in {"base", "base_ce", "ce"}


def collect_kd_rows(missing: list[dict]) -> list[dict]:
    root = RESULTS_ROOT / "distill"
    indexed: list[dict] = []
    if not root.is_dir():
        add_missing(missing, scope="KD", protocol="fewshot", missing_item="distillation result root missing",
                    detail=str(root))
        return indexed
    for csv_path in sorted(root.rglob("results_per_run.csv")):
        if "loso" in csv_path.relative_to(root).parts:
            continue
        rel_dir = csv_path.parent.relative_to(root)
        series = source_series(rel_dir)
        fallback_teacher, fallback_student = config_pair_for_dir(csv_path.parent)
        try:
            csv_rows = load_csv(csv_path)
        except OSError as exc:
            add_missing(missing, scope="KD", protocol="fewshot", missing_item="cannot read per-run manifest",
                        detail=f"{csv_path}: {exc}")
            continue
        for row in csv_rows:
            dataset = canonical_dataset(field(row, "dataset"))
            if dataset not in DISPLAY_BY_CANONICAL:
                continue
            protocol = field(row, "protocol") or "fewshot"
            if protocol.casefold() != "fewshot":
                continue
            condition = condition_value(row)
            category = method_class(condition)
            if not is_ce_control(condition) and category is None:
                continue
            seed = int_field(row, "seed")
            subject_key = row_subject_key(row)
            if seed is None or subject_key is None:
                continue
            teacher = field(row, "teacher", "ft_teacher").lower() or fallback_teacher
            student = field(row, "student").lower() or fallback_student
            if student not in STUDENTS:
                continue
            if teacher not in TEACHERS and not (is_ce_control(condition) and teacher == "shared"):
                if is_ce_control(condition):
                    teacher = "shared"
                else:
                    add_missing(missing, scope="KD", protocol="fewshot", dataset=dataset,
                                method=condition, seed_set=str(seed), missing_item="teacher metadata absent",
                                detail="row and provenance config do not identify the teacher",
                                source_expected=str(csv_path))
                    continue
            if not is_ce_control(condition) and teacher not in TEACHERS:
                continue
            status = field(row, "status").casefold()
            if status and status not in {"complete", "completed", "ok"}:
                add_missing(missing, scope="KD", protocol="fewshot", dataset=dataset,
                            teacher=teacher, student=student, method=condition,
                            seed_set=str(seed), missing_item="run not complete",
                            detail=status, source_expected=str(csv_path))
                continue
            pred_path_text = field(row, "prediction_path")
            pred_path = external_path(pred_path_text) if pred_path_text else None
            if pred_path is None or not pred_path.is_file():
                if not is_ce_control(condition) and category:
                    add_missing(missing, scope="KD", protocol="fewshot", dataset=dataset,
                                teacher=teacher, student=student, method=condition,
                                seed_set=str(seed), missing_item="test logits artifact missing",
                                detail=pred_path_text, source_expected=str(csv_path))
                continue
            indexed.append({
                "dataset": dataset, "teacher": teacher, "student": student,
                "condition": condition, "category": category or "CE_control",
                "seed": seed, "subject_key": subject_key,
                "subject": subject_key + 1, "series": series,
                "row": row, "csv_path": csv_path, "prediction_path": pred_path,
                "uid_hash": field(row, "test_uid_hash", "test_uid_sha256"),
            })
    return indexed


def load_kd_prediction(record: dict) -> dict:
    arrays = npz_dict(record["prediction_path"], KD_TEST_ARRAY_FIELDS)
    logits = next((arrays[name] for name in ("test_logits", "logits") if name in arrays), None)
    if logits is None:
        raise ValueError("prediction NPZ has no test_logits")
    uid = next((arrays[name] for name in ("test_uid", "test_sample_uid", "trial_uid", "sample_uid")
                if name in arrays), None)
    y = next((arrays[name] for name in ("test_y", "y", "labels") if name in arrays), None)
    return {"logits": np.asarray(logits), "uid": uid,
            "y": np.asarray(y) if y is not None else None,
            "arrays": arrays}


def find_matching_ce(record: dict, all_records: list[dict]) -> dict | None:
    candidates = [r for r in all_records if r["category"] == "CE_control"
                  and r["dataset"] == record["dataset"]
                  and r["student"] == record["student"]
                  and r["seed"] == record["seed"]
                  and r["subject_key"] == record["subject_key"]]
    target_uid_hash = record["uid_hash"]
    target_init = initial_hash(record["row"])
    target_train_uid = field(record["row"], "train_uid_hash", "split_uid_hash")
    ranked = []
    for candidate in candidates:
        row = candidate["row"]
        if target_uid_hash and candidate["uid_hash"] and target_uid_hash != candidate["uid_hash"]:
            continue
        candidate_init = initial_hash(row)
        same_series = candidate["series"] == record["series"]
        same_teacher = candidate["teacher"] in {record["teacher"], "shared"}
        if target_init and candidate_init:
            if target_init != candidate_init:
                continue
            identity_match = 3
        elif same_series:
            identity_match = 2
        else:
            continue
        candidate_train_uid = field(row, "train_uid_hash", "split_uid_hash")
        if target_train_uid and candidate_train_uid and target_train_uid != candidate_train_uid:
            continue
        same_batch = (field(row, "batch_order_hash", "batch_order_hashes") ==
                      field(record["row"], "batch_order_hash", "batch_order_hashes"))
        ranked.append((identity_match, int(same_series), int(same_teacher), int(same_batch), candidate))
    if not ranked:
        return None
    ranked.sort(key=lambda item: item[:-1], reverse=True)
    return ranked[0][-1]


def add_kd_group_rows(summary: list[dict], group_rows: list[dict], metric_key: str,
                      metric: str, detail: str, source: str,
                      denominator_key: str | None = None) -> None:
    by_seed: list[tuple[int, int, float]] = []
    pooled_num = pooled_den = 0
    for item in group_rows:
        value = item.get(metric_key)
        if value is None or not math.isfinite(float(value)):
            continue
        by_seed.append((item["seed"], item["subject"], float(value)))
        if denominator_key:
            pooled_num += int(item.get("numerator", 0))
            pooled_den += int(item.get(denominator_key, 0))
    if not by_seed:
        return
    first = group_rows[0]
    add_summary(summary, analysis="C", protocol=first["protocol"],
                dataset=first["dataset"], teacher=first["teacher"],
                student=first["student"], method=first["method"], metric=metric,
                values=by_seed, pooled_count=pooled_num if denominator_key else None,
                pooled_denominator=pooled_den if denominator_key else None,
                detail=detail, source=source)


def fewshot_kd_analysis(summary: list[dict], missing: list[dict]) -> list[dict]:
    records = collect_kd_rows(missing)
    teacher_cache: dict[tuple, dict] = {}
    loaded_predictions: dict[tuple, dict] = {}
    for record in records:
        try:
            loaded_predictions[id(record)] = load_kd_prediction(record)
        except (OSError, ValueError, KeyError) as exc:
            if record["category"] != "CE_control":
                add_missing(missing, scope="KD", protocol="fewshot",
                            dataset=record["dataset"], teacher=record["teacher"],
                            student=record["student"], method=record["condition"],
                            seed_set=str(record["seed"]), missing_item="invalid test logits artifact",
                            detail=f"{record['prediction_path']}: {exc}",
                            source_expected=str(record["csv_path"]))
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    for record in records:
        if record["category"] == "CE_control":
            continue
        pred_data = loaded_predictions.get(id(record))
        if pred_data is None:
            continue
        ce = find_matching_ce(record, records)
        if ce is None or id(ce) not in loaded_predictions:
            add_missing(missing, scope="KD", protocol="fewshot", dataset=record["dataset"],
                        teacher=record["teacher"], student=record["student"],
                        method=record["condition"], seed_set=str(record["seed"]),
                        missing_item="matched ordinary CE student logits",
                        detail="No CE control with matching split and student initialization was exported.",
                        source_expected=str(record["csv_path"]))
            continue
        ce_data = loaded_predictions[id(ce)]
        cache_key = (record["dataset"], record["teacher"], record["subject_key"], record["seed"])
        try:
            teacher_test = teacher_cache.get(cache_key)
            if teacher_test is None:
                teacher_test = load_light_artifact(record["dataset"], record["teacher"],
                                                   record["subject_key"], record["seed"], "test")
                if "sample_uid" not in teacher_test:
                    raise ValueError("teacher few-shot artifact lacks sample_uid")
                teacher_cache[cache_key] = teacher_test
            uid_ref, y_ref = teacher_test["sample_uid"], teacher_test["y"]
            kd_logits = align_to_reference(pred_data["logits"], pred_data["y"],
                                           pred_data["uid"], record["uid_hash"],
                                           uid_ref, y_ref)
            ce_logits = align_to_reference(ce_data["logits"], ce_data["y"],
                                           ce_data["uid"], ce["uid_hash"],
                                           uid_ref, y_ref)
        except (FileNotFoundError, ValueError, KeyError) as exc:
            add_missing(missing, scope="KD", protocol="fewshot", dataset=record["dataset"],
                        teacher=record["teacher"], student=record["student"],
                        method=record["condition"], seed_set=str(record["seed"]),
                        missing_item="test logits UID/label alignment failed",
                        detail=str(exc), source_expected=str(record["prediction_path"]))
            continue
        t_pred = np.asarray(teacher_test["logits"]).argmax(axis=1)
        kd_pred, ce_pred = kd_logits.argmax(axis=1), ce_logits.argmax(axis=1)
        agreement_delta = (kd_pred == t_pred).mean() - (ce_pred == t_pred).mean()
        ce_wrong_t_right = (ce_pred != y_ref) & (t_pred == y_ref)
        corrected = int(np.sum((kd_pred == y_ref) & ce_wrong_t_right))
        denominator = int(np.sum(ce_wrong_t_right))
        record_result = {
            "protocol": "fewshot", "dataset": record["dataset"],
            "teacher": record["teacher"], "student": record["student"],
            "method": f"{record['condition']} [{record['series']}]",
            "condition": record["condition"], "series": record["series"],
            "seed": record["seed"], "subject": record["subject"],
            "agreement_delta": float(agreement_delta * 100.0),
            "accuracy_delta": float((np.mean(kd_pred == y_ref) - np.mean(ce_pred == y_ref)) * 100.0),
            "correction_rate": (100.0 * corrected / denominator if denominator else None),
            "numerator": corrected, "denominator": denominator,
            "batch_order_match": bool(field(record["row"], "batch_order_hash") and
                                       field(record["row"], "batch_order_hash") ==
                                       field(ce["row"], "batch_order_hash")),
            "row": record["row"], "csv_path": record["csv_path"],
        }
        grouped[(record_result["dataset"], record_result["teacher"],
                 record_result["student"], record_result["method"])].append(record_result)
    for key, items in sorted(grouped.items()):
        dataset, teacher, student, method = key
        nseeds = sorted({item["seed"] for item in items})
        base_detail = (f"per-subject per-seed paired to exported CE control; seed set {seed_string(nseeds)}; "
                       f"batch-order exact in {sum(x['batch_order_match'] for x in items)}/{len(items)} units")
        source = "; ".join(sorted({str(item["csv_path"].parent.relative_to(RESULTS_ROOT / "distill"))
                                   for item in items}))
        add_kd_group_rows(summary, items, "agreement_delta",
                          "teacher_prediction_agreement_change_vs_CE_pp",
                          base_detail, source)
        add_kd_group_rows(summary, items, "accuracy_delta",
                          "student_accuracy_change_vs_CE_pp",
                          "KD/variant accuracy minus the matched CE accuracy, paired by fold and seed",
                          source)
        add_kd_group_rows(summary, items, "correction_rate",
                          "CE_wrong_teacher_right_corrected_pct",
                          "per-subject correction rate; pooled numerator/denominator count the CE-wrong, teacher-right trials",
                          source, denominator_key="denominator")
    return records


def condition_metric_value(record: dict) -> tuple[float | None, str]:
    """Select only explicit exported final training feature-cosine metrics."""
    row = record["row"]
    candidates = [
        ("transition_epoch100_feature_cosine", "train_feature_cosine_epoch100"),
        ("feature_cosine_epoch_100", "train_feature_cosine_epoch100"),
        ("final_ft_cosine", "train_feature_cosine_final_ft"),
    ]
    for column, name in candidates:
        value = csv_float(row.get(column))
        if value is not None:
            return value, name
    metrics_path = record["csv_path"].with_name("feature_transition_metrics.csv")
    if metrics_path.is_file():
        matches = []
        for metric_row in load_csv(metrics_path):
            if canonical_dataset(field(metric_row, "dataset")) != record["dataset"]:
                continue
            if int_field(metric_row, "subject", "subject_index") not in {
                record["subject"], record["subject_key"]
            }:
                continue
            student = field(metric_row, "student").lower()
            teacher = field(metric_row, "teacher").lower()
            condition = field(metric_row, "condition")
            if student and student != record["student"]:
                continue
            if teacher and teacher not in {record["teacher"], "shared"}:
                continue
            if condition and condition.upper() != record["condition"].upper():
                continue
            feature_column = next((name for name in ("feature_cosine", "ft_cosine")
                                   if csv_float(metric_row.get(name)) is not None), None)
            value = csv_float(metric_row.get(feature_column)) if feature_column else None
            epoch = int_field(metric_row, "epoch")
            if value is not None and epoch is not None:
                metric_name = "train_feature_cosine_ft" if feature_column == "ft_cosine" else "train_feature_cosine"
                matches.append((epoch, value, metric_name))
        if matches:
            epoch, value, metric_name = max(matches)
            return value, f"{metric_name}_epoch{epoch}"
    return None, ""


def feature_cosine_rows(summary: list[dict], missing: list[dict], records: list[dict]) -> None:
    feature_groups: dict[tuple, list[tuple[int, int, float]]] = defaultdict(list)
    missing_groups: set[tuple] = set()
    for record in records:
        if record["category"] != "feature_alignment" and record["condition"].upper() not in {
            "PREALIGN_ONLY", "PREALIGN_THEN_KD", "PREALIGN_THEN_TASK_FEAT_LOGIT_KD"
        }:
            continue
        value, metric = condition_metric_value(record)
        if value is None:
            missing_groups.add((record["dataset"], record["teacher"], record["student"],
                                record["condition"], record["series"]))
            continue
        key = (record["dataset"], record["teacher"], record["student"],
               record["condition"], record["series"], metric)
        feature_groups[key].append((record["seed"], record["subject"], value))
    for key, values in feature_groups.items():
        dataset, teacher, student, condition, series, metric = key
        add_summary(summary, analysis="C", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student,
                    method=f"{condition} [{series}]", metric=metric,
                    values=values,
                    detail="exported training-set cosine scalar from per-run metrics; no feature arrays were recomputed",
                    source=series)
        seeds = seed_string(seed for seed, _, _ in values)
        add_missing(missing, scope="KD feature", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student, method=condition,
                    seed_set=seeds, missing_item="test-set student feature vectors / cosine or CKA",
                    detail="Prediction NPZ exports test logits but no post-KD student test features; source-model features are not a substitute.",
                    source_expected=series)
    for dataset, teacher, student, condition, series in sorted(missing_groups):
        add_missing(missing, scope="KD feature", protocol="fewshot", dataset=dataset,
                    teacher=teacher, student=student, method=condition,
                    seed_set="available in run manifest", missing_item="training feature similarity scalar",
                    detail="No recognized final train cosine field in the exported per-run row.",
                    source_expected=series)


def discover_loso_files() -> dict[tuple, Path]:
    """Discover already-exported current LOSO predictions in registered runs."""
    roots_by_dataset = {
        "BNCI2014001": [RESULTS_ROOT / "reproductions/loso_config_alignment_v2/wideband_npy_v3"],
        "BNCI2014001-4": [RESULTS_ROOT / "reproductions/loso_config_alignment_v2/wideband_npy_v3"],
        "BNCI2014004": [RESULTS_ROOT / "reproductions/loso_source_refresh_004_5001_v1"],
        "BNCI2015001": [RESULTS_ROOT / "reproductions/loso_source_refresh_004_5001_v1"],
        "AlexMI": [RESULTS_ROOT / "reproductions/loso_five_settings_canonical4s_v1"],
    }
    result: dict[tuple, Path] = {}
    for dataset, roots in roots_by_dataset.items():
        for root in roots:
            if not root.is_dir():
                continue
            for pattern in ("result.npz", "test_predictions.npz"):
                for path in root.rglob(pattern):
                    parts = path.parts
                    path_dataset = next((canonical_dataset(part) for part in parts
                                         if canonical_dataset(part) in DISPLAY_BY_CANONICAL), None)
                    if path_dataset != dataset:
                        continue
                    model = next((part.lower() for part in parts if part.lower() in TEACHERS + STUDENTS), None)
                    if not model:
                        continue
                    subject = None
                    for part in parts:
                        found = re.fullmatch(r"subject_(\d+)", part)
                        if found:
                            subject = int(found.group(1))
                            break
                    seed = None
                    for part in parts:
                        found = re.fullmatch(r"seed_(\d+)", part)
                        if found:
                            seed = int(found.group(1))
                            break
                    if subject is None or seed is None:
                        continue
                    key = (dataset, model, seed, subject)
                    result.setdefault(key, path)
    return result


def select_uid(arrays: dict[str, np.ndarray]) -> np.ndarray | None:
    for name in ("trial_uid", "trial_uids", "sample_uid"):
        if name in arrays:
            return np.asarray(arrays[name])
    return None


def select_logits(arrays: dict[str, np.ndarray]) -> np.ndarray | None:
    for name in ("logits", "test_logits"):
        if name in arrays and np.asarray(arrays[name]).ndim == 2:
            return np.asarray(arrays[name])
    return None


def select_labels(arrays: dict[str, np.ndarray]) -> np.ndarray | None:
    for name in ("y", "labels", "test_y"):
        if name in arrays and np.asarray(arrays[name]).ndim == 1:
            return np.asarray(arrays[name], dtype=np.int64)
    return None


def uid_tokens_for_baseline(arrays: dict[str, np.ndarray], dataset: str,
                            model: str, root_path: Path) -> list:
    uid = select_uid(arrays)
    if uid is None:
        raise ValueError("LOSO prediction lacks per-trial UID")
    if uid.ndim == 2 or (uid.ndim == 1 and (uid.dtype.kind in "USO" or uid.dtype.kind in "iu")):
        return uid_keys(uid)
    raise ValueError("unsupported UID shape")


def reorder_pair_by_uid(big: dict, small: dict) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    uid_big, uid_small = select_uid(big), select_uid(small)
    if uid_big is None or uid_small is None:
        raise ValueError("teacher or student LOSO result lacks trial UIDs")
    keys_big, keys_small = uid_keys(uid_big), uid_keys(uid_small)
    if len(set(keys_big)) != len(keys_big) or len(set(keys_small)) != len(keys_small):
        raise ValueError("duplicate LOSO trial UIDs")
    if set(keys_big) != set(keys_small):
        raise ValueError("teacher/student LOSO UID sets differ")
    small_map = {key: i for i, key in enumerate(keys_small)}
    order = np.asarray([small_map[key] for key in keys_big], dtype=np.int64)
    by = select_labels(big)
    sy = select_labels(small)
    blogits, slogits = select_logits(big), select_logits(small)
    if by is None or sy is None or blogits is None or slogits is None:
        raise ValueError("missing LOSO logits or labels")
    if not np.array_equal(by, sy[order]):
        raise ValueError("teacher/student LOSO labels differ after UID alignment")
    return blogits, slogits[order], by, uid_big


def loso_analysis(summary: list[dict], missing: list[dict]) -> tuple[dict, dict]:
    files = discover_loso_files()
    pair_units: dict[tuple, dict] = defaultdict(dict)
    gate_cache = {}
    for display, dataset in DATASETS:
        for teacher, student in PAIRS:
            seeds = sorted({seed for ds, model, seed, _ in files
                            if ds == dataset and model == teacher} &
                           {seed for ds, model, seed, _ in files
                            if ds == dataset and model == student})
            valid_seed_units: dict[int, dict] = defaultdict(dict)
            for seed in seeds:
                subject_results = {}
                errors = []
                for subject in range(1, EXPECTED_SUBJECTS[dataset] + 1):
                    bp, sp = files.get((dataset, teacher, seed, subject)), files.get((dataset, student, seed, subject))
                    if bp is None or sp is None:
                        errors.append(f"subject {subject}: teacher/student result file missing")
                        continue
                    try:
                        big = npz_dict(bp, LOSO_ARRAY_FIELDS)
                        small = npz_dict(sp, LOSO_ARRAY_FIELDS)
                        blogits, slogits, y, uid = reorder_pair_by_uid(big, small)
                        pprob = (softmax(blogits) + softmax(slogits)) / 2.0
                        plogit = softmax((blogits + slogits) / 2.0)
                        values = {
                            "teacher_only": percent_accuracy(blogits, y),
                            "student_only": percent_accuracy(slogits, y),
                            "mean_prob": float(np.mean(pprob.argmax(axis=1) == y) * 100.0),
                            "mean_logit": float(np.mean(plogit.argmax(axis=1) == y) * 100.0),
                        }
                        subject_results[subject] = {
                            "values": values, "tlogits": blogits, "slogits": slogits,
                            "y": y, "uid": uid, "teacher_path": bp,
                            "student_path": sp,
                        }
                    except (OSError, ValueError, KeyError) as exc:
                        errors.append(f"subject {subject}: {exc}")
                if errors or len(subject_results) != EXPECTED_SUBJECTS[dataset]:
                    add_missing(missing, scope="fusion", protocol="loso", dataset=dataset,
                                teacher=teacher, student=student,
                                method="teacher/student/mean_prob/mean_logit",
                                seed_set=str(seed), missing_item="complete UID-aligned LOSO fold set",
                                detail="; ".join(errors[:5]) or f"found {len(subject_results)}/{EXPECTED_SUBJECTS[dataset]} subjects")
                    continue
                for subject, detail in subject_results.items():
                    valid_seed_units[seed][subject] = detail
            if not valid_seed_units:
                continue
            available_seeds = sorted(valid_seed_units)
            for method in ("teacher_only", "student_only", "mean_prob", "mean_logit"):
                vals = [(seed, subject, unit["values"][method])
                        for seed, subjects in valid_seed_units.items()
                        for subject, unit in subjects.items()]
                definition = {
                    "teacher_only": "teacher argmax(logits)",
                    "student_only": "student argmax(logits)",
                    "mean_prob": "argmax(0.5×softmax(teacher logits)+0.5×softmax(student logits))",
                    "mean_logit": "argmax(0.5×teacher logits+0.5×student logits)",
                }[method]
                add_summary(summary, analysis="A", protocol="loso", dataset=dataset,
                            teacher=teacher, student=student, method=method,
                            metric="accuracy_pct", values=vals,
                            detail=f"UID-aligned current LOSO held-out predictions; mean across subjects per seed then mean±sample SD across seeds; {definition}",
                            source="registered LOSO result.npz / test_predictions.npz")
                add_summary(summary, analysis="B", protocol="loso", dataset=dataset,
                            teacher=teacher, student=student, method=method,
                            metric="accuracy_pct", values=vals,
                            detail=f"accuracy context for complementarity/oracle comparison; {definition}",
                            source="registered LOSO result.npz / test_predictions.npz")
            # LOSO fusion result exports were not generated in the registered fusion output path.
            add_missing(missing, scope="fusion", protocol="loso", dataset=dataset,
                        teacher=teacher, student=student,
                        method="concat_mlp / gate_conf_acc",
                        seed_set=seed_string(available_seeds),
                        missing_item="existing LOSO concat/gate fusion output",
                        detail="No saved LOSO fusion CSV or same-fold fusion-head export was found; no fusion head was trained for this report.",
                        source_expected=str(RESULTS_ROOT / "fusion"))
            # Complementarity and oracle from exact UID-aligned LOSO folds.
            states: dict[str, list[tuple[int, int, float]]] = defaultdict(list)
            counts: dict[str, int] = defaultdict(int)
            denominator = 0
            for seed, subjects in valid_seed_units.items():
                for subject, unit in subjects.items():
                    t = unit["tlogits"].argmax(axis=1) == unit["y"]
                    s = unit["slogits"].argmax(axis=1) == unit["y"]
                    masks = {"both_correct": t & s, "teacher_only_correct": t & ~s,
                             "student_only_correct": ~t & s, "both_wrong": ~t & ~s}
                    denominator += len(t)
                    for name, mask in masks.items():
                        states[name].append((seed, subject, float(mask.mean() * 100.0)))
                        counts[name] += int(mask.sum())
                    states["oracle_accuracy"].append((seed, subject, float((t | s).mean() * 100.0)))
            for name in ("both_correct", "teacher_only_correct", "student_only_correct", "both_wrong"):
                add_summary(summary, analysis="B", protocol="loso", dataset=dataset,
                            teacher=teacher, student=student, method="teacher_vs_student",
                            metric=f"{name}_pct_of_test_trials", values=states[name],
                            pooled_count=counts[name], pooled_denominator=denominator,
                            detail="2×2 correctness cell; exact UID alignment; pooled counts across all folds and seeds",
                            source="registered LOSO result artifacts")
            add_summary(summary, analysis="B", protocol="loso", dataset=dataset,
                        teacher=teacher, student=student, method="oracle_teacher_or_student",
                        metric="accuracy_pct", values=states["oracle_accuracy"],
                        detail="per-trial oracle is correct when either teacher or student is correct; not deployable",
                        source="registered LOSO result artifacts")
            pair_units[(dataset, teacher, student)] = valid_seed_units
    return pair_units, files


def loso_uid_hash_for_current(dataset: str, student: str, seed: int,
                              subject: int, small_arrays: dict,
                              small_path: Path) -> str:
    local = small_arrays.get("local_sample_uid")
    if local is not None:
        local = np.asarray(local, dtype=np.int64)
        if local.ndim == 2 and local.shape[1] == 2:
            return sha256_array(local)
    # Some refreshed-source exports carry string trial UIDs only. Reconstruct
    # the loader's (zero-based subject, within-subject trial index) UID from the
    # already-exported model-input metadata, without loading X or EEG samples.
    input_root = (PROJECT_DATA_ROOT / "cache/artifact-loso/model_inputs" /
                  dataset / student)
    subjects_path = input_root / "subjects.npy"
    trials_path = input_root / "trials.csv"
    current_uid = select_uid(small_arrays)
    if subjects_path.is_file() and trials_path.is_file() and current_uid is not None:
        subjects = np.load(subjects_path, allow_pickle=False).astype(np.int64)
        trial_rows = load_csv(trials_path)
        if len(trial_rows) != len(subjects):
            raise ValueError("model-input subjects/trials row counts differ")
        counts: dict[int, int] = defaultdict(int)
        uid_to_local = {}
        for i, s in enumerate(subjects.tolist()):
            uid = field(trial_rows[i], "trial_uid")
            if not uid:
                raise ValueError("model-input trials.csv lacks trial_uid")
            uid_to_local[uid] = (int(s), counts[int(s)])
            counts[int(s)] += 1
        tokens = uid_keys(current_uid)
        if not tokens or not isinstance(tokens[0], str):
            raise ValueError("no local_sample_uid; current result must carry physical string UIDs")
        try:
            actual_order = np.asarray([uid_to_local[token] for token in tokens], dtype=np.int64)
        except KeyError as exc:
            raise ValueError(f"current result UID absent from model-input trials: {exc}") from exc
        if not np.all(actual_order[:, 0] == subject - 1):
            raise ValueError("current result contains UIDs outside its held-out subject")
        return sha256_array(actual_order)
    raise FileNotFoundError(f"no local_sample_uid or existing subjects metadata for {dataset}/{student}")


def teacher_checkpoint_matches(row: dict[str, str], teacher_path: Path,
                               dataset: str, teacher: str, seed: int,
                               subject: int) -> bool:
    expected_hash = field(row, "teacher_checkpoint_sha256")
    current_hash = ""
    for name in ("result.json", "manifest.json"):
        candidate = teacher_path.parent / name
        if not candidate.is_file():
            continue
        try:
            metadata = json.loads(candidate.read_text(encoding="utf-8"))
            current_hash = str(metadata.get("final_model_sha256", ""))
            if current_hash:
                break
        except (OSError, ValueError):
            pass
    if expected_hash and current_hash:
        return expected_hash == current_hash
    manifest_text = field(row, "teacher_target_manifest_path")
    if manifest_text:
        manifest_path = external_path(manifest_text)
        if manifest_path.is_file():
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                cache_hash = str(manifest.get("teacher_checkpoint_sha256", ""))
                if expected_hash and cache_hash and cache_hash != expected_hash:
                    return False
                checkpoint = manifest.get("teacher_checkpoint") or manifest.get("teacher_checkpoint_path")
                if isinstance(checkpoint, dict):
                    checkpoint_path = str(checkpoint.get("path", ""))
                    checkpoint_hash = str(checkpoint.get("sha256", ""))
                    if expected_hash and checkpoint_hash and checkpoint_hash != expected_hash:
                        return False
                    if checkpoint_path and teacher_path.parent.joinpath("model.pt").exists():
                        return external_path(checkpoint_path).resolve() == teacher_path.parent.joinpath("model.pt").resolve()
                    if checkpoint_path and teacher_path.parent.joinpath("final_model.pt").exists():
                        return external_path(checkpoint_path).resolve() == teacher_path.parent.joinpath("final_model.pt").resolve()
                if isinstance(checkpoint, str) and checkpoint:
                    return external_path(checkpoint).resolve() in {
                        (teacher_path.parent / "model.pt").resolve(),
                        (teacher_path.parent / "final_model.pt").resolve(),
                    }
            except (OSError, ValueError):
                return False
    # For AlexMI the teacher-target manifest identifies the same fold-specific
    # model.pt path; it is matched above without loading the model weights.
    return not expected_hash


def loso_kd_analysis(summary: list[dict], missing: list[dict],
                     baseline_files: dict[tuple, Path]) -> None:
    path = RESULTS_ROOT / "loso/distill/distill_kd_wide_fre_band.csv"
    if not path.is_file():
        add_missing(missing, scope="KD", protocol="loso", missing_item="KD_MMD per-trial export missing",
                    detail=str(path))
        return
    rows = load_csv(path)
    index = {}
    for row in rows:
        dataset = canonical_dataset(field(row, "dataset"))
        teacher, student = field(row, "teacher").lower(), field(row, "student").lower()
        seed, subject = int_field(row, "seed"), int_field(row, "subject")
        method = field(row, "method")
        if dataset in DISPLAY_BY_CANONICAL and teacher in TEACHERS and student in STUDENTS and seed is not None and subject is not None:
            index[(dataset, teacher, student, seed, subject, method)] = row
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    teacher_cache: dict[tuple, dict] = {}
    for dataset in [canonical for _, canonical in DATASETS]:
        for teacher, student in PAIRS:
            for seed in (666, 667, 668):
                for subject in range(1, EXPECTED_SUBJECTS[dataset] + 1):
                    ce = index.get((dataset, teacher, student, seed, subject, "Base"))
                    kd = index.get((dataset, teacher, student, seed, subject, "KD_MMD"))
                    if ce is None or kd is None:
                        continue
                    student_path = baseline_files.get((dataset, student, seed, subject))
                    teacher_path = baseline_files.get((dataset, teacher, seed, subject))
                    if student_path is None or teacher_path is None:
                        add_missing(missing, scope="KD", protocol="loso", dataset=dataset,
                                    teacher=teacher, student=student, method="KD_MMD",
                                    seed_set=str(seed), missing_item="UID-aligned current teacher/student test result",
                                    detail=f"subject {subject}")
                        continue
                    try:
                        current_student = npz_dict(student_path, LOSO_ARRAY_FIELDS)
                        current_teacher = npz_dict(teacher_path, LOSO_ARRAY_FIELDS)
                        uid_hash = loso_uid_hash_for_current(dataset, student, seed,
                                                             subject, current_student,
                                                             student_path)
                        if uid_hash != field(ce, "test_uid_hash") or uid_hash != field(kd, "test_uid_hash"):
                            raise ValueError("stored test_uid_hash differs from current held-out UID order")
                        if field(ce, "test_uid_hash") != field(kd, "test_uid_hash"):
                            raise ValueError("Base and KD_MMD test UID hashes differ")
                        if not teacher_checkpoint_matches(kd, teacher_path, dataset,
                                                          teacher, seed, subject):
                            raise ValueError("current teacher prediction checkpoint does not match KD teacher target")
                        t_logits, s_logits = select_logits(current_teacher), select_logits(current_student)
                        y_t, y_s = select_labels(current_teacher), select_labels(current_student)
                        t_uid, s_uid = select_uid(current_teacher), select_uid(current_student)
                        if t_logits is None or s_logits is None or y_t is None or y_s is None or t_uid is None or s_uid is None:
                            raise ValueError("current LOSO teacher/student result lacks logits, labels, or UIDs")
                        tkeys, skeys = uid_keys(t_uid), uid_keys(s_uid)
                        if set(tkeys) != set(skeys):
                            raise ValueError("current teacher/student UIDs or labels differ")
                        t_map = {key: i for i, key in enumerate(tkeys)}
                        order = np.asarray([t_map[key] for key in skeys], dtype=np.int64)
                        t_logits = t_logits[order]
                        y_t = y_t[order]
                        if not np.array_equal(y_t, y_s):
                            raise ValueError("current teacher/student labels differ after UID join")
                        ce_pred_path = external_path(field(ce, "prediction_path"))
                        kd_pred_path = external_path(field(kd, "prediction_path"))
                        ce_pred = npz_dict(ce_pred_path, KD_TEST_ARRAY_FIELDS)
                        kd_pred = npz_dict(kd_pred_path, KD_TEST_ARRAY_FIELDS)
                        ce_logits = select_logits(ce_pred)
                        kd_logits = select_logits(kd_pred)
                        if ce_logits is None or kd_logits is None:
                            raise ValueError("LOSO distillation prediction NPZ lacks test logits")
                        if len(ce_logits) != len(y_t) or len(kd_logits) != len(y_t):
                            raise ValueError("LOSO distillation test logits length mismatch")
                        tp = t_logits.argmax(axis=1)
                        cp, kp = ce_logits.argmax(axis=1), kd_logits.argmax(axis=1)
                        agree_delta = (kp == tp).mean() - (cp == tp).mean()
                        ce_wrong_t_right = (cp != y_t) & (tp == y_t)
                        n = int(ce_wrong_t_right.sum())
                        corrected = int(((kp == y_t) & ce_wrong_t_right).sum())
                        item = {
                            "protocol": "loso", "dataset": dataset, "teacher": teacher,
                            "student": student, "method": "KD_MMD [distill_kd_wide_fre_band]",
                            "condition": "KD_MMD", "series": "distill_kd_wide_fre_band",
                            "seed": seed, "subject": subject,
                            "agreement_delta": float(agree_delta * 100.0),
                            "accuracy_delta": float((np.mean(kp == y_t) - np.mean(cp == y_t)) * 100.0),
                            "correction_rate": 100.0 * corrected / n if n else None,
                            "numerator": corrected, "denominator": n,
                        }
                        grouped[(dataset, teacher, student, item["method"])].append(item)
                    except (OSError, ValueError, KeyError) as exc:
                        add_missing(missing, scope="KD", protocol="loso", dataset=dataset,
                                    teacher=teacher, student=student, method="KD_MMD",
                                    seed_set=str(seed), missing_item="KD_MMD test UID/logit/teacher identity validation failed",
                                    detail=f"subject {subject}: {exc}",
                                    source_expected=field(kd, "prediction_path"))
    for key, items in sorted(grouped.items()):
        dataset, teacher, student, method = key
        source = str(path)
        detail = "KD_MMD test logits vs same-run Base CE; teacher logits from matching fold/seed checkpoint; exact test UID hash checked"
        add_kd_group_rows(summary, items, "agreement_delta",
                          "teacher_prediction_agreement_change_vs_CE_pp", detail, source)
        add_kd_group_rows(summary, items, "accuracy_delta",
                          "student_accuracy_change_vs_CE_pp",
                          "KD_MMD accuracy minus same-run Base CE accuracy, paired by fold/seed",
                          source)
        add_kd_group_rows(summary, items, "correction_rate",
                          "CE_wrong_teacher_right_corrected_pct",
                          "per-fold correction rate on CE-wrong, teacher-right trials; exact UID validation",
                          source, denominator_key="denominator")


def fusion_implementation_rows(summary: list[dict], gate_diag: list[dict]) -> None:
    by_group: dict[tuple, list[dict]] = defaultdict(list)
    for item in gate_diag:
        by_group[(item["dataset"], item["teacher"], item["student"])].append(item)
    for (dataset, teacher, student), items in sorted(by_group.items()):
        for source_key, label in (("teacher_train_acc_pct", "gate_train_accuracy_teacher_pct"),
                                  ("student_train_acc_pct", "gate_train_accuracy_student_pct")):
            vals = [(r["seed"], r["subject"], r[source_key]) for r in items]
            add_summary(summary, analysis="D", protocol="fewshot", dataset=dataset,
                        teacher=teacher, student=student, method="gate_conf_acc",
                        metric=label, values=vals,
                        detail="accuracy on the exact same training/support artifact used by the fusion runner",
                        source="few-shot train logits and labels")
        w_metrics = {
            "gate_teacher_weight_mean": lambda x: float(np.mean(x)),
            "gate_teacher_weight_std_within_test": lambda x: float(np.std(x, ddof=1)) if len(x) > 1 else 0.0,
            "gate_teacher_weight_p10": lambda x: float(np.percentile(x, 10)),
            "gate_teacher_weight_median": lambda x: float(np.percentile(x, 50)),
            "gate_teacher_weight_p90": lambda x: float(np.percentile(x, 90)),
        }
        for metric, fn in w_metrics.items():
            vals = [(r["seed"], r["subject"], fn(r["w_teacher"])) for r in items]
            add_summary(summary, analysis="D", protocol="fewshot", dataset=dataset,
                        teacher=teacher, student=student, method="gate_conf_acc",
                        metric=metric, values=vals,
                        detail="teacher weight w_big over test trials; alpha=beta=1, temperature=1",
                        source="recomputed from exported train/test logits")
    summary.append({
        "analysis": "D", "protocol": "fewshot", "dataset": "all",
        "teacher": "all", "student": "all", "method": "concat_mlp",
        "metric": "implementation_check", "seed_set": "varies by saved fusion CSV",
        "n_seeds": None, "n_subjects": None, "n_units": None,
        "value_mean": None, "value_std": None, "delta_vs_gate_mean": None,
        "delta_vs_gate_std": None, "pooled_count": None, "pooled_denominator": None,
        "detail": ("collab.fusion.train_concat_mlp standardizes from training features, then runs a fixed epoch loop; "
                   "no validation set and no early stopping. run_fusion passes the same train/support artifact split "
                   "used to fine-tune the teacher/student. Existing CSVs do not save resolved per-run concat hyperparameters."),
        "source": "experiments/fusion/run_fusion.py + collab/fusion.py + saved fusion CSV schema",
    })


def missing_method_coverage(missing: list[dict], kd_records: list[dict]) -> None:
    # Explicitly enumerate absent standard methods across the requested few-shot grid.
    available = defaultdict(set)
    for record in kd_records:
        if record["category"] == "CE_control":
            continue
        method = record["condition"].upper()
        if method in {"KD_ALL", "KD"}:
            category = "KL"
        elif method in {"KD_MASKED", "KD_MASK"}:
            category = "KL_mask"
        elif method == "MMD":
            category = "MMD"
        elif method == "KD_MMD":
            category = "KD+MMD"
        elif method in {"TASK_FEAT_LOGIT_KD", "PREALIGN_THEN_TASK_FEAT_LOGIT_KD"}:
            category = "feature_alignment"
        elif "WARMUP" in method or "DELAYED" in method or "PREALIGN_THEN_KD" in method:
            category = "warmup_or_delayed_KD"
        else:
            category = "specialized_KD_variant"
        available[(category, record["dataset"], record["teacher"], record["student"])].add(record["seed"])
    for category in ("KL", "KL_mask", "MMD", "feature_alignment", "warmup_or_delayed_KD"):
        for _, dataset in DATASETS:
            for teacher, student in PAIRS:
                seeds = sorted(available.get((category, dataset, teacher, student), set()))
                if category == "KL_mask" and not seeds:
                    add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="KD_masked", seed_set="666,667,668",
                                missing_item="per-trial test logits export",
                                detail="No standard KD_masked condition was found in the exported per-run manifests; specialized teacher-correct/prototype masks are listed separately by exact condition.")
                elif category == "MMD" and not seeds:
                    add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="MMD", seed_set="666,667,668",
                                missing_item="pure MMD per-trial test logits export",
                                detail="No few-shot pure MMD condition was found. KD+MMD LOSO exports are analyzed separately.")
                elif category == "KL" and not seeds:
                    add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="KL / KD_ALL", seed_set="666,667,668",
                                missing_item="per-trial test logits export",
                                detail="No ordinary KL export for this pair/dataset was found.")
                elif category == "feature_alignment" and not seeds:
                    add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="TASK_FEAT_LOGIT_KD", seed_set="666,667,668",
                                missing_item="per-trial test logits export",
                                detail="No feature-aligned KD export for this pair/dataset was found.")
                elif category == "warmup_or_delayed_KD" and not seeds:
                    add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                dataset=dataset, teacher=teacher, student=student,
                                method="warmup/delayed KD", seed_set="666,667,668",
                                missing_item="per-trial test logits export",
                                detail="No warmup/delayed KD export for this pair/dataset was found.")
                elif seeds and len(seeds) < 3:
                    missing_seeds = sorted(set(FEWSHOT_SEEDS) - set(seeds))
                    if missing_seeds:
                        add_missing(missing, scope="KD method coverage", protocol="fewshot",
                                    dataset=dataset, teacher=teacher, student=student,
                                    method=category, seed_set=seed_string(missing_seeds),
                                    missing_item="additional seeds not exported",
                                    detail=f"Available seeds: {seed_string(seeds)}")


def make_conclusions(summary: list[dict], missing: list[dict]) -> list[str]:
    lines = []
    fewshot_a = [r for r in summary if r["analysis"] == "A" and r["protocol"] == "fewshot"
                 and r["method"] in {"mean_prob", "mean_logit"}]
    if fewshot_a:
        diffs = [r for r in fewshot_a if r["delta_vs_gate_mean"] is not None]
        for method in ("mean_prob", "mean_logit"):
            vals = [r["delta_vs_gate_mean"] for r in diffs if r["method"] == method]
            if vals:
                lines.append(f"Few-shot {method} minus gate: {np.mean(vals):+.2f} pp average across {len(vals)} dataset×pair cells (n_seeds=3 per cell).")
    b_rows = [r for r in summary if r["analysis"] == "B" and r["protocol"] == "fewshot"
              and r["metric"] == "teacher_only_correct_pct_of_test_trials"]
    if b_rows:
        lines.append(f"Few-shot complementarity: teacher-only-correct trials average {np.mean([r['value_mean'] for r in b_rows if r['value_mean'] is not None]):.2f}% across {len(b_rows)} dataset×pair cells (n_seeds=3 per cell).")
    oracle = [r for r in summary if r["analysis"] == "B" and r["protocol"] == "fewshot"
              and r["method"] == "oracle_teacher_or_student"]
    if oracle:
        lines.append(f"Few-shot oracle accuracy averages {np.mean([r['value_mean'] for r in oracle if r['value_mean'] is not None]):.2f}% across {len(oracle)} dataset×pair cells (n_seeds=3 per cell); it is an upper bound, not a deployable method.")
    loso_a = [r for r in summary if r["analysis"] == "A" and r["protocol"] == "loso" and r["method"] == "mean_prob"]
    if loso_a:
        lines.append(f"LOSO mean_prob was calculated for {len(loso_a)} UID-aligned dataset×pair cells; n_seeds=3 for three datasets and 6 for 001/001-4.")
    kd_rows = [r for r in summary if r["analysis"] == "C" and r["metric"] == "teacher_prediction_agreement_change_vs_CE_pp"]
    if kd_rows:
        lines.append(f"KD agreement diagnostics include {len(kd_rows)} grouped condition rows with exported test logits and matched CE controls; n_seeds is recorded per row.")
    missing_categories = defaultdict(int)
    for row in missing:
        missing_categories[row["missing_item"]] += 1
    if missing_categories.get("per-trial test logits export") or missing_categories.get("pure MMD per-trial test logits export"):
        lines.append("No few-shot standard KD_masked or pure MMD test-logit export was found; missing pair/dataset cells are listed in the workbook.")
    if missing_categories.get("test-set student feature vectors / cosine or CKA"):
        lines.append("Feature-alignment exports provide some training cosine scalars, but no post-KD test feature vectors for test cosine/CKA.")
    lines.append("Concat MLP uses fixed-epoch training on frozen features with no validation/early stopping; saved fusion CSVs omit resolved hyperparameters.")
    if missing_categories.get("existing LOSO concat/gate fusion output"):
        lines.append("No saved LOSO concat/gate results were found; LOSO fusion-head training was not performed.")
    return lines[:10]


def write_workbook(path: Path, summary: list[dict], missing: list[dict], conclusions: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(SUMMARY_FIELDS)
    for row in summary:
        ws.append([row.get(name) for name in SUMMARY_FIELDS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if ws.max_row > 1:
        table = Table(displayName="ArtifactAnalysis", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False)
        ws.add_table(table)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in {"A": 12, "B": 12, "C": 12, "D": 14, "E": 12, "F": 32,
                       "G": 45, "H": 20, "I": 10, "J": 12, "K": 10, "L": 14,
                       "M": 13, "N": 16, "O": 15, "P": 16, "Q": 18,
                       "R": 68, "S": 48}.items():
        ws.column_dimensions[col].width = width
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (7, 18, 19))
        for col in (12, 13, 14, 15):
            row[col - 1].number_format = "0.00"

    wm = wb.create_sheet("missing_artifacts")
    wm.append(MISSING_FIELDS)
    for row in missing:
        wm.append([row.get(name) for name in MISSING_FIELDS])
    wm.freeze_panes = "A2"
    wm.auto_filter.ref = wm.dimensions
    if wm.max_row > 1:
        table = Table(displayName="MissingArtifacts", ref=wm.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False)
        wm.add_table(table)
    wm.sheet_view.showGridLines = False
    for cell in wm[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in {"A": 24, "B": 12, "C": 12, "D": 14, "E": 12,
                       "F": 34, "G": 20, "H": 42, "I": 78, "J": 58}.items():
        wm.column_dimensions[col].width = width
    for row in wm.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wc = wb.create_sheet("conclusions")
    wc.append(["#", "Conclusion"])
    for i, line in enumerate(conclusions, 1):
        wc.append([i, line])
    wc.column_dimensions["A"].width = 8
    wc.column_dimensions["B"].width = 140
    wc.freeze_panes = "A2"
    wc.sheet_view.showGridLines = False
    for cell in wc[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in wc.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_markdown(path: Path, workbook: Path, summary: list[dict],
                   missing: list[dict], conclusions: list[str]) -> None:
    sections = [
        "# Exported artifact analysis",
        "",
        f"Workbook: [{workbook.name}]({workbook.name})",
        "",
        "The `summary` worksheet is the single combined table; `missing_artifacts` records unavailable or unalignable exports. All accuracy/correction/agreement values are percentages or percentage points as named. Aggregation is equal-weight subjects within seed, then mean and sample SD across seeds. Seed lists and counts are recorded per row.",
        "",
        "## Conclusions",
        "",
    ]
    sections += [f"{i}. {line}" for i, line in enumerate(conclusions, 1)]
    sections += ["", f"Summary rows: {len(summary)}. Missing/unalignable artifact entries: {len(missing)}."]
    path.write_text("\n".join(sections) + "\n", encoding="utf-8")


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="results/exported_artifact_diagnostics.xlsx",
                        help="checkout-local workbook path (must be under results/)")
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    out = require_report_output(args.output)
    md = require_report_output(out.with_suffix(".md"))
    out.parent.mkdir(parents=True, exist_ok=True)
    summary: list[dict] = []
    missing: list[dict] = []
    print("[stage] few-shot fusion and complementarity", flush=True)
    fusion_cells, gate_diag = fewshot_analysis(summary, missing)
    print(f"[stage] few-shot complete: {len(summary)} summary rows", flush=True)
    print("[stage] UID-aligned LOSO baseline fusion analysis", flush=True)
    _, loso_files = loso_analysis(summary, missing)
    print(f"[stage] LOSO baseline complete: {len(summary)} summary rows", flush=True)
    print("[stage] few-shot KD transfer diagnostics", flush=True)
    kd_records = fewshot_kd_analysis(summary, missing)
    print(f"[stage] few-shot KD complete: {len(kd_records)} exported rows indexed", flush=True)
    feature_cosine_rows(summary, missing, kd_records)
    print("[stage] LOSO KD_MMD identity and transfer diagnostics", flush=True)
    loso_kd_analysis(summary, missing, loso_files)
    fusion_implementation_rows(summary, gate_diag)
    missing_method_coverage(missing, kd_records)
    conclusions = make_conclusions(summary, missing)
    summary.sort(key=lambda r: (str(r.get("analysis", "")), str(r.get("protocol", "")),
                                str(r.get("dataset", "")), str(r.get("teacher", "")),
                                str(r.get("student", "")), str(r.get("method", "")),
                                str(r.get("metric", ""))))
    write_workbook(out, summary, missing, conclusions)
    write_markdown(md, out, summary, missing, conclusions)
    print(f"[report] {out}")
    print(f"[conclusions] {md}")
    print(f"[summary_rows] {len(summary)} [missing_entries] {len(missing)}")
    for line in conclusions:
        print(f"[conclusion] {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
