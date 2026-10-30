"""CBraMod channel-aware MCC teacher adaptation and DiffKD IFNet LOSO pilot."""
from __future__ import annotations

import os
for _name in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS',
              'NUMEXPR_NUM_THREADS', 'BLIS_NUM_THREADS', 'VECLIB_MAXIMUM_THREADS',
              'BIGSMALLCOLLAB_CPU_THREADS', 'TORCH_NUM_THREADS',
              'TORCH_NUM_INTEROP_THREADS'):
    os.environ[_name] = '1'

# Configure PyTorch before NumPy/SciPy/model imports can start parallel work.
import torch
torch.set_num_threads(1)

import argparse
import copy
import csv
import hashlib
import json
from pathlib import Path
import random
import shutil
import subprocess
import sys
import time
import traceback

import numpy as np
import pandas as pd
from torch.nn import functional as F
import yaml

import config as project_config
from collab import cbramod_mcc_diffkd as channel_core
from collab import mcc_diffkd as common_core
from collab.student_feature_prior import apply_pca_whitener, fit_pca_whitener
from experiments.distill import run_student_aware_diffusion_loso as reference
from experiments.storage import (DATA_CACHE_ROOT, REPORTS_ROOT, RESULTS_ROOT,
                                 WEIGHTS_ROOT, require_external_output,
                                 require_report_output, resolve_local_file)
from models import get_adapter


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CONFIG = ROOT / 'configs/experiments/loso/cbramod_mcc_diffkd_loso_v1.yaml'
METHODS = ('CE_FEATURE', 'MCC_FEATURE', 'CE_DIFFKD', 'MCC_DIFFKD')
TEACHERS = ('T0', 'T_CE', 'T_MCC')
CONDITION_TEACHER = {'CE_FEATURE': 'T_CE', 'CE_DIFFKD': 'T_CE',
                     'MCC_FEATURE': 'T_MCC', 'MCC_DIFFKD': 'T_MCC'}
CONDITION_DIFFUSION = {'CE_FEATURE': False, 'MCC_FEATURE': False,
                       'CE_DIFFKD': True, 'MCC_DIFFKD': True}
RESOLVED_HASHES = {
    666: '6dfe184f7f6a5ddd8b75e57e6d67c77ce4e68f1bdf4aa349d558c0240539342f',
    667: '775d37f4ee707b0593270a5c25892135171790c4365b3717b900745bf0b5ffdf',
    668: '1f11b668046af82a4d3bbea0012f05e5fb8258862f383e668e37c8d25e79cde2',
}


def _sha(path) -> str:
    digest = hashlib.sha256()
    with Path(resolve_local_file(path)).open('rb') as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def _state_hash(state) -> str:
    return reference._state_hash(state)


def _json_hash(value) -> str:
    return reference._json_hash(value)


def _read_json(path: Path) -> dict:
    return reference._read_json(path)


def _atomic_json(path: Path, value: dict) -> None:
    reference._atomic_json(Path(path), value)


def _atomic_npz(path: Path, **arrays) -> None:
    reference._atomic_npz(Path(path), **arrays)


def _atomic_torch(path: Path, value: dict) -> None:
    reference._atomic_torch(Path(path), value)


def _atomic_csv(path: Path, rows: list[dict]) -> None:
    reference._atomic_csv(Path(path), rows)


def _write_progress(path: Path | None, value: dict) -> None:
    """Queue progress is checkout-local metadata, unlike training artifacts."""
    if path is None:
        return
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True) + '\n',
                         encoding='utf-8')
    os.replace(temporary, target)


def _read_config(path: Path) -> dict:
    with Path(resolve_local_file(path)).open(encoding='utf-8') as stream:
        return yaml.safe_load(stream)


def _validate_config(config: dict, seed: int) -> None:
    if seed not in (666, 667, 668) or seed not in config.get('seeds', []):
        raise ValueError(f'unsupported CBraMod pilot seed {seed}')
    if config['reference_run']['resolved_config_sha256_by_seed'].get(seed) != RESOLVED_HASHES[seed]:
        raise ValueError(f'seed {seed}: resolved IFNet reference hash differs from audited identity')
    if (config['experiment_id'] != 'CBRAMOD_MCC_DIFFKD_LOSO_V1'
            or config['dataset'] != 'BNCI2014004'
            or config['conditions'] != list(METHODS)
            or config['protocol']['folds'] != 9
            or config['feature_space']['channels'] != 3
            or config['diffusion']['ddim_timesteps'] != [250, 200, 150, 100, 50, 0]):
        raise ValueError('fixed CBraMod experiment configuration was changed')


def _device(gpu: int) -> torch.device:
    if not torch.cuda.is_available():
        raise RuntimeError('formal CBraMod MCC-DiffKD training requires CUDA')
    torch.cuda.set_device(int(gpu))
    return torch.device('cuda', int(gpu))


def _seed_all(seed: int, device: torch.device) -> None:
    random.seed(int(seed))
    np.random.seed(int(seed) % (2**32))
    torch.manual_seed(int(seed))
    if device.type == 'cuda':
        torch.cuda.manual_seed_all(int(seed))
    torch.set_num_threads(1)


def _roots(run_id: str, seed: int, fold: int) -> dict:
    roots = reference._run_roots(run_id, seed, fold)
    for key in ('run_result', 'run_cache', 'run_weights', 'run_seed_weights'):
        roots[key] = require_external_output(roots[key])
        roots[key].mkdir(parents=True, exist_ok=True)
    return roots


def _resolved(config: dict, seed: int):
    ref_config = dict(config['reference_run'])
    historical = dict(config)
    historical['seed'] = seed
    historical['reference_run'] = {
        'protocol_id': ref_config['protocol_id'], 'seed': seed,
        'resolved_config_sha256': ref_config['resolved_config_sha256_by_seed'][seed],
    }
    return reference._load_resolved_reference(historical)


def _check_cbramod_manifest(root: Path, identity: dict) -> dict:
    path = root / 'manifest.json'
    manifest = _read_json(path)
    if (manifest.get('dataset') != 'BNCI2014004'
            or manifest.get('model') != 'cbramod'
            or manifest.get('source_manifest_sha256') != identity['source_manifest_sha256']
            or manifest.get('selected_session') not in ('session_3', '3test')):
        raise RuntimeError(f'CBraMod model-input manifest does not match split source: {path}')
    for name in ('X.npy', 'y.npy', 'subjects.npy', 'trials.csv'):
        if not (root / name).is_file():
            raise FileNotFoundError(f'missing CBraMod input artifact {root / name}')
    shape = manifest.get('shape')
    if shape and tuple(shape[1:]) != (3, 4, 200):
        raise RuntimeError(f'expected cached CBraMod input [N,3,4,200], got {shape}')
    return manifest


def _load_data(identity: dict, roots: dict, student_cfg: dict, device: torch.device) -> dict:
    required = identity['uids']
    expected_y = {name: identity['labels'][name] for name in ('train', 'feedback')
                  if name in identity['labels']}
    input_root = roots['input_root']
    cbra_root, ifnet_root = input_root / 'cbramod', input_root / 'ifnet'
    cbra_manifest = _check_cbramod_manifest(cbra_root, identity)
    teacher_rows, teacher_info = reference._input_rows(
        cbra_root, required, expected_y, 'cbramod')
    student_rows, student_info = reference._input_rows(
        ifnet_root, required, expected_y, 'ifnet')
    cbra_source = np.load(resolve_local_file(cbra_root / 'X.npy'), mmap_mode='r',
                          allow_pickle=False)
    ifnet_source = np.load(resolve_local_file(ifnet_root / 'X.npy'), mmap_mode='r',
                           allow_pickle=False)
    if tuple(cbra_source.shape[1:]) != (3, 4, 200) or tuple(ifnet_source.shape[1:]) != (3, 1000):
        raise RuntimeError(f'unexpected cached input shapes: CBraMod={cbra_source.shape}, IFNet={ifnet_source.shape}')
    teacher_x = {part: np.asarray(cbra_source[teacher_rows[part]], dtype=np.float32)
                 for part in ('train', 'feedback')}
    student_adapter = get_adapter('ifnet', device='cpu', **student_cfg)
    student_x = {}
    for part in ('train', 'feedback'):
        raw = np.asarray(ifnet_source[student_rows[part]], dtype=np.float32)
        student_x[part] = reference._preprocess_student(student_adapter, raw)
    prepared = roots['split_dir'] / 'student_inputs.npz'
    if not prepared.is_file():
        raise FileNotFoundError(f'audited paired IFNet input cache is missing: {prepared}')
    with np.load(resolve_local_file(prepared), allow_pickle=False) as saved:
        for part in ('train', 'feedback'):
            if (not np.array_equal(saved[f'{part}_uids'], identity['uids'][part])
                    or not np.array_equal(saved[f'{part}_y'], identity['labels'][part])):
                raise RuntimeError(f'prepared IFNet {part} UID/labels differ from split')
            cached = np.asarray(saved[f'{part}_x'], dtype=np.float32)
            if cached.shape != student_x[part].shape or not np.allclose(
                    cached, student_x[part], rtol=1e-5, atol=1e-5):
                raise RuntimeError(f'prepared IFNet {part} inputs differ from verified source')
            student_x[part] = cached.copy()
    if not np.isfinite(teacher_x['train']).all() or not np.isfinite(student_x['train']).all():
        raise FloatingPointError('non-finite cached model input')
    return {
        'teacher_x': teacher_x, 'student_x': student_x,
        'teacher_rows': teacher_rows, 'student_rows': student_rows,
        'teacher_input_manifest': cbra_manifest,
        'student_input_manifest': student_info['manifest'],
        'teacher_trials_sha256': teacher_info['trials_sha256'],
        'student_trials_sha256': student_info['trials_sha256'],
        'teacher_input_mmap': cbra_source, 'student_input_mmap': ifnet_source,
        'teacher_input_root': cbra_root, 'student_input_root': ifnet_root,
        'student_adapter': student_adapter,
    }


def _metric_values(labels: np.ndarray, logits: np.ndarray) -> dict:
    from sklearn.metrics import (accuracy_score, balanced_accuracy_score,
                                 cohen_kappa_score, confusion_matrix, log_loss)
    y = np.asarray(labels, dtype=np.int64)
    z = np.asarray(logits, dtype=np.float64)
    p = 1.0 / (1.0 + np.exp(-np.clip(z[:, 1] - z[:, 0], -50, 50)))
    pred = np.argmax(z, axis=1)
    return {
        'accuracy': float(accuracy_score(y, pred)),
        'balanced_accuracy': float(balanced_accuracy_score(y, pred)),
        'kappa': float(cohen_kappa_score(y, pred)),
        'cross_entropy': float(log_loss(y, np.column_stack((1-p, p)), labels=[0, 1])),
        'prediction_count_0': int((pred == 0).sum()),
        'prediction_count_1': int((pred == 1).sum()),
        'recall_0': float(confusion_matrix(y, pred, labels=[0, 1])[0, 0]
                          / max(1, (y == 0).sum())),
        'recall_1': float(confusion_matrix(y, pred, labels=[0, 1])[1, 1]
                          / max(1, (y == 1).sum())),
        'collapse': bool((pred == 0).all() or (pred == 1).all()),
    }


def _predict_teacher(model, x: np.ndarray, device: torch.device,
                     batch_size: int = 32, include_features: bool = False):
    model.eval()
    logits_out, features_out = [], []
    with torch.no_grad():
        for start in range(0, len(x), batch_size):
            xb = torch.as_tensor(x[start:start + batch_size], dtype=torch.float32, device=device)
            backbone = model.backbone(xb)
            if backbone.ndim != 4 or tuple(backbone.shape[1:]) != (3, 4, 200):
                raise RuntimeError(f'CBraMod backbone output shape changed: {tuple(backbone.shape)}')
            flat = model.flatten(backbone)
            logits = model.classifier(model.dropout(flat))
            logits_out.append(logits.float().cpu().numpy())
            if include_features:
                # Keep channel identity; pool only the four within-channel temporal patches.
                features_out.append(backbone.mean(dim=2).float().cpu().numpy())
    logits_np = np.concatenate(logits_out, axis=0)
    features_np = np.concatenate(features_out, axis=0) if include_features else None
    return features_np, logits_np


def _build_teacher(teacher_cfg: dict, device: torch.device,
                   seed: int) -> tuple[object, torch.nn.Module]:
    _seed_all(seed, device)
    adapter = get_adapter('cbramod', device=device, **teacher_cfg)
    model = adapter.build(2)
    model.eval()
    return adapter, model


def _fit_source_teacher(config: dict, identity: dict, roots: dict,
                        teacher_cfg: dict, data: dict, fingerprint: str,
                        device: torch.device, smoke: bool = False) -> tuple[object, torch.nn.Module, dict]:
    result_path = roots['run_result'] / 'teacher_T0.json'
    weight_path = roots['run_weights'] / 'teacher_T0.pt'
    seed = int(config['seed'])
    if result_path.is_file() or weight_path.is_file():
        if not (result_path.is_file() and weight_path.is_file()):
            raise RuntimeError('T0 checkpoint/manifest is incomplete')
        manifest = _read_json(result_path)
        if (manifest.get('fingerprint') != fingerprint
                or manifest.get('checkpoint_sha256') != _sha(weight_path)):
            raise RuntimeError('T0 checkpoint provenance/hash mismatch')
        adapter, model = _build_teacher(teacher_cfg, device,
                                        int(manifest['initialization_seed']))
        payload = torch.load(resolve_local_file(weight_path), map_location='cpu', weights_only=False)
        model.load_state_dict(payload['model'], strict=True)
        return adapter, model, manifest

    train_seed = int(config['randomness']['teacher_seed_base']) + 1009 * int(identity['manifest']['target_subject_id']) + seed
    adapter, model = _build_teacher(teacher_cfg, device, train_seed)
    initial_hash = _state_hash(model.state_dict())
    x = torch.as_tensor(data['teacher_x']['train'], dtype=torch.float32, device=device)
    y = torch.as_tensor(identity['labels']['train'], dtype=torch.long, device=device)
    epochs = 2 if smoke else int(teacher_cfg['epochs'])
    batch_size = int(teacher_cfg['batch_size'])
    n_batches = (len(x) + batch_size - 1) // batch_size
    lr_values = adapter._cosine_schedule(
        float(teacher_cfg['lr']), float(teacher_cfg['min_lr']), epochs, n_batches,
        warmup_epochs=int(teacher_cfg['warmup_epochs']))
    optimizer = torch.optim.AdamW(
        model.parameters(), lr=float(teacher_cfg['lr']),
        weight_decay=float(teacher_cfg['weight_decay']), eps=1e-8)
    train_rng = np.random.default_rng(train_seed + 1)
    history, step = [], 0
    model.train()
    criterion = torch.nn.CrossEntropyLoss(
        label_smoothing=float(teacher_cfg['label_smoothing']))
    for epoch in range(epochs):
        order = train_rng.permutation(len(x))
        epoch_loss, seen = 0.0, 0
        for start in range(0, len(order), batch_size):
            ids = torch.as_tensor(order[start:start + batch_size], dtype=torch.long, device=device)
            optimizer.zero_grad(set_to_none=True)
            _, logits = adapter.forward(model, x[ids])
            loss = criterion(logits, y[ids])
            if not torch.isfinite(loss):
                raise FloatingPointError(f'T0 non-finite loss at epoch {epoch + 1}')
            loss.backward()
            optimizer.step()
            count = len(ids)
            epoch_loss += float(loss.detach().cpu()) * count
            seen += count
            step += 1
        adapter._apply_lr_schedule(optimizer, lr_values, step)
        history.append({'epoch': epoch + 1, 'train_ce': epoch_loss / seen,
                        'lr_next': float(optimizer.param_groups[0]['lr']), 'n_train': seen})
        _atomic_csv(roots['run_result'] / 'teacher_T0_history.csv', history)
    model.eval()
    state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    _atomic_torch(weight_path, {'model': state, 'fingerprint': fingerprint,
                                'initial_model_sha256': initial_hash,
                                'final_model_sha256': _state_hash(state), 'seed': seed})
    manifest = {
        'status': 'complete', 'role': 'seven-source T0', 'fingerprint': fingerprint,
        'checkpoint': str(weight_path), 'checkpoint_sha256': _sha(weight_path),
        'model_sha256': _state_hash(state), 'initial_model_sha256': initial_hash,
        'initialization_seed': train_seed, 'pretrained_checkpoint': teacher_cfg['pretrain'],
        'pretrained_checkpoint_sha256': _sha(teacher_cfg['pretrain']),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'train_subject_ids': identity['manifest']['train_subject_ids'],
        'epochs': epochs, 'batch_size': batch_size,
        'optimizer': {'name': 'AdamW', 'lr': teacher_cfg['lr'],
                      'weight_decay': teacher_cfg['weight_decay'],
                      'label_smoothing': teacher_cfg['label_smoothing'],
                      'scheduler': teacher_cfg.get('lr_schedule')},
        'training_schedule': 'seeded per-epoch numpy permutations; DataLoader workers=0',
        'history': str(roots['run_result'] / 'teacher_T0_history.csv'),
    }
    _atomic_json(result_path, manifest)
    return adapter, model, manifest


def _adapt_teacher(name: str, config: dict, identity: dict, roots: dict,
                   teacher_cfg: dict, adapter, t0_model: torch.nn.Module, data: dict,
                   fingerprint: str, device: torch.device,
                   smoke: bool = False) -> tuple[torch.nn.Module, dict]:
    result_path = roots['run_result'] / f'teacher_{name}.json'
    weight_path = roots['run_weights'] / f'teacher_{name}.pt'
    if result_path.is_file() or weight_path.is_file():
        if not (result_path.is_file() and weight_path.is_file()):
            raise RuntimeError(f'{name} checkpoint/manifest is incomplete')
        manifest = _read_json(result_path)
        if (manifest.get('fingerprint') != fingerprint
                or manifest.get('checkpoint_sha256') != _sha(weight_path)):
            raise RuntimeError(f'{name} checkpoint identity/hash mismatch')
        model = copy.deepcopy(t0_model)
        payload = torch.load(resolve_local_file(weight_path), map_location='cpu', weights_only=False)
        model.load_state_dict(payload['model'], strict=True)
        model.eval()
        return model, manifest

    model = copy.deepcopy(t0_model).to(device)
    model.eval()  # fixed dropout/normalization behavior, with autograd enabled
    for parameter in model.parameters():
        parameter.requires_grad_(True)
        parameter.grad = None
    optimizer_cfg = config['teacher_adaptation']
    optimizer = torch.optim.Adam(
        model.parameters(), lr=float(optimizer_cfg['lr']),
        betas=tuple(optimizer_cfg['betas']), eps=float(optimizer_cfg['eps']),
        weight_decay=float(optimizer_cfg['weight_decay']))
    x = torch.as_tensor(data['teacher_x']['train'], dtype=torch.float32, device=device)
    y = torch.as_tensor(identity['labels']['train'], dtype=torch.long, device=device)
    epochs = 2 if smoke else int(optimizer_cfg['epochs'])
    batch_size = int(optimizer_cfg['batch_size'])
    schedule_seed = int(config['randomness']['adaptation_schedule_seed_base']) \
        + 997 * int(identity['manifest']['target_subject_id']) + int(config['seed'])
    schedule_rng = np.random.default_rng(schedule_seed)
    history = []
    model.eval()
    for epoch in range(epochs):
        order = schedule_rng.permutation(len(x))
        total_ce = total_mcc = total_loss = 0.0
        seen = 0
        for start in range(0, len(order), batch_size):
            ids = torch.as_tensor(order[start:start + batch_size], dtype=torch.long, device=device)
            optimizer.zero_grad(set_to_none=True)
            _, logits = adapter.forward(model, x[ids])
            ce = F.cross_entropy(logits, y[ids])
            if name == 'T_MCC':
                mcc, _ = common_core.minimum_class_confusion(
                    logits, temperature=float(optimizer_cfg['mcc']['temperature']),
                    entropy_epsilon=float(optimizer_cfg['mcc']['entropy_epsilon']),
                    row_epsilon=float(optimizer_cfg['mcc']['row_denominator_epsilon']))
                loss = ce + float(optimizer_cfg['mcc']['weight']) * mcc
            else:
                mcc = logits.new_zeros(())
                loss = ce
            if not torch.isfinite(loss):
                raise FloatingPointError(f'{name} non-finite objective at epoch {epoch + 1}')
            loss.backward()
            grad_norm = torch.nn.utils.clip_grad_norm_(model.parameters(),
                                                       float(optimizer_cfg['gradient_clip_norm']))
            if not torch.isfinite(grad_norm):
                raise FloatingPointError(f'{name} non-finite gradient at epoch {epoch + 1}')
            optimizer.step()
            count = len(ids)
            total_ce += float(ce.detach().cpu()) * count
            total_mcc += float(mcc.detach().cpu()) * count
            total_loss += float(loss.detach().cpu()) * count
            seen += count
        history.append({'teacher': name, 'epoch': epoch + 1,
                        'train_ce': total_ce / seen, 'train_mcc': total_mcc / seen,
                        'train_total': total_loss / seen, 'n_train': seen,
                        'schedule_seed': schedule_seed})
        _atomic_csv(roots['run_result'] / f'teacher_{name}_history.csv', history)
    model.eval()
    state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    _atomic_torch(weight_path, {'model': state, 'fingerprint': fingerprint,
                                'model_sha256': _state_hash(state), 'source_T0_sha256': _state_hash(t0_model.state_dict())})
    manifest = {
        'status': 'complete', 'teacher': name, 'fingerprint': fingerprint,
        'checkpoint': str(weight_path), 'checkpoint_sha256': _sha(weight_path),
        'model_sha256': _state_hash(state), 'source_T0_sha256': _state_hash(t0_model.state_dict()),
        'train_uid_sha256': identity['uid_sha256']['train'], 'epochs': epochs,
        'objective': 'CE' if name == 'T_CE' else 'CE + 1.0*MCC',
        'mcc_temperature': optimizer_cfg['mcc']['temperature'],
        'optimizer': {'name': 'Adam', 'lr': optimizer_cfg['lr'],
                      'betas': optimizer_cfg['betas'], 'eps': optimizer_cfg['eps'],
                      'weight_decay': optimizer_cfg['weight_decay'], 'scheduler': 'none'},
        'forward_mode': 'eval with autograd; dropout disabled; normalization fixed',
        'history': str(roots['run_result'] / f'teacher_{name}_history.csv'),
    }
    _atomic_json(result_path, manifest)
    return model, manifest


def _pca_for_fold(t0_features: np.ndarray, roots: dict, fingerprint: str,
                  config: dict) -> dict:
    path = roots['run_cache'] / 'fixed_T0_channel_pca.npz'
    manifest_path = roots['run_cache'] / 'fixed_T0_channel_pca.json'
    if path.is_file() or manifest_path.is_file():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete fixed CBraMod channel PCA artifacts')
        manifest = _read_json(manifest_path)
        if manifest.get('fingerprint') != fingerprint or manifest.get('file_sha256') != _sha(path):
            raise RuntimeError('fixed channel PCA identity/hash mismatch')
        with np.load(resolve_local_file(path), allow_pickle=False) as saved:
            pca = {key: np.asarray(saved[key]) for key in
                   ('mean', 'components', 'eigenvalues', 'scales')}
        pca.update(manifest['pca'])
        return {'pca': pca, 'manifest': manifest}
    flat = np.asarray(t0_features, dtype=np.float32).reshape(-1, 200)
    cfg = config['feature_space']
    pca = fit_pca_whitener(
        flat, max_dimensions=int(cfg['pca_max_dimensions']),
        rank_relative_threshold=float(cfg['rank_relative_threshold']),
        eigenvalue_floor_relative=float(cfg['whitening_eigenvalue_floor_relative']))
    if int(pca['dimension']) < 1:
        raise RuntimeError('CBraMod D_train channel-token PCA has zero effective rank')
    arrays = {key: pca[key] for key in ('mean', 'components', 'eigenvalues', 'scales')}
    _atomic_npz(path, **arrays)
    summary = {key: (int(value) if key in ('effective_rank', 'dimension') else float(value))
               for key, value in pca.items()
               if key not in ('mean', 'components', 'eigenvalues', 'scales')}
    manifest = {
        'fingerprint': fingerprint, 'file_sha256': _sha(path),
        'fit_partition': 'D_train only; all three channel tokens per trial',
        'feature_tap': 'CBraMod backbone [B,3,4,200], mean over patch axis only -> [B,3,200]',
        'train_trials': int(t0_features.shape[0]), 'train_channel_tokens': int(flat.shape[0]),
        'pca': summary,
    }
    _atomic_json(manifest_path, manifest)
    return {'pca': pca, 'manifest': manifest}


def _teacher_targets(identity: dict, roots: dict, data: dict, pca_record: dict,
                     teachers: dict[str, torch.nn.Module], adapter,
                     fingerprint: str, device: torch.device) -> tuple[dict, dict]:
    path = roots['run_cache'] / 'teacher_channel_targets.npz'
    manifest_path = roots['run_cache'] / 'teacher_channel_targets.json'
    if path.is_file() or manifest_path.is_file():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete CBraMod teacher target cache')
        manifest = _read_json(manifest_path)
        if (manifest.get('fingerprint') != fingerprint or manifest.get('file_sha256') != _sha(path)
                or manifest.get('train_uid_sha256') != identity['uid_sha256']['train']):
            raise RuntimeError('CBraMod teacher target cache identity/hash mismatch')
        with np.load(resolve_local_file(path), allow_pickle=False) as saved:
            arrays = {key: np.asarray(saved[key]) for key in saved.files}
        return arrays, manifest
    pca = pca_record['pca']
    arrays = {
        'train_uids': identity['uids']['train'], 'train_y': identity['labels']['train'],
        'feedback_uids': identity['uids']['feedback'], 'feedback_y': identity['labels']['feedback'],
    }
    t0_train, t0_train_logits = _predict_teacher(
        teachers['T0'], data['teacher_x']['train'], device, include_features=True)
    t0_feedback, t0_feedback_logits = _predict_teacher(
        teachers['T0'], data['teacher_x']['feedback'], device, include_features=True)
    arrays.update({'T0_train_features_native': t0_train,
                   'T0_feedback_features_native': t0_feedback,
                   'T0_train_logits': t0_train_logits,
                   'T0_feedback_logits': t0_feedback_logits})
    for name in ('T_CE', 'T_MCC'):
        for part in ('train', 'feedback'):
            native, logits = _predict_teacher(teachers[name], data['teacher_x'][part],
                                               device, include_features=True)
            whitened = apply_pca_whitener(native.reshape(-1, 200), pca).reshape(
                len(native), 3, int(pca['dimension'])).astype(np.float32)
            arrays[f'{name}_{part}_features_native'] = native.astype(np.float32)
            arrays[f'{name}_{part}_features_whitened'] = whitened
            arrays[f'{name}_{part}_logits'] = logits.astype(np.float32)
    _atomic_npz(path, **arrays)
    manifest = {
        'fingerprint': fingerprint, 'file_sha256': _sha(path),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'teacher_checkpoint_sha256': {
            name: _sha(roots['run_weights'] / f'teacher_{name}.pt') for name in TEACHERS},
        'pca_sha256': pca_record['manifest']['file_sha256'],
        'feature_tap': 'per-channel patch mean; no cross-channel pooling or flattening',
        'training_samples': 'D_train only; feedback used for diagnostics; no test target exported',
        'whitened_dimension': int(pca['dimension']),
    }
    _atomic_json(manifest_path, manifest)
    return arrays, manifest


def _channel_prior_init(dimension: int, seed: int, config: dict) -> tuple[dict, str]:
    diffusion = config['diffusion']
    with torch.random.fork_rng(devices=[]):
        torch.manual_seed(int(seed))
        model = channel_core.ChannelEpsilonPredictor(
            dimension, hidden_dimensions=int(diffusion['hidden_dimensions']),
            time_dimensions=int(diffusion['time_embedding_dimensions']),
            channel_dimensions=int(diffusion['channel_embedding_dimensions']), channels=3)
    state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    return state, _state_hash(state)


def _fit_channel_prior(name: str, config: dict, identity: dict, roots: dict,
                       targets: dict, fingerprint: str, initial_state: dict,
                       initial_hash: str, device: torch.device,
                       smoke: bool = False) -> tuple[torch.nn.Module, dict]:
    final_path = roots['run_weights'] / f'prior_{name}.pt'
    manifest_path = roots['run_result'] / f'prior_{name}.json'
    history_path = roots['run_result'] / f'prior_{name}_history.csv'
    features = np.asarray(targets[f'{name}_train_features_whitened'], dtype=np.float32)
    feedback = np.asarray(targets[f'{name}_feedback_features_whitened'], dtype=np.float32)
    dimension = int(features.shape[-1])
    diffusion = config['diffusion']
    updates = 4 if smoke else int(diffusion['training_updates'])
    method_fingerprint = _json_hash({'run_fingerprint': fingerprint, 'teacher': name,
        'initial_prior_sha256': initial_hash, 'updates': updates,
        'teacher_targets_sha256': _sha(roots['run_cache'] / 'teacher_channel_targets.npz')})
    if final_path.is_file() or manifest_path.is_file():
        if not (final_path.is_file() and manifest_path.is_file()):
            raise RuntimeError(f'{name} channel prior checkpoint/manifest is incomplete')
        manifest = _read_json(manifest_path)
        if manifest.get('method_fingerprint') != method_fingerprint or manifest.get('checkpoint_sha256') != _sha(final_path):
            raise RuntimeError(f'{name} channel prior checkpoint identity mismatch')
        model = channel_core.ChannelEpsilonPredictor(
            dimension, int(diffusion['hidden_dimensions']),
            int(diffusion['time_embedding_dimensions']),
            int(diffusion['channel_embedding_dimensions']), 3).to(device)
        payload = torch.load(resolve_local_file(final_path), map_location='cpu', weights_only=False)
        model.load_state_dict(payload['model'], strict=True)
        model.eval()
        for parameter in model.parameters():
            parameter.requires_grad_(False)
            parameter.grad = None
        return model, manifest
    model = channel_core.ChannelEpsilonPredictor(
        dimension, int(diffusion['hidden_dimensions']),
        int(diffusion['time_embedding_dimensions']),
        int(diffusion['channel_embedding_dimensions']), 3).to(device)
    model.load_state_dict(initial_state, strict=True)
    optimizer = torch.optim.AdamW(
        model.parameters(), lr=float(diffusion['lr']),
        betas=tuple(diffusion['betas']), eps=float(diffusion['eps']),
        weight_decay=float(diffusion['weight_decay']))
    _, alpha, sigma = channel_core.channel_ddpm_schedule(
        int(diffusion['steps']), float(diffusion['beta_start']),
        float(diffusion['beta_end']), device)
    train = torch.as_tensor(features.reshape(-1, dimension), dtype=torch.float32, device=device)
    seed = int(config['randomness']['diffusion_seed_base']) + 101 * int(identity['manifest']['target_subject_id']) + int(config['seed'])
    generator = torch.Generator(device=device).manual_seed(seed)
    batch_size = min(int(diffusion['batch_size_cap']), len(train))
    fixed_tables = {}
    for part, values in (('train', features), ('feedback', feedback)):
        rng = np.random.default_rng(seed + (0 if part == 'train' else 1))
        flat = values.reshape(-1, dimension)
        count = min(96, len(flat))
        ids = rng.choice(len(flat), count, replace=False)
        fixed_tables[part] = (ids, rng.integers(1, 1001, count),
                              rng.integers(0, 3, count),
                              rng.standard_normal((count, dimension)).astype(np.float32), flat)
    history = []
    def measure(which):
        ids, times, channels, eps, values = fixed_tables[which]
        t = torch.as_tensor(times, dtype=torch.long, device=device)
        c = torch.as_tensor(channels, dtype=torch.long, device=device)
        e = torch.as_tensor(eps, dtype=torch.float32, device=device)
        clean = torch.as_tensor(values[ids], dtype=torch.float32, device=device)
        noisy = alpha[t, None] * clean + sigma[t, None] * e
        with torch.no_grad():
            predicted = model(noisy, t, c)
            return float(F.mse_loss(predicted, e).cpu()), float(e.square().mean().cpu())
    train_mse, train_zero = measure('train')
    feedback_mse, feedback_zero = measure('feedback')
    history.append({'teacher': name, 'update': 0, 'train_mse': train_mse,
                    'train_zero_mse': train_zero, 'feedback_mse': feedback_mse,
                    'feedback_zero_mse': feedback_zero})
    model.train()
    for update in range(1, updates + 1):
        ids = torch.randint(len(train), (batch_size,), device=device, generator=generator)
        channels = ids.remainder(3)
        timesteps = torch.randint(1, int(diffusion['steps']) + 1, (batch_size,),
                                  dtype=torch.long, device=device, generator=generator)
        eps = torch.randn((batch_size, dimension), dtype=torch.float32,
                          device=device, generator=generator)
        noisy = alpha[timesteps, None] * train[ids] + sigma[timesteps, None] * eps
        prediction = model(noisy, timesteps, channels)
        loss = F.mse_loss(prediction, eps)
        if not torch.isfinite(loss):
            raise FloatingPointError(f'{name} prior non-finite noise loss at update {update}')
        optimizer.zero_grad(set_to_none=True)
        loss.backward()
        grad_norm = torch.nn.utils.clip_grad_norm_(model.parameters(), float(diffusion['gradient_clip_norm']))
        if not torch.isfinite(grad_norm):
            raise FloatingPointError(f'{name} prior non-finite gradient at update {update}')
        optimizer.step()
        if update % int(diffusion['diagnostic_interval']) == 0 or update == updates:
            train_mse, train_zero = measure('train')
            feedback_mse, feedback_zero = measure('feedback')
            history.append({'teacher': name, 'update': update,
                            'train_loss_batch': float(loss.detach().cpu()),
                            'gradient_norm': float(grad_norm.detach().cpu()),
                            'train_mse': train_mse, 'train_zero_mse': train_zero,
                            'feedback_mse': feedback_mse,
                            'feedback_zero_mse': feedback_zero})
            _atomic_csv(history_path, history)
    model.eval()
    for parameter in model.parameters():
        parameter.requires_grad_(False)
        parameter.grad = None
    state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    _atomic_torch(final_path, {'model': state, 'method_fingerprint': method_fingerprint,
                               'model_sha256': _state_hash(state)})
    manifest = {
        'status': 'complete', 'teacher': name, 'method_fingerprint': method_fingerprint,
        'fingerprint': fingerprint, 'checkpoint': str(final_path),
        'checkpoint_sha256': _sha(final_path), 'model_sha256': _state_hash(state),
        'initial_model_sha256': initial_hash, 'final_update': updates,
        'conditioning': 'CBraMod channel ID only; no class label, logits, or IFNet feature',
        'training_partition': 'D_train channel-token targets only',
        'noise_schedule': {'steps': diffusion['steps'], 'beta_start': diffusion['beta_start'],
                           'beta_end': diffusion['beta_end']},
        'history': str(history_path),
    }
    _atomic_json(manifest_path, manifest)
    return model, manifest


def _load_student_reference(config: dict, seed: int, fold: int,
                            identity: dict, roots: dict, resolved: dict,
                            cfg: dict, data: dict, device: torch.device) -> dict:
    student_cfg = dict(resolved['student_config'])
    adapter = get_adapter('ifnet', device=device, **student_cfg)
    from experiments.distill import sample_utility_protocol as pilot_protocol
    from experiments.distill import run_sample_utility_adaptive_rl_loso as pilot_runner
    schedule = pilot_protocol.epoch_schedule(
        len(identity['uids']['train']), int(student_cfg['batch_size']),
        int(student_cfg['epochs']), seed)
    schedule_sha = pilot_runner._schedule_hash(schedule, identity['uids']['train'])

    baseline_root = roots['base_results'] / 'BASE_CE'
    baseline_manifest_path = baseline_root / 'manifest.json'
    baseline_checkpoint = baseline_root / 'checkpoint_final.pt'
    if not baseline_manifest_path.is_file() or not baseline_checkpoint.is_file():
        raise FileNotFoundError(f'paired IFNet CE baseline is missing under {baseline_root}')
    base_manifest = _read_json(baseline_manifest_path)
    split = identity['manifest']
    checks = {
        'status': 'complete', 'condition': 'BASE_CE',
        'final_epoch': int(student_cfg['epochs']), 'student_seed': seed,
        'target_subject_id': int(split['target_subject_id']),
        'feedback_subject_id': int(split['feedback_subject_id']),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'resolved_config_sha256': cfg['resolved_config_sha256'],
        'schedule_sha256': schedule_sha,
    }
    for key, expected in checks.items():
        if base_manifest.get(key) != expected:
            raise RuntimeError(f'BASE_CE baseline {key} mismatch: {base_manifest.get(key)!r} != {expected!r}')
    base_payload = torch.load(resolve_local_file(baseline_checkpoint), map_location='cpu',
                              weights_only=False)
    base_state = base_payload.get('model')
    if not base_state or _state_hash(base_state) != base_manifest.get('final_model_sha256'):
        raise RuntimeError('BASE_CE model state does not match its manifest')

    reuse_id = ('student_aware_diffusion_prior_loso_v1_seed666'
                if seed == 666 else 'mcc_diffkd_ifnet_initializations')
    init_roots = reference._run_roots(reuse_id, seed, fold)
    for key in ('run_result', 'run_seed_weights'):
        require_external_output(init_roots[key]).mkdir(parents=True, exist_ok=True)
    initial = reference._ensure_student_initial_state(init_roots, resolved, cfg, seed)
    warmup_path = roots['base_teacher'] / 'warmup_epoch10.pt'
    if not warmup_path.is_file():
        raise FileNotFoundError(f'paired CE warmup checkpoint is missing: {warmup_path}')
    warmup = torch.load(resolve_local_file(warmup_path), map_location='cpu', weights_only=False)
    if (int(warmup.get('epoch', -1)) != 10
            or warmup.get('train_uid_sha256') != identity['uid_sha256']['train']
            or warmup.get('resolved_config_sha256') != cfg['resolved_config_sha256']):
        raise RuntimeError('historical ten-epoch CE checkpoint does not match this fold')
    audit_path = init_roots['run_result'] / 'student_initialization_audit.json'
    if not audit_path.is_file():
        raise FileNotFoundError(f'paired IFNet initialization audit is missing: {audit_path}')
    audit = _read_json(audit_path)
    if (not audit.get('matched')
            or audit.get('initial_model_sha256') != initial['model_sha256']
            or audit.get('train_uid_sha256') != identity['uid_sha256']['train']
            or audit.get('schedule_sha256') != schedule_sha
            or audit.get('historical_warmup_sha256') != warmup.get('warmup_state_sha256')):
        raise RuntimeError('existing IFNet initialization audit does not match this fold')
    return {
        'student_adapter': adapter, 'student_cfg': student_cfg,
        'optimizer_defaults': resolved['optimizer_defaults'],
        'student_schedule': schedule, 'student_schedule_sha256': schedule_sha,
        'initial_student': initial, 'baseline_manifest': base_manifest,
        'baseline_checkpoint': str(baseline_checkpoint),
        'baseline_checkpoint_sha256': _sha(baseline_checkpoint),
        'initialization_audit': audit,
    }


def _make_projection(dimension: int, seed: int, roots: dict,
                     fingerprint: str) -> tuple[torch.Tensor, dict]:
    path = roots['run_cache'] / 'fixed_IFNet_channel_Q.npz'
    manifest_path = roots['run_cache'] / 'fixed_IFNet_channel_Q.json'
    if path.is_file() or manifest_path.is_file():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete fixed IFNet channel projection')
        manifest = _read_json(manifest_path)
        if (manifest.get('fingerprint') != fingerprint or manifest.get('seed') != seed
                or manifest.get('dimension') != dimension or manifest.get('file_sha256') != _sha(path)):
            raise RuntimeError('fixed IFNet channel projection identity/hash mismatch')
        with np.load(resolve_local_file(path), allow_pickle=False) as saved:
            q = np.asarray(saved['Q'], dtype=np.float32)
        return torch.as_tensor(q.reshape(3, dimension, 512)), manifest
    q_joint = common_core.fixed_student_projection(3 * dimension, seed)
    q = q_joint.reshape(3, dimension, 512)
    _atomic_npz(path, Q=q_joint, Q_float32=q.astype(np.float32))
    manifest = {
        'fingerprint': fingerprint, 'file_sha256': _sha(path), 'seed': seed,
        'dimension': dimension, 'joint_projection_shape': list(q_joint.shape),
        'channel_view_shape': list(q.shape),
        'construction': 'float64 Gaussian [512,3d] thin QR; positive R diagonal; reshape Q.T',
        'joint_row_orthogonality_error': float(np.max(np.abs(
            q_joint @ q_joint.T - np.eye(3 * dimension)))),
        'per_channel_views_are_distinct': True,
    }
    _atomic_json(manifest_path, manifest)
    return torch.as_tensor(q.astype(np.float32)), manifest


def _student_feedback(model, adapter, x: np.ndarray, y: np.ndarray,
                      device: torch.device) -> dict:
    model.eval()
    logits = []
    with torch.no_grad():
        for start in range(0, len(x), 64):
            xb = torch.as_tensor(x[start:start + 64], dtype=torch.float32, device=device)
            _, out = adapter.forward(model, xb)
            logits.append(out.float().cpu().numpy())
    values = np.concatenate(logits, axis=0)
    return _metric_values(y, values)


def _fit_student(condition: str, config: dict, identity: dict, roots: dict,
                  models: dict, data: dict, initial: dict,
                  teacher_targets: dict, projection: torch.Tensor,
                  prior: torch.nn.Module | None, fingerprint: str,
                  device: torch.device, smoke: bool = False) -> dict:
    final_path = roots['run_weights'] / f'student_{condition}.pt'
    manifest_path = roots['run_result'] / f'student_{condition}.json'
    feature_weight = float(config['student']['feature_weight'])
    teacher_name = CONDITION_TEACHER[condition]
    use_diffusion = CONDITION_DIFFUSION[condition]
    epochs = 2 if smoke else int(models['student_cfg']['epochs'])
    method_fingerprint = _json_hash({
        'run_fingerprint': fingerprint, 'condition': condition, 'teacher': teacher_name,
        'uses_diffusion': use_diffusion, 'initial_student_sha256': initial['model_sha256'],
        'teacher_target_sha256': _sha(roots['run_cache'] / 'teacher_channel_targets.npz'),
        'projection_sha256': _sha(roots['run_cache'] / 'fixed_IFNet_channel_Q.npz'),
        'prior_sha256': _sha(roots['run_weights'] / f'prior_{teacher_name}.pt') if use_diffusion else None,
        'epochs': epochs,
    })
    if final_path.is_file() or manifest_path.is_file():
        if not (final_path.is_file() and manifest_path.is_file()):
            raise RuntimeError(f'{condition} student checkpoint/manifest is incomplete')
        manifest = _read_json(manifest_path)
        if manifest.get('method_fingerprint') != method_fingerprint or manifest.get('checkpoint_sha256') != _sha(final_path):
            raise RuntimeError(f'{condition} student checkpoint identity/hash mismatch')
        return manifest
    model, optimizer, scheduler, _ = reference._new_paired_student(
        models, initial, int(config['seed']), device)
    train_x = torch.as_tensor(data['student_x']['train'], dtype=torch.float32, device=device)
    train_y = torch.as_tensor(identity['labels']['train'], dtype=torch.long, device=device)
    targets = torch.as_tensor(teacher_targets[f'{teacher_name}_train_features_whitened'],
                              dtype=torch.float32, device=device)
    q = projection.to(device)
    if prior is not None:
        prior.eval()
        for parameter in prior.parameters():
            parameter.requires_grad_(False)
            parameter.grad = None
    noise_seed = int(config['randomness']['student_diffusion_noise_seed_base']) \
        + 1009 * int(identity['manifest']['target_subject_id']) + int(config['seed'])
    noise_generator = torch.Generator(device=device).manual_seed(noise_seed) if use_diffusion else None
    alpha = sigma = None
    if use_diffusion:
        _, alpha, sigma = channel_core.channel_ddpm_schedule(
            int(config['diffusion']['steps']), float(config['diffusion']['beta_start']),
            float(config['diffusion']['beta_end']), device)
    schedule = models['student_schedule']
    history = []
    eval_rows = []
    first_gradients = None
    for epoch_index in range(epochs):
        model.train()
        ce_total = feature_total = total_loss = grad_total = 0.0
        seen = steps = 0
        for ids_np in schedule[epoch_index]:
            ids = torch.as_tensor(ids_np, dtype=torch.long, device=device)
            xb, yb = train_x[ids], train_y[ids]
            optimizer.zero_grad(set_to_none=True)
            native, logits = models['student_adapter'].forward(model, xb)
            native.retain_grad()
            projected = channel_core.student_channel_views(
                native, q, float(config['feature_space']['student_layer_norm_epsilon']))
            if use_diffusion:
                noise = torch.randn(projected.shape, dtype=torch.float32,
                                   device=device, generator=noise_generator)
                matched = channel_core.channel_ddim_denoise(
                    projected, noise, prior, alpha, sigma,
                    tuple(config['diffusion']['ddim_timesteps']))
            else:
                matched = projected
            feature_loss = F.mse_loss(
                matched, targets[ids].detach(), reduction='mean')
            ce = F.cross_entropy(logits, yb)
            loss = ce + feature_weight * feature_loss
            if not torch.isfinite(loss):
                raise FloatingPointError(f'{condition} non-finite student loss, epoch={epoch_index+1}')
            loss.backward()
            if first_gradients is None:
                native_grad = float(native.grad.norm().detach().cpu()) if native.grad is not None else 0.0
                first_param_grad = any(p.grad is not None and torch.isfinite(p.grad).all()
                                       and float(p.grad.norm().detach().cpu()) > 0
                                       for p in model.parameters())
                first_gradients = {'native_feature_gradient_norm': native_grad,
                                   'student_parameter_gradient_reached': bool(first_param_grad)}
                if native_grad <= 0 or not first_param_grad:
                    raise RuntimeError(f'{condition}: feature loss did not backpropagate into IFNet')
            grad_parts = [p.grad.detach().norm() for p in model.parameters() if p.grad is not None]
            grad_norm = torch.linalg.vector_norm(torch.stack(grad_parts))
            if not torch.isfinite(grad_norm):
                raise FloatingPointError(f'{condition} non-finite IFNet gradient')
            optimizer.step()
            count = len(ids)
            ce_total += float(ce.detach().cpu()) * count
            feature_total += float(feature_loss.detach().cpu()) * count
            total_loss += float(loss.detach().cpu()) * count
            grad_total += float(grad_norm.detach().cpu())
            seen += count
            steps += 1
        lr_used = float(optimizer.param_groups[0]['lr'])
        scheduler.step()
        row = {'condition': condition, 'teacher': teacher_name,
               'epoch': epoch_index + 1, 'train_ce': ce_total / seen,
               'train_feature_mse': feature_total / seen,
               'weighted_feature_mse': feature_weight * feature_total / seen,
               'train_total': total_loss / seen, 'mean_gradient_norm': grad_total / steps,
               'lr_used': lr_used, 'lr_next': float(scheduler.get_last_lr()[0]),
               'train_uid_sha256': identity['uid_sha256']['train'],
               'schedule_sha256': models['student_schedule_sha256'],
               **(first_gradients or {})}
        history.append(row)
        if (epoch_index + 1) in (1, 10, 50, 100) or (smoke and epoch_index + 1 == epochs):
            metrics = _student_feedback(model, models['student_adapter'],
                                        data['student_x']['feedback'],
                                        identity['labels']['feedback'], device)
            eval_rows.append({'condition': condition, 'epoch': epoch_index + 1,
                              **{f'feedback_{key}': value for key, value in metrics.items()}})
        _atomic_csv(roots['run_result'] / f'student_{condition}_history.csv', history)
        if eval_rows:
            _atomic_csv(roots['run_result'] / f'student_{condition}_feedback.csv', eval_rows)
    model.eval()
    state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    _atomic_torch(final_path, {'model': state, 'method_fingerprint': method_fingerprint,
                               'model_sha256': _state_hash(state), 'condition': condition})
    manifest = {
        'status': 'complete', 'condition': condition, 'teacher': teacher_name,
        'uses_diffusion': use_diffusion, 'method_fingerprint': method_fingerprint,
        'run_fingerprint': fingerprint, 'checkpoint': str(final_path),
        'checkpoint_sha256': _sha(final_path), 'model_sha256': _state_hash(state),
        'initial_model_sha256': initial['model_sha256'], 'final_epoch': epochs,
        'train_uid_sha256': identity['uid_sha256']['train'],
        'schedule_sha256': models['student_schedule_sha256'],
        'loss': {'ce_weight': 1.0, 'feature_weight': feature_weight,
                 'feature': 'mean squared error across batch, 3 channel views and d',
                 'logits_kd': False, 'start_epoch': 1},
        'fixed_Q_sha256': _sha(roots['run_cache'] / 'fixed_IFNet_channel_Q.npz'),
        'prior_checkpoint_sha256': (_sha(roots['run_weights'] / f'prior_{teacher_name}.pt')
                                    if use_diffusion else None),
        'history': str(roots['run_result'] / f'student_{condition}_history.csv'),
        'feedback_diagnostics': str(roots['run_result'] / f'student_{condition}_feedback.csv'),
    }
    _atomic_json(manifest_path, manifest)
    return manifest


def _fold_fingerprint(config_path: Path, config: dict, seed: int, fold: int,
                      identity: dict, roots: dict, cfg: dict,
                      teacher_cfg: dict, student_models: dict) -> str:
    values = {
        'experiment_id': config['experiment_id'], 'seed': seed, 'fold': fold,
        'config_sha256': _sha(config_path), 'implementation_sha256': _sha(__file__),
        'channel_core_sha256': _sha(ROOT / 'collab/cbramod_mcc_diffkd.py'),
        'mcc_core_sha256': _sha(ROOT / 'collab/mcc_diffkd.py'),
        'reference_runner_sha256': _sha(reference.__file__),
        'resolved_config_sha256': cfg['resolved_config_sha256'],
        'split_manifest_sha256': identity['sha256'],
        'source_manifest_sha256': identity['source_manifest_sha256'],
        'source_trials_sha256': identity['source_trial_manifest_sha256'],
        'cbramod_input_manifest_sha256': _sha(roots['input_root'] / 'cbramod/manifest.json'),
        'ifnet_input_manifest_sha256': _sha(roots['input_root'] / 'ifnet/manifest.json'),
        'prepared_ifnet_input_sha256': _sha(roots['split_dir'] / 'student_inputs.npz'),
        'pretrained_cbramod_sha256': _sha(teacher_cfg['pretrain']),
        'student_initial_sha256': student_models['initial_student']['model_sha256'],
        'student_baseline_manifest_sha256': _sha(roots['base_results'] / 'BASE_CE/manifest.json'),
        'student_schedule_sha256': student_models['student_schedule_sha256'],
        'teacher_config': teacher_cfg,
    }
    return _json_hash(values)


def _save_fold_resolved(roots: dict, config_path: Path, config: dict, seed: int,
                        fold: int, identity: dict, cfg: dict, teacher_cfg: dict,
                        student_models: dict, fingerprint: str) -> None:
    path = roots['run_result'] / 'resolved_run.json'
    if path.is_file():
        saved = _read_json(path)
        if saved.get('run_fingerprint') != fingerprint:
            raise RuntimeError(f'fold {fold + 1}: resolved run identity changed')
        return
    _atomic_json(path, {
        'status': 'running', 'experiment_id': config['experiment_id'],
        'run_fingerprint': fingerprint, 'config': str(config_path),
        'config_sha256': _sha(config_path), 'seed': seed, 'fold_zero_based': fold,
        'target_subject_id': fold + 1,
        'feedback_subject_id': int(identity['manifest']['feedback_subject_id']),
        'train_subject_ids': identity['manifest']['train_subject_ids'],
        'split_manifest': identity['path'], 'split_manifest_sha256': identity['sha256'],
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'test_uid_sha256': identity['uid_sha256']['test'],
        'historical_resolved_yaml': cfg['resolved_config_path'],
        'historical_resolved_yaml_sha256': cfg['resolved_yaml_sha256'],
        'resolved_config_sha256': cfg['resolved_config_sha256'],
        'cbramod_config': teacher_cfg,
        'ifnet_config': student_models['student_cfg'],
        'student_schedule_sha256': student_models['student_schedule_sha256'],
        'student_initialization_audit': student_models['initialization_audit'],
        'source_manifest_sha256': identity['source_manifest_sha256'],
        'source_trial_manifest_sha256': identity['source_trial_manifest_sha256'],
    })


def _train_fold(config_path: Path, config: dict, seed: int, run_id: str,
                fold: int, device: torch.device, smoke: bool = False,
                progress_path: Path | None = None) -> dict:
    _seed_all(int(config['randomness']['teacher_seed_base']) + 1009 * fold + seed, device)
    roots = _roots(run_id, seed, fold)
    historical, cfg = _resolved(config, seed)
    run_config = dict(config)
    run_config['seed'] = seed
    run_config['reference_run'] = {
        'protocol_id': config['reference_run']['protocol_id'], 'seed': seed,
        'resolved_config_sha256': cfg['resolved_config_sha256'],
    }
    identity = reference._load_fold_identity(
        run_config, fold, roots, cfg['resolved_config_sha256'], include_test_labels=False)
    data = _load_data(identity, roots, historical['student_config'], device)
    student_models = _load_student_reference(
        run_config, seed, fold, identity, roots, historical, cfg, data, device)
    teacher_cfg = project_config.load_model_config('cbramod', 'BNCI2014004', 'loso')
    teacher_cfg['pretrain'] = str(resolve_local_file(project_config.weight_path('cbramod')))
    if teacher_cfg.get('feature_head') != 'flatten' or teacher_cfg.get('skip_preprocess') is not True:
        raise RuntimeError('CBraMod resolved input/head recipe differs from this protocol')
    fingerprint = _fold_fingerprint(
        config_path, config, seed, fold, identity, roots, cfg, teacher_cfg, student_models)
    _save_fold_resolved(roots, config_path, config, seed, fold, identity,
                        cfg, teacher_cfg, student_models, fingerprint)
    complete_path = roots['run_result'] / 'fold_complete.json'
    if complete_path.is_file():
        complete = _read_json(complete_path)
        if complete.get('run_fingerprint') != fingerprint or complete.get('status') != 'trained':
            raise RuntimeError(f'fold {fold + 1}: completion record belongs to a different run')
        return {'status': 'trained', 'fold': fold, 'resumed_existing': True,
                'run_fingerprint': fingerprint}

    # T0 is a fresh seven-source fine-tune from the canonical CBraMod checkpoint.
    teacher_adapter, t0, t0_manifest = _fit_source_teacher(
        run_config, identity, roots, teacher_cfg, data, fingerprint, device, smoke=smoke)
    t0_train, t0_train_logits = _predict_teacher(
        t0, data['teacher_x']['train'], device, include_features=True)
    t0_feedback, t0_feedback_logits = _predict_teacher(
        t0, data['teacher_x']['feedback'], device, include_features=True)
    if t0_train.shape != (len(identity['uids']['train']), 3, 200):
        raise RuntimeError(f'CBraMod channel feature tap shape mismatch: {t0_train.shape}')
    pca_record = _pca_for_fold(t0_train, roots, fingerprint, config)
    pca = pca_record['pca']
    t0_diagnostics = []
    for part, logits in (('train', t0_train_logits), ('feedback', t0_feedback_logits)):
        t0_diagnostics.append({'teacher': 'T0', 'partition': part,
                               **_metric_values(identity['labels'][part], logits)})
    _atomic_csv(roots['run_result'] / 'teacher_T0_diagnostics.csv', t0_diagnostics)

    teachers = {'T0': t0}
    teacher_manifests = {'T0': t0_manifest}
    for name in ('T_CE', 'T_MCC'):
        model, manifest = _adapt_teacher(
            name, run_config, identity, roots, teacher_cfg, teacher_adapter,
            t0, data, fingerprint, device, smoke=smoke)
        teachers[name] = model
        teacher_manifests[name] = manifest

    targets, target_manifest = _teacher_targets(
        identity, roots, data, pca_record, teachers, teacher_adapter, fingerprint, device)
    teacher_diag = []
    for name in ('T_CE', 'T_MCC'):
        for part in ('train', 'feedback'):
            logits = targets[f'{name}_{part}_logits']
            teacher_diag.append({'teacher': name, 'partition': part,
                                 **_metric_values(identity['labels'][part], logits)})
    _atomic_csv(roots['run_result'] / 'teacher_adaptation_diagnostics.csv', teacher_diag)

    dim = int(pca['dimension'])
    q_seed = int(config['feature_space']['student_projection_seed_base']) \
        + 100 * (seed - 666) + fold
    projection, projection_manifest = _make_projection(dim, q_seed, roots, fingerprint)
    prior_seed = int(config['randomness']['diffusion_seed_base']) + 31 * fold + seed
    prior_initial, prior_initial_hash = _channel_prior_init(dim, prior_seed, config)
    initial_path = roots['run_weights'] / 'prior_common_initial.pt'
    _atomic_torch(initial_path, {'model': prior_initial, 'model_sha256': prior_initial_hash,
                                 'fingerprint': fingerprint, 'seed': prior_seed})
    priors, prior_manifests = {}, {}
    for name in ('T_CE', 'T_MCC'):
        prior, manifest = _fit_channel_prior(
            name, run_config, identity, roots, targets, fingerprint,
            prior_initial, prior_initial_hash, device, smoke=smoke)
        priors[name] = prior
        prior_manifests[name] = manifest

    student_manifests = {}
    for condition in METHODS:
        teacher_name = CONDITION_TEACHER[condition]
        manifest = _fit_student(
            condition, run_config, identity, roots, student_models, data,
            student_models['initial_student'], targets, projection,
            priors[teacher_name] if CONDITION_DIFFUSION[condition] else None,
            fingerprint, device, smoke=smoke)
        student_manifests[condition] = manifest
    if smoke:
        smoke_record = {
            'status': 'passed', 'run_fingerprint': fingerprint,
            'seed': seed, 'fold': fold, 'target_subject_id': fold + 1,
            'teacher_T0_epochs': t0_manifest['epochs'],
            'teacher_adaptation_epochs': {k: v['epochs'] for k, v in teacher_manifests.items()
                                          if k != 'T0'},
            'prior_updates': {k: v['final_update'] for k, v in prior_manifests.items()},
            'student_epochs': {k: v['final_epoch'] for k, v in student_manifests.items()},
            'pca_dimension': dim, 'Q_sha256': projection_manifest['file_sha256'],
            'teacher_targets_sha256': target_manifest['file_sha256'],
            'target_evaluated': False,
        }
        _atomic_json(roots['run_result'] / 'smoke_result.json', smoke_record)
        return {'status': 'smoke_passed', 'fold': fold, 'run_fingerprint': fingerprint}

    completion = {
        'status': 'trained', 'experiment_id': config['experiment_id'],
        'run_fingerprint': fingerprint, 'seed': seed, 'fold_zero_based': fold,
        'target_subject_id': fold + 1,
        'train_uid_sha256': identity['uid_sha256']['train'],
        'teacher_checkpoints': {k: v['checkpoint_sha256'] for k, v in teacher_manifests.items()},
        'prior_checkpoints': {k: v['checkpoint_sha256'] for k, v in prior_manifests.items()},
        'student_checkpoints': {k: v['checkpoint_sha256'] for k, v in student_manifests.items()},
        'pca_sha256': pca_record['manifest']['file_sha256'],
        'Q_sha256': projection_manifest['file_sha256'],
        'teacher_targets_sha256': target_manifest['file_sha256'],
        'initial_IFNet_sha256': student_models['initial_student']['model_sha256'],
        'completed_at_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
    }
    _atomic_json(complete_path, completion)
    _atomic_json(roots['run_result'] / 'resolved_run.json', {
        **_read_json(roots['run_result'] / 'resolved_run.json'),
        'status': 'trained', 'fold_complete_sha256': _sha(complete_path)})
    _write_progress(progress_path, {'status': 'running', 'seed': seed,
                                    'records': [{'fold': fold + 1, 'status': 'trained'}]})
    return {'status': 'trained', 'fold': fold, 'run_fingerprint': fingerprint}


def _predict_student(model, adapter, x: np.ndarray, device: torch.device,
                     batch_size: int = 64) -> np.ndarray:
    model.eval()
    output = []
    with torch.no_grad():
        for start in range(0, len(x), batch_size):
            xb = torch.as_tensor(x[start:start + batch_size], dtype=torch.float32, device=device)
            _, logits = adapter.forward(model, xb)
            output.append(logits.float().cpu().numpy())
    return np.concatenate(output, axis=0)


def _predict_student_preserving_state(model, adapter, x: np.ndarray,
                                      device: torch.device) -> tuple[np.ndarray, str]:
    before = reference._cpu_tree(model.state_dict())
    before_hash = _state_hash(before)
    logits = _predict_student(model, adapter, x, device)
    after_hash = _state_hash(model.state_dict())
    if after_hash != before_hash:
        # IFNet's constrained Linear layer may renormalize its weights during
        # eval forward. Restore the checkpoint so an audit pass cannot mutate it.
        model.load_state_dict(before, strict=True)
    if _state_hash(model.state_dict()) != before_hash:
        raise RuntimeError('IFNet evaluation failed to restore its checkpoint state')
    return logits, before_hash


def _load_teacher_checkpoint(adapter, checkpoint: Path, device: torch.device,
                             initialization_seed: int) -> torch.nn.Module:
    _, model = _build_teacher(adapter.cfg, device, initialization_seed)
    payload = torch.load(resolve_local_file(checkpoint), map_location='cpu', weights_only=False)
    model.load_state_dict(payload['model'], strict=True)
    model.eval()
    return model


def _evaluate_fold(config: dict, seed: int, run_id: str, fold: int,
                   device: torch.device) -> tuple[list[dict], list[dict]]:
    roots = _roots(run_id, seed, fold)
    eval_path = roots['run_result'] / 'fold_evaluation.json'
    subject_path = roots['run_result'] / 'fold_evaluation_metrics.csv'
    prediction_path = roots['run_result'] / 'fold_predictions.csv'
    if eval_path.is_file() and subject_path.is_file() and prediction_path.is_file():
        record = _read_json(eval_path)
        if record.get('status') != 'complete':
            raise RuntimeError(f'fold {fold + 1}: evaluation record is incomplete')
        metrics = pd.read_csv(resolve_local_file(subject_path)).to_dict('records')
        predictions = pd.read_csv(resolve_local_file(prediction_path)).to_dict('records')
        shutil.rmtree(roots['run_weights'], ignore_errors=True)
        return metrics, predictions
    complete_path = roots['run_result'] / 'fold_complete.json'
    if not complete_path.is_file() or _read_json(complete_path).get('status') != 'trained':
        raise RuntimeError(f'fold {fold + 1}: target evaluation requested before training completed')
    complete = _read_json(complete_path)
    if complete.get('seed') != seed or complete.get('fold_zero_based') != fold:
        raise RuntimeError(f'fold {fold + 1}: training completion identity mismatch')
    historical, cfg = _resolved(config, seed)
    run_config = dict(config)
    run_config['seed'] = seed
    run_config['reference_run'] = {'protocol_id': config['reference_run']['protocol_id'],
                                   'seed': seed,
                                   'resolved_config_sha256': cfg['resolved_config_sha256']}
    identity = reference._load_fold_identity(
        run_config, fold, roots, cfg['resolved_config_sha256'], include_test_labels=True)
    data = _load_data(identity, roots, historical['student_config'], device)
    target_rows = data['teacher_rows']['test']
    teacher_x = np.asarray(data['teacher_input_mmap'][target_rows], dtype=np.float32)
    student_raw = np.asarray(data['student_input_mmap'][data['student_rows']['test']], dtype=np.float32)
    student_x = reference._preprocess_student(data['student_adapter'], student_raw)
    labels = np.asarray(identity['labels']['test'], dtype=np.int64)
    uids = np.asarray(identity['uids']['test'], dtype=np.int64)
    if len(labels) != len(teacher_x) or len(labels) != len(student_x):
        raise RuntimeError('target data/label length mismatch')

    teacher_cfg = project_config.load_model_config('cbramod', 'BNCI2014004', 'loso')
    teacher_cfg['pretrain'] = str(resolve_local_file(project_config.weight_path('cbramod')))
    teacher_adapter = get_adapter('cbramod', device=device, **teacher_cfg)
    student_adapter = get_adapter('ifnet', device=device, **historical['student_config'])
    model_logits = {}
    for name in TEACHERS:
        manifest = _read_json(roots['run_result'] / f'teacher_{name}.json')
        model = _load_teacher_checkpoint(teacher_adapter,
            roots['run_weights'] / f'teacher_{name}.pt', device,
            int(manifest['initialization_seed']) if name == 'T0'
            else int(config['randomness']['teacher_seed_base']) + 1009 * (fold + 1) + seed)
        _, logits = _predict_teacher(model, teacher_x, device)
        model_logits[name] = logits

    base_root = roots['base_results'] / 'BASE_CE'
    baseline_payload = torch.load(resolve_local_file(base_root / 'checkpoint_final.pt'),
                                  map_location='cpu', weights_only=False)
    baseline_model = student_adapter.build(2)
    baseline_model.load_state_dict(baseline_payload['model'], strict=True)
    baseline_hash = _state_hash(baseline_model.state_dict())
    if baseline_hash != _read_json(base_root / 'manifest.json').get('final_model_sha256'):
        raise RuntimeError('IFNet_CE state changed since baseline manifest')
    model_logits['IFNet_CE'], restored_baseline_hash = _predict_student_preserving_state(
        baseline_model, student_adapter, student_x, device)
    if restored_baseline_hash != baseline_hash:
        raise RuntimeError('IFNet_CE evaluation changed checkpoint state')
    for condition in METHODS:
        path = roots['run_weights'] / f'student_{condition}.pt'
        payload = torch.load(resolve_local_file(path), map_location='cpu', weights_only=False)
        model = student_adapter.build(2)
        model.load_state_dict(payload['model'], strict=True)
        if _state_hash(model.state_dict()) != payload.get('model_sha256'):
            raise RuntimeError(f'{condition}: final IFNet checkpoint hash mismatch')
        model_logits[condition], _ = _predict_student_preserving_state(
            model, student_adapter, student_x, device)

    target_subject = int(identity['manifest']['target_subject_id'])
    metric_rows, prediction_rows = [], []
    for model_name, logits in model_logits.items():
        metric_rows.append({'seed': seed, 'target_subject': target_subject,
                            'model': model_name,
                            'method': ('none' if model_name == 'IFNet_CE' else
                                       'teacher_reference' if model_name.startswith('T') else model_name),
                            'n_trials': len(labels), **_metric_values(labels, logits)})
    for row_idx, (uid, label) in enumerate(zip(uids, labels)):
        row = {'seed': seed, 'target_subject': target_subject,
               'subject_zero_based': int(uid[0]), 'legacy_raw_row': int(uid[1]),
               'trial_uid': f'{int(uid[0])}:{int(uid[1])}', 'label': int(label)}
        for model_name, logits in model_logits.items():
            row[f'{model_name}_logit_0'] = float(logits[row_idx, 0])
            row[f'{model_name}_logit_1'] = float(logits[row_idx, 1])
            row[f'{model_name}_prediction'] = int(np.argmax(logits[row_idx]))
        prediction_rows.append(row)
    _atomic_csv(subject_path, metric_rows)
    _atomic_csv(prediction_path, prediction_rows)
    _atomic_json(eval_path, {
        'status': 'complete', 'seed': seed, 'target_subject': target_subject,
        'run_fingerprint': complete['run_fingerprint'],
        'test_uid_sha256': identity['uid_sha256']['test'],
        'test_rows': len(labels), 'metrics_sha256': _sha(subject_path),
        'predictions_sha256': _sha(prediction_path),
        'target_labels_used_for_evaluation_only': True,
    })
    shutil.rmtree(roots['run_weights'], ignore_errors=True)
    return metric_rows, prediction_rows


def _report_write(path: Path, rows: list[dict]) -> None:
    target = require_report_output(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + '.tmp')
    fields = list(rows[0]) if rows else []
    with temporary.open('w', encoding='utf-8', newline='') as stream:
        if fields:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
    os.replace(temporary, target)


def _report_json(path: Path, value: dict) -> None:
    target = require_report_output(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    indent=2) + '\n', encoding='utf-8')
    os.replace(temporary, target)


def _evaluate_seed(config: dict, seed: int, run_id: str,
                   device: torch.device, progress_path: Path | None = None) -> dict:
    all_metrics, all_predictions = [], []
    for fold in range(9):
        metrics, predictions = _evaluate_fold(config, seed, run_id, fold, device)
        all_metrics.extend(metrics)
        all_predictions.extend(predictions)
        _write_progress(progress_path, {'status': 'evaluating', 'seed': seed,
                                       'records': [{'fold': i + 1, 'status': 'evaluated'}
                                                   for i in range(fold + 1)]})
    expected = 9 * (3 + 1 + len(METHODS))
    if len(all_metrics) != expected:
        raise RuntimeError(f'seed {seed}: expected {expected} model/fold metric rows, got {len(all_metrics)}')
    report_dir = REPORTS_ROOT / 'distill' / 'cbramod_mcc_diffkd_loso_v1' / f'seed_{seed}'
    _report_write(report_dir / 'subject_metrics.csv', all_metrics)
    _report_write(report_dir / 'trial_predictions.csv', all_predictions)
    means = []
    for model_name in dict.fromkeys(row['model'] for row in all_metrics):
        rows = [row for row in all_metrics if row['model'] == model_name]
        if len(rows) != 9 or len({row['target_subject'] for row in rows}) != 9:
            raise RuntimeError(f'{model_name}: incomplete or duplicate LOSO subjects in seed {seed}')
        means.append({'seed': seed, 'model': model_name,
                      **{metric: float(np.mean([row[metric] for row in rows]))
                         for metric in ('accuracy', 'balanced_accuracy', 'kappa', 'cross_entropy')},
                      'n_subjects': 9})
    _report_write(report_dir / 'seed_means.csv', means)
    _report_json(report_dir / 'completion.json', {
        'status': 'complete', 'seed': seed, 'n_subjects': 9,
        'n_metrics': len(all_metrics), 'n_predictions': len(all_predictions),
        'aggregation': 'equal-weight mean over nine target subjects within this seed',
        'training_run_ids': [run_id],
        'subject_metrics_sha256': _sha(report_dir / 'subject_metrics.csv'),
        'trial_predictions_sha256': _sha(report_dir / 'trial_predictions.csv'),
    })
    _write_progress(progress_path, {'status': 'completed', 'seed': seed,
                                   'records': [{'fold': i + 1, 'status': 'evaluated'}
                                               for i in range(9)]})
    return {'metrics': all_metrics, 'predictions': all_predictions, 'seed_means': means}


def run_seed(config_path: Path, config: dict, seed: int, run_id: str,
             gpu: int = 0, folds: list[int] | None = None,
             smoke: bool = False, progress_path: Path | None = None,
             stop_file: Path | None = None) -> int:
    _validate_config(config, seed)
    device = _device(gpu)
    selected = [0] if smoke else list(folds if folds is not None else range(9))
    if not selected or len(set(selected)) != len(selected) or any(f < 0 or f > 8 for f in selected):
        raise ValueError('folds must be a unique non-empty subset of zero-based 0..8')
    records = []
    for fold in selected:
        record = _train_fold(config_path, config, seed, run_id, fold, device,
                             smoke=smoke, progress_path=progress_path)
        records.append({'fold': fold + 1, 'status': record['status']})
        _write_progress(progress_path, {'status': 'running', 'seed': seed, 'records': records})
        if smoke:
            continue
        if stop_file is not None and Path(stop_file).exists():
            _write_progress(progress_path, {'status': 'paused', 'seed': seed, 'records': records})
            return 75
    if smoke:
        _write_progress(progress_path, {'status': 'completed', 'seed': seed,
                                       'records': records})
        return 0
    if len(selected) != 9:
        raise ValueError('formal seed run must train all nine folds before target evaluation')
    if stop_file is not None and Path(stop_file).exists():
        _write_progress(progress_path, {'status': 'paused', 'seed': seed,
                                       'records': records})
        return 75
    _evaluate_seed(config, seed, run_id, device, progress_path)
    _atomic_json(RESULTS_ROOT / 'distill' / run_id / f'seed_{seed}' / 'seed_complete.json', {
        'status': 'complete', 'seed': seed, 'run_id': run_id,
        'completed_at_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())})
    return 0


def _aggregate_three_seed(config_path: Path, config: dict) -> dict:
    report_root = REPORTS_ROOT / 'distill' / 'cbramod_mcc_diffkd_loso_v1'
    subject_rows, seed_rows, prediction_rows = [], [], []
    for seed in (666, 667, 668):
        seed_dir = report_root / f'seed_{seed}'
        completion_path = seed_dir / 'completion.json'
        metrics_path = seed_dir / 'subject_metrics.csv'
        predictions_path = seed_dir / 'trial_predictions.csv'
        if not all(path.is_file() for path in (completion_path, metrics_path, predictions_path)):
            raise FileNotFoundError(f'seed {seed}: completed LOSO evaluation report is incomplete')
        completion = _read_json(completion_path)
        if completion.get('status') != 'complete' or completion.get('n_subjects') != 9:
            raise RuntimeError(f'seed {seed}: report completion record is invalid')
        current = pd.read_csv(resolve_local_file(metrics_path)).to_dict('records')
        if len(current) != 72 or {int(row['target_subject']) for row in current} != set(range(1, 10)):
            raise RuntimeError(f'seed {seed}: expected 72 model/fold rows covering subjects 1..9')
        subject_rows.extend(current)
        prediction_rows.extend(pd.read_csv(resolve_local_file(predictions_path)).to_dict('records'))
        for model_name in dict.fromkeys(row['model'] for row in current):
            rows = [row for row in current if row['model'] == model_name]
            if len(rows) != 9 or len({int(row['target_subject']) for row in rows}) != 9:
                raise RuntimeError(f'seed {seed}/{model_name}: duplicate or missing fold')
            seed_rows.append({'seed': seed, 'model': model_name,
                **{metric: float(np.mean([float(row[metric]) for row in rows]))
                   for metric in ('accuracy', 'balanced_accuracy', 'kappa', 'cross_entropy')},
                'n_subjects': 9})
    summary = []
    model_names = list(dict.fromkeys(row['model'] for row in seed_rows))
    for model_name in model_names:
        rows = [row for row in seed_rows if row['model'] == model_name]
        if {int(row['seed']) for row in rows} != {666, 667, 668}:
            raise RuntimeError(f'{model_name}: expected all three seeds')
        record = {'model': model_name, 'n_seed': 3, 'seeds': '666,667,668', 'n_subjects_per_seed': 9}
        for metric in ('accuracy', 'balanced_accuracy', 'kappa', 'cross_entropy'):
            values = np.asarray([float(row[metric]) for row in rows], dtype=np.float64)
            record[f'{metric}_mean'] = float(values.mean())
            record[f'{metric}_sd_across_seed_means'] = float(values.std(ddof=1))
        summary.append(record)

    comparisons = []
    pairs = (('primary_MCC_DIFFKD_minus_MCC_FEATURE', 'MCC_DIFFKD', 'MCC_FEATURE'),
             ('MCC_effect_with_diffusion', 'MCC_DIFFKD', 'CE_DIFFKD'),
             ('MCC_effect_direct_feature', 'MCC_FEATURE', 'CE_FEATURE'),
             ('DiffKD_effect_under_CE_teacher', 'CE_DIFFKD', 'CE_FEATURE'),
             ('full_combo_minus_CE_FEATURE', 'MCC_DIFFKD', 'CE_FEATURE'))
    for label, left, right in pairs:
        per_seed = []
        paired_differences = []
        for seed in (666, 667, 668):
            for metric in ('accuracy', 'balanced_accuracy'):
                lookup = {(int(row['target_subject']), row['model']): float(row[metric])
                          for row in subject_rows if int(row['seed']) == seed}
                fold_deltas = [lookup[(subject, left)] - lookup[(subject, right)]
                               for subject in range(1, 10)]
                paired_differences.extend({'seed': seed, 'subject': subject,
                    'metric': metric, 'left': left, 'right': right,
                    'difference': lookup[(subject, left)] - lookup[(subject, right)]}
                    for subject in range(1, 10))
                per_seed.append({'seed': seed, 'metric': metric,
                                 'mean_subject_difference': float(np.mean(fold_deltas))})
        for metric in ('accuracy', 'balanced_accuracy'):
            values = np.asarray([r['mean_subject_difference'] for r in per_seed
                                 if r['metric'] == metric])
            comparisons.append({'comparison': label, 'left': left, 'right': right,
                'metric': metric, 'mean_difference': float(values.mean()),
                'sd_across_seed_means': float(values.std(ddof=1)), 'n_seed': 3,
                'seeds': '666,667,668', 'paired_subjects_per_seed': 9})
    report_root.mkdir(parents=True, exist_ok=True)
    _report_write(report_root / 'summary_3seed.csv', summary)
    _report_write(report_root / 'seed_means_3seed.csv', seed_rows)
    _report_write(report_root / 'paired_differences_3seed.csv', comparisons)
    _report_write(report_root / 'subject_metrics_3seed.csv', subject_rows)
    _report_write(report_root / 'trial_predictions_3seed.csv', prediction_rows)
    _write_workbook(report_root / 'cbramod_mcc_diffkd_loso_3seed.xlsx', summary)
    result = {
        'status': 'complete', 'experiment_id': config['experiment_id'],
        'seeds': [666, 667, 668], 'subjects_per_seed': 9,
        'aggregation': 'equal-weight subject mean within each seed, then equal-weight mean over seeds',
        'student_conditions': list(METHODS),
        'primary_comparison': 'MCC_DIFFKD - MCC_FEATURE',
        'n_seed_subject_pairs': 27,
        'summary_csv_sha256': _sha(report_root / 'summary_3seed.csv'),
        'paired_differences_csv_sha256': _sha(report_root / 'paired_differences_3seed.csv'),
        'workbook': str(report_root / 'cbramod_mcc_diffkd_loso_3seed.xlsx'),
        'completed_at_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
    }
    _report_json(report_root / 'completion_3seed.json', result)
    return result


def _write_workbook(path: Path, summary: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    lookup = {row['model']: row for row in summary}
    teacher_map = {'CE_FEATURE': 'T_CE', 'CE_DIFFKD': 'T_CE',
                   'MCC_FEATURE': 'T_MCC', 'MCC_DIFFKD': 'T_MCC'}
    workbook = Workbook()
    workbook.remove(workbook.active)
    specs = (('Accuracy (%)', 'accuracy_mean'),
             ('Balanced Accuracy (%)', 'balanced_accuracy_mean'))
    for title, metric_key in specs:
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        headers = ['student', 'method', 'BNCI2014004']
        sheet.append(headers)
        teacher0 = lookup['T0'][metric_key] * 100
        for model, method, value, teacher_model in (
            ('CBraMod', 'T0', teacher0, None),
            ('CBraMod', 'CE adapted', lookup['T_CE'][metric_key] * 100, None),
            ('CBraMod', 'MCC adapted', lookup['T_MCC'][metric_key] * 100, None),
            ('IFNet', 'none', lookup['IFNet_CE'][metric_key] * 100, None),
            ('IFNet', 'CE_FEATURE', lookup['CE_FEATURE'][metric_key] * 100, 'T_CE'),
            ('IFNet', 'MCC_FEATURE', lookup['MCC_FEATURE'][metric_key] * 100, 'T_MCC'),
            ('IFNet', 'CE_DIFFKD', lookup['CE_DIFFKD'][metric_key] * 100, 'T_CE'),
            ('IFNet', 'MCC_DIFFKD', lookup['MCC_DIFFKD'][metric_key] * 100, 'T_MCC'),
        ):
            sheet.append([model, method, round(float(value), 2)])
            cell = sheet.cell(sheet.max_row, 3)
            if model == 'CBraMod':
                cell.fill = PatternFill('solid', fgColor='FAC090')
            elif method == 'none':
                cell.fill = PatternFill('solid', fgColor='93CDDD')
            elif teacher_model:
                score_teacher = float(lookup[teacher_model][metric_key]) * 100
                score_student = float(lookup['IFNet_CE'][metric_key]) * 100
                score = float(value)
                if score > score_teacher and score > score_student:
                    color = 'FFC7CE'
                elif score > score_teacher:
                    color = 'C6EFCE'
                elif score > score_student:
                    color = 'FFF2CC'
                else:
                    color = None
                if color:
                    cell.fill = PatternFill('solid', fgColor=color)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            row[2].number_format = '0.00'
        sheet.column_dimensions['A'].width = 20
        sheet.column_dimensions['B'].width = 24
        sheet.column_dimensions['C'].width = 20
        legend_row = sheet.max_row + 2
        legend = [('大于小模型', 'FFF2CC'), ('大于大模型', 'C6EFCE'), ('大于二者', 'FFC7CE')]
        for offset, (label, color) in enumerate(legend):
            c = sheet.cell(legend_row, 1 + offset, label)
            c.fill = PatternFill('solid', fgColor=color)
            c.alignment = Alignment(horizontal='center')
        notes = [
            '数值为百分比；seed 内对 9 个 LOSO 被试等权平均，再对 seed 666/667/668 等权平均。',
            '括号中的 seed SD 在 summary_3seed.csv；不是跨被试 SD。方法与对应适配教师及 IFNet CE baseline 比较。',
            'CE_FEATURE/MCC_FEATURE 为三通道配对特征 MSE；*_DIFFKD 使用同通道条件去噪后的配对特征 MSE。所有学生损失均为 CE + 0.5×feature loss，无 logits KD。',
        ]
        note_row = legend_row + 2
        for index, note in enumerate(notes):
            sheet.merge_cells(start_row=note_row + index, start_column=1,
                              end_row=note_row + index, end_column=3)
            cell = sheet.cell(note_row + index, 1, note)
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            sheet.row_dimensions[note_row + index].height = 30
    out = require_report_output(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)


def enqueue_run(config_path: Path, *, queue_root: Path | None = None,
                smoke: bool = False, smoke_id: str | None = None) -> list[str]:
    from experiments.distill.queue import DEFAULT_ROOT, enqueue_snapshot
    config_path = Path(config_path).resolve()
    config = _read_config(config_path)
    selected_seeds = [666] if smoke else [666, 667, 668]
    for seed in selected_seeds:
        _validate_config(config, seed)
    root = Path(queue_root or DEFAULT_ROOT).resolve()
    jobs = []
    for seed in selected_seeds:
        run_id = (f'cbramod_mcc_diffkd_loso_v1_smoke_{smoke_id or int(time.time())}_seed{seed}' if smoke
                  else f'cbramod_mcc_diffkd_loso_v1_seed{seed}')
        snapshot = {
            'kind': 'mcc_diffkd_cbramod_loso_v1',
            'stage': 'smoke' if smoke else 'train_seed',
            'config': str(config_path), 'config_sha256': _sha(config_path),
            'seed': seed, 'run_id': run_id,
            'folds': [0] if smoke else list(range(9)),
            'args': {'dataset': 'BNCI2014004', 'teacher': 'cbramod',
                     'student': 'ifnet', 'protocol': 'loso_7_1_1',
                     'keys': [0] if smoke else list(range(9)), 'seeds': [seed],
                     'methods': list(METHODS), 'source_variant': 'wide_fre_band',
                     'selected_session': 'session_3', 'epochs': 100},
        }
        jobs.append(enqueue_snapshot(
            snapshot, root=root,
            label=(f'CBraMod MCC-DiffKD smoke seed={seed} target=01' if smoke else
                   f'CBraMod MCC-DiffKD seed={seed}, 9-fold LOSO')))
    if not smoke:
        dispatch_file = require_report_output(
            REPORTS_ROOT / 'distill' / 'cbramod_mcc_diffkd_loso_v1'
            / 'dispatch' / 'dispatch.json')
        dispatch_dir = dispatch_file.parent
        dispatch_dir.mkdir(parents=True, exist_ok=True)
        watcher_log = dispatch_dir / 'watcher.log'
        command = [sys.executable, '-B', '-m', 'experiments.distill.run_cbramod_mcc_diffkd_loso',
                   '--watch-jobs', '--queue-root', str(root), '--job-ids', *jobs,
                   '--config', str(config_path)]
        environment = os.environ.copy()
        for name in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS',
                     'NUMEXPR_NUM_THREADS', 'BLIS_NUM_THREADS',
                     'VECLIB_MAXIMUM_THREADS', 'BIGSMALLCOLLAB_CPU_THREADS',
                     'TORCH_NUM_THREADS'):
            environment[name] = '1'
        with watcher_log.open('a', encoding='utf-8') as stream:
            process = subprocess.Popen(command, cwd=ROOT, stdin=subprocess.DEVNULL,
                stdout=stream, stderr=subprocess.STDOUT, start_new_session=True,
                env=environment, close_fds=True)
        _report_json(dispatch_dir / 'dispatch.json', {
            'status': 'submitted', 'experiment_id': config['experiment_id'],
            'config': str(config_path), 'config_sha256': _sha(config_path),
            'seeds': selected_seeds, 'training_job_ids': jobs,
            'queue_root': str(root), 'watcher_pid': process.pid,
            'target_evaluation': 'each seed evaluates only after all 9 folds train'})
    return jobs


def _watch_jobs(config_path: Path, queue_root: Path, job_ids: list[str]) -> int:
    from experiments.distill.queue import QueueStore, validated_root
    store = QueueStore(validated_root(queue_root))
    while True:
        states = [store.load(job_id) for job_id in job_ids]
        failed = [state for state in states if state['status'] in ('failed', 'cancelled')]
        if failed:
            report = REPORTS_ROOT / 'distill' / 'cbramod_mcc_diffkd_loso_v1' / 'dispatch' / 'watcher_state.json'
            _report_json(report, {'status': 'failed', 'jobs': [
                {'id': s['id'], 'status': s['status'], 'message': s.get('message')}
                for s in failed]})
            return 1
        if all(state['status'] == 'completed' for state in states):
            result = _aggregate_three_seed(config_path, _read_config(config_path))
            _report_json(REPORTS_ROOT / 'distill' / 'cbramod_mcc_diffkd_loso_v1'
                         / 'dispatch' / 'watcher_state.json', {
                'status': 'complete', 'jobs': job_ids,
                'aggregate_completion': result})
            return 0
        time.sleep(30)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, default=DEFAULT_CONFIG)
    parser.add_argument('--seed', type=int, default=666)
    parser.add_argument('--run-id', default=None)
    parser.add_argument('--gpu', type=int, default=0)
    parser.add_argument('--folds', nargs='+', type=int, default=None)
    parser.add_argument('--smoke', action='store_true')
    parser.add_argument('--enqueue-smoke', action='store_true')
    parser.add_argument('--smoke-id', default=None)
    parser.add_argument('--enqueue', action='store_true')
    parser.add_argument('--queue-root', type=Path, default=None)
    parser.add_argument('--progress', type=Path, default=None)
    parser.add_argument('--stop-file', type=Path, default=None)
    parser.add_argument('--aggregate', action='store_true')
    parser.add_argument('--watch-jobs', action='store_true')
    parser.add_argument('--job-ids', nargs='*', default=[])
    args = parser.parse_args(argv)
    try:
        config_path = args.config.resolve()
        config = _read_config(config_path)
        if args.enqueue_smoke:
            jobs = enqueue_run(config_path, queue_root=args.queue_root,
                               smoke=True, smoke_id=args.smoke_id)
            print(json.dumps({'smoke_job_ids': jobs}, indent=2), flush=True)
            return 0
        if args.enqueue:
            jobs = enqueue_run(config_path, queue_root=args.queue_root, smoke=False)
            print(json.dumps({'training_job_ids': jobs}, indent=2), flush=True)
            return 0
        if args.aggregate:
            print(json.dumps(_aggregate_three_seed(config_path, config), indent=2), flush=True)
            return 0
        if args.watch_jobs:
            if not args.job_ids:
                raise ValueError('--watch-jobs requires --job-ids')
            return _watch_jobs(config_path, args.queue_root or Path(
                '/home/lixinli/BigSmallCollab/results/loso/distill/queue'), args.job_ids)
        _validate_config(config, args.seed)
        run_id = args.run_id or f'cbramod_mcc_diffkd_loso_v1_seed{args.seed}'
        return run_seed(config_path, config, args.seed, run_id, args.gpu,
                        args.folds, args.smoke, args.progress, args.stop_file)
    except Exception as exc:
        traceback.print_exc()
        print(f'[cbramod-mcc-diffkd] error: {type(exc).__name__}: {exc}',
              file=sys.stderr, flush=True)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())

