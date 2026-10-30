"""Run the preregistered MCC-teacher / DiffKD IFNet LOSO pilot."""
from __future__ import annotations

# Respect the shared-host CPU budget before importing NumPy, Torch, or sklearn.
import os
for _name in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS',
              'NUMEXPR_NUM_THREADS', 'BLIS_NUM_THREADS', 'VECLIB_MAXIMUM_THREADS',
              'BIGSMALLCOLLAB_CPU_THREADS', 'TORCH_NUM_THREADS'):
    os.environ[_name] = '1'

import argparse
import copy
import csv
import hashlib
import json
import math
from pathlib import Path
import random
import shutil
import subprocess
import sys
import time
import traceback

import numpy as np
import torch
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from torch.nn import functional as F

from collab import mcc_diffkd as core
from collab.student_feature_prior import apply_pca_whitener, fit_pca_whitener
from experiments.distill import run_student_aware_diffusion_loso as reference
from experiments.storage import (
    DATA_CACHE_ROOT, REPORTS_ROOT, RESULTS_ROOT, WEIGHTS_ROOT,
    require_external_output, require_report_output, resolve_local_file,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CONFIG = PROJECT_ROOT / 'configs/experiments/loso/mcc_diffkd_loso_v1.yaml'
METHODS = ('CE_FEATURE', 'MCC_FEATURE', 'CE_DIFFKD', 'MCC_DIFFKD')
TEACHER_NAMES = ('T_CE', 'T_MCC')
NOISE_BINS = ((1, 250), (251, 500), (501, 750), (751, 1000))
CONDITION_TEACHER = {
    'CE_FEATURE': 'T_CE', 'CE_DIFFKD': 'T_CE',
    'MCC_FEATURE': 'T_MCC', 'MCC_DIFFKD': 'T_MCC',
}
CONDITION_DIFFUSION = {
    'CE_FEATURE': False, 'MCC_FEATURE': False,
    'CE_DIFFKD': True, 'MCC_DIFFKD': True,
}


def _sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(resolve_local_file(path)).open('rb') as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def _state_hash(state: dict) -> str:
    return reference._state_hash(state)


def _json_hash(value) -> str:
    return reference._json_hash(value)


def _read_json(path: str | Path) -> dict:
    return reference._read_json(path)


def _atomic_json(path: str | Path, value: dict) -> None:
    reference._atomic_json(Path(path), value)


def _atomic_npz(path: str | Path, **arrays) -> None:
    reference._atomic_npz(Path(path), **arrays)


def _atomic_torch(path: str | Path, value: dict) -> None:
    reference._atomic_torch(Path(path), value)


def _atomic_csv(path: str | Path, rows: list[dict]) -> None:
    reference._atomic_csv(Path(path), rows)


def _read_yaml(path: str | Path) -> dict:
    with Path(resolve_local_file(path)).open(encoding='utf-8') as stream:
        return yaml.safe_load(stream)


def _validate_config(config: dict) -> None:
    expected = {
        'experiment_id': 'MCC_TEACHER_ADAPT_DIFFKD_LOSO_V1',
        'dataset': 'BNCI2014004', 'selected_session': 'session_3',
        'source_variant': 'wide_fre_band', 'seed': 666,
        'reference_run': {
            'protocol_id': 'sample_utility_adaptive_weight_wideband_v1', 'seed': 666,
            'resolved_config_sha256':
                '6dfe184f7f6a5ddd8b75e57e6d67c77ce4e68f1bdf4aa349d558c0240539342f',
        },
        'protocol': {'folds': 9, 'train_subjects': 7, 'feedback_subjects': 1,
                     'target_subjects': 1,
                     'class_mapping': {'left_hand': 0, 'right_hand': 1},
                     'training_subject_data_only': True,
                     'feedback_is_diagnostic_only': True,
                     'target_evaluation_after_all_training': True},
        'teacher_adaptation': {
            'epochs': 10, 'optimizer': 'adam', 'lr': 1e-4,
            'betas': [0.9, 0.999], 'eps': 1e-8, 'weight_decay': 1e-6,
            'batch_size': 8, 'scheduler': 'none',
            'forward_mode': 'eval_with_autograd', 'precision': 'fp32',
            'gradient_clip_norm': 1.0,
            'mcc': {'temperature': 2.0, 'weight': 1.0,
                    'entropy_epsilon': 1e-5,
                    'row_denominator_epsilon': 1e-8},
        },
        'feature_space': {'pca_max_dimensions': 32,
                          'rank_relative_threshold': 1e-5,
                          'whitening_eigenvalue_floor_relative': 1e-4,
                          'teacher_native_dimension': 256,
                          'student_native_dimension': 512,
                          'student_layer_norm_epsilon': 1e-5,
                          'projection': 'fixed_gaussian_thin_qr',
                          'projection_seed_base': 20261030},
        'diffusion': {'conditioning': 'unconditional', 'steps': 1000,
                      'beta_start': 1e-4, 'beta_end': 0.02,
                      'training_updates': 5000, 'batch_size_cap': 128,
                      'hidden_dimensions': 128, 'time_embedding_dimensions': 32,
                      'optimizer': 'adamw', 'lr': 3e-4,
                      'betas': [0.9, 0.999], 'eps': 1e-8,
                      'weight_decay': 1e-4, 'gradient_clip_norm': 1.0,
                      'diagnostic_interval': 500,
                      'diagnostic_examples_per_noise_bin': 64,
                      'ddim_start_timestep': 250,
                      'ddim_timesteps': [250, 200, 150, 100, 50, 0],
                      'ddim_eta': 0.0},
        'student': {'architecture': 'ifnet', 'epochs': 100, 'batch_size': 16,
                    'optimizer': 'adamw', 'lr': 0.001, 'weight_decay': 0.01,
                    'scheduler': 'cosine_annealing_epoch_end',
                    'scheduler_t_max': 100, 'scheduler_eta_min': 0.0,
                    'feature_weight': 0.5, 'feature_start_epoch': 1,
                    'gradient_clip_norm': None, 'precision': 'fp32',
                    'augmentation': 'none'},
        'conditions': list(METHODS),
        'randomness': {
            'teacher_schedule_seed_base': 101666,
            'diffusion_seed_base': 202666,
            'student_schedule_seed': 666,
            'student_initialization_seed': 666,
            'student_feature_noise_seed_base': 303666,
            'paired_conditions_share_dropout_stream': True,
            'paired_diffkd_conditions_share_noise_stream': True,
        },
        'report': {'primary_metric': 'balanced_accuracy',
                   'bootstrap_draws': 10000, 'bootstrap_seed': 20261030,
                   'tie_threshold_pp': 1e-6},
        'storage': {
            'external_run_id': 'mcc_diffkd_loso_v1_seed666',
            'reuse_initial_ifnet_run_id': 'student_aware_diffusion_prior_loso_v1_seed666',
        },
    }

    def check(actual, wanted, route=''):
        if isinstance(wanted, dict):
            if not isinstance(actual, dict):
                raise ValueError(f'{route} must be a mapping')
            for key, value in wanted.items():
                check(actual.get(key), value, f'{route}.{key}' if route else key)
        elif actual != wanted:
            raise ValueError(f'fixed setting {route} must be {wanted!r}, got {actual!r}')

    check(config, expected)
    if config.get('reference_run', {}).get('protocol_id') != \
            'sample_utility_adaptive_weight_wideband_v1':
        raise ValueError('unexpected historical reference protocol')
    if config.get('reference_run', {}).get('resolved_config_sha256') != \
            '6dfe184f7f6a5ddd8b75e57e6d67c77ce4e68f1bdf4aa349d558c0240539342f':
        raise ValueError('historical resolved-config hash changed')


def _device(gpu: int) -> torch.device:
    if not torch.cuda.is_available():
        raise RuntimeError('formal MCC-DiffKD training requires CUDA')
    device = torch.device('cuda', int(gpu))
    torch.cuda.set_device(device)
    return device


def _run_roots(run_id: str, seed: int, fold: int) -> dict[str, Path]:
    roots = reference._run_roots(run_id, seed, fold)
    for name in ('run_result', 'run_cache', 'run_weights', 'run_seed_weights'):
        roots[name] = require_external_output(roots[name])
        roots[name].mkdir(parents=True, exist_ok=True)
    return roots


def _fold_fingerprint(config_path: Path, config: dict, cfg: dict,
                      identity: dict, roots: dict, models: dict,
                      initial: dict, run_id: str) -> str:
    relevant = {
        'experiment_id': config['experiment_id'],
        'run_id': run_id,
        'config_sha256': _sha256(config_path),
        'implementation_sha256': _sha256(__file__),
        'core_sha256': _sha256(PROJECT_ROOT / 'collab/mcc_diffkd.py'),
        'reference_runner_sha256': _sha256(reference.__file__),
        'pca_helper_sha256': _sha256(PROJECT_ROOT / 'collab/student_feature_prior.py'),
        'resolved_yaml_sha256': cfg['resolved_yaml_sha256'],
        'resolved_config_sha256': cfg['resolved_config_sha256'],
        'split_manifest_sha256': identity['sha256'],
        'source_manifest_sha256': identity['source_manifest_sha256'],
        'source_trials_sha256': identity['source_trial_manifest_sha256'],
        'teacher_input_manifest_sha256': _sha256(roots['input_root'] / 'mirepnet/manifest.json'),
        'student_input_manifest_sha256': _sha256(roots['input_root'] / 'ifnet/manifest.json'),
        'prepared_student_inputs_sha256': _sha256(roots['split_dir'] / 'student_inputs.npz'),
        'T0_checkpoint_sha256': models['teacher_checkpoint_sha256'],
        'T0_provenance_sha256': _sha256(roots['base_teacher'] / 'teacher_provenance.json'),
        'baseline_checkpoint_sha256': models['baseline_checkpoint_sha256'],
        'baseline_manifest_sha256': _sha256(roots['base_results'] / 'BASE_CE/manifest.json'),
        'initial_student_file_sha256': initial['file_sha256'],
        'initial_student_model_sha256': initial['model_sha256'],
        'seed': int(config['seed']),
        'fold': int(fold_from_identity(identity)),
    }
    return _json_hash(relevant)


def fold_from_identity(identity: dict) -> int:
    return int(identity['manifest']['target_subject_id']) - 1


def _load_initial_student(config: dict, identity: dict, roots: dict,
                          models: dict) -> dict:
    seed = int(config['seed'])
    reuse_id = config['storage']['reuse_initial_ifnet_run_id']
    path = WEIGHTS_ROOT / reuse_id / f'seed_{seed}' / f'ifnet_initial_seed_{seed}.pt'
    if not Path(resolve_local_file(path)).is_file():
        raise FileNotFoundError(f'canonical fresh IFNet initialization is missing: {path}')
    payload = torch.load(resolve_local_file(path), map_location='cpu', weights_only=False)
    if (int(payload.get('seed', -1)) != seed
            or payload.get('resolved_config_sha256') != models['baseline_manifest']['resolved_config_sha256']
            or payload.get('student_config_sha256') != _json_hash(models['student_cfg'])
            or _state_hash(payload.get('model', {})) != payload.get('model_sha256')):
        raise RuntimeError('canonical IFNet initialization provenance/hash mismatch')
    audit_path = (RESULTS_ROOT / 'distill' / reuse_id / f'seed_{seed}' / 'folds'
                  / roots['split_dir'].name / 'student_initialization_audit.json')
    audit = _read_json(audit_path)
    if (not audit.get('matched')
            or audit.get('initial_model_sha256') != payload.get('model_sha256')
            or audit.get('train_uid_sha256') != identity['uid_sha256']['train']
            or audit.get('schedule_sha256') != models['student_schedule_sha256']
            or audit.get('historical_warmup_sha256')
            != models['warmup_payload'].get('warmup_state_sha256')):
        raise RuntimeError('previous fresh-IFNet audit does not verify this fold/base CE trajectory')
    return {'path': str(path), 'file_sha256': _sha256(path),
            'model_sha256': payload['model_sha256'], 'payload': payload,
            'baseline_initialization_audit': str(audit_path),
            'baseline_initialization_audit_sha256': _sha256(audit_path)}


def _save_resolved(roots: dict, config_path: Path, config: dict,
                   cfg: dict, identity: dict, models: dict,
                   initial: dict, fingerprint: str) -> None:
    path = roots['run_result'] / 'resolved_run.json'
    record = {
        'status': 'running', 'experiment_id': config['experiment_id'],
        'run_fingerprint': fingerprint,
        'config_path': str(config_path), 'config_sha256': _sha256(config_path),
        'resolved_reference_yaml': cfg['resolved_config_path'],
        'resolved_reference_yaml_sha256': cfg['resolved_yaml_sha256'],
        'resolved_config_sha256': cfg['resolved_config_sha256'],
        'seed': int(config['seed']), 'target_subject_id': int(identity['manifest']['target_subject_id']),
        'feedback_subject_id': int(identity['manifest']['feedback_subject_id']),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'test_uid_sha256': identity['uid_sha256']['test'],
        'split_manifest': identity['path'], 'split_manifest_sha256': identity['sha256'],
        'source_manifest_sha256': identity['source_manifest_sha256'],
        'source_trial_manifest_sha256': identity['source_trial_manifest_sha256'],
        'T0_checkpoint': models['teacher_checkpoint'],
        'T0_checkpoint_sha256': models['teacher_checkpoint_sha256'],
        'T0_provenance': str(roots['base_teacher'] / 'teacher_provenance.json'),
        'initial_IFNet': initial['path'], 'initial_IFNet_file_sha256': initial['file_sha256'],
        'initial_IFNet_model_sha256': initial['model_sha256'],
        'baseline_CE_checkpoint': models['baseline_checkpoint'],
        'baseline_CE_checkpoint_sha256': models['baseline_checkpoint_sha256'],
        'student_schedule_sha256': models['student_schedule_sha256'],
        'student_config': models['student_cfg'], 'optimizer_defaults': models['optimizer_defaults'],
        'implementation_sha256': _sha256(__file__),
        'core_sha256': _sha256(PROJECT_ROOT / 'collab/mcc_diffkd.py'),
        'reference_runner_sha256': _sha256(reference.__file__),
        'created_at_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
    }
    if path.is_file():
        old = _read_json(path)
        if old.get('run_fingerprint') != fingerprint:
            raise RuntimeError(f'run identity changed for completed/in-progress fold: {path}')
        return
    _atomic_json(path, record)


def _progress(path: str | None, run_id: str, records: list[dict], total: int,
              status: str = 'running', message: str = '') -> None:
    if not path:
        return
    _atomic_json(path, {'run_id': run_id, 'status': status,
                        'total_cells': int(total), 'records': records,
                        'message': message, 'updated_at': time.time()})


def _stop_requested(stop_file: str | None) -> bool:
    return bool(stop_file and Path(stop_file).exists())


def _teacher_metrics(adapter, model, x: np.ndarray, y: np.ndarray,
                     device: torch.device) -> tuple[dict, np.ndarray, np.ndarray]:
    features, logits = reference._infer(adapter, model, x, device,
                                        batch_size=32, return_features=True)
    metrics = reference._metric_values(y, logits)
    return metrics, features, logits


def _capture_rng(device: torch.device) -> dict:
    return reference._capture_rng(device)


def _restore_rng(state: dict, device: torch.device) -> None:
    reference._restore_rng(state, device)


def _atomic_checkpoint(path: Path, run_fingerprint: str,
                       model: torch.nn.Module, **metadata) -> str:
    state = {name: value.detach().cpu().clone()
             for name, value in model.state_dict().items()}
    model_hash = _state_hash(state)
    reference._atomic_torch(path, {'run_fingerprint': run_fingerprint,
                                   'model_sha256': model_hash,
                                   'model': state, **metadata})
    return _sha256(path)


def _load_model_state(model: torch.nn.Module, path: Path,
                      run_fingerprint: str) -> dict:
    payload = torch.load(resolve_local_file(path), map_location='cpu', weights_only=False)
    if payload.get('run_fingerprint') != run_fingerprint:
        raise RuntimeError(f'checkpoint identity mismatch: {path}')
    if _state_hash(payload.get('model', {})) != payload.get('model_sha256'):
        raise RuntimeError(f'checkpoint state hash mismatch: {path}')
    model.load_state_dict(payload['model'], strict=True)
    model.eval()
    return payload


def _fit_t0_pca(t0_features: np.ndarray, roots: dict, config: dict,
                run_fingerprint: str) -> dict:
    path = roots['run_cache'] / 'fixed_T0_pca.npz'
    manifest_path = roots['run_cache'] / 'fixed_T0_pca.json'
    if path.exists() or manifest_path.exists():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete fixed PCA artifact; use the same run identity to resume')
        manifest = _read_json(manifest_path)
        if (manifest.get('run_fingerprint') != run_fingerprint
                or manifest.get('file_sha256') != _sha256(path)):
            raise RuntimeError('fixed PCA artifact identity/hash mismatch')
        with np.load(resolve_local_file(path), allow_pickle=False) as saved:
            transform = {key: np.asarray(saved[key]) for key in
                         ('mean', 'components', 'eigenvalues', 'scales')}
        transform.update({key: manifest['pca'][key] for key in
                          ('effective_rank', 'dimension', 'total_variance',
                           'retained_variance', 'retained_variance_fraction',
                           'maximum_eigenvalue')})
        return {'pca': transform, 'manifest': manifest}
    pca_cfg = config['feature_space']
    transform = fit_pca_whitener(
        t0_features,
        max_dimensions=int(pca_cfg['pca_max_dimensions']),
        rank_relative_threshold=float(pca_cfg['rank_relative_threshold']),
        eigenvalue_floor_relative=float(pca_cfg['whitening_eigenvalue_floor_relative']))
    if int(transform['dimension']) < 1:
        raise RuntimeError('T0 D_train PCA has zero effective dimension')
    reference._atomic_npz(path, **{key: transform[key] for key in
                                  ('mean', 'components', 'eigenvalues', 'scales')})
    manifest = {
        'run_fingerprint': run_fingerprint, 'file_sha256': _sha256(path),
        'fit_partition': 'D_train only', 'source_features': 'frozen T0 MIRepNet [N,256]',
        'train_feature_count': len(t0_features),
        'pca': {key: (int(value) if key in ('effective_rank', 'dimension') else float(value))
                for key, value in transform.items() if key not in
                ('mean', 'components', 'eigenvalues', 'scales')},
        'mean_sha256': hashlib.sha256(np.ascontiguousarray(transform['mean']).tobytes()).hexdigest(),
        'components_sha256': hashlib.sha256(
            np.ascontiguousarray(transform['components']).tobytes()).hexdigest(),
    }
    _atomic_json(manifest_path, manifest)
    return {'pca': transform, 'manifest': manifest}


def _teacher_diagnostic_rows(name: str, model: torch.nn.Module, adapter,
                             data: dict, identity: dict, device: torch.device,
                             epoch: int) -> list[dict]:
    rows = []
    for part in ('train', 'feedback'):
        metrics, _, logits = _teacher_metrics(
            adapter, model, data['teacher_x'][part], identity['labels'][part], device)
        mcc, stats = core.minimum_class_confusion(logits=torch.as_tensor(
            logits, dtype=torch.float32, device=device),
            temperature=2.0, entropy_epsilon=1e-5, row_epsilon=1e-8)
        counts = np.asarray(metrics['prediction_counts'], dtype=np.int64)
        rows.append({
            'teacher': name, 'epoch': int(epoch), 'partition': part,
            'accuracy': metrics['accuracy'], 'balanced_accuracy': metrics['balanced_accuracy'],
            'cross_entropy': metrics['cross_entropy'], 'kappa': metrics['kappa'],
            'prediction_count_0': int(counts[0]), 'prediction_count_1': int(counts[1]),
            'collapse': metrics['collapse'], 'mcc': float(mcc.detach().cpu()),
            'confusion_raw': json.dumps(stats['confusion'].cpu().tolist()),
            'confusion_row_normalized': json.dumps(
                stats['normalized_confusion'].cpu().tolist()),
            'confusion_row_mass': json.dumps(stats['row_mass'].cpu().tolist()),
            'clamped_rows': int(stats['clamped_rows'].cpu()),
            'mean_entropy': float(stats['entropy'].mean().cpu()),
            'mean_max_probability': float(stats['probabilities'].max(dim=1).values.mean().cpu()),
        })
    return rows


def _teacher_objective(logits: torch.Tensor, labels: torch.Tensor,
                       mcc_weight: float, mcc_config: dict) -> tuple[
                           torch.Tensor, torch.Tensor, torch.Tensor, dict | None]:
    """Compute CE and add MCC only when its configured weight is nonzero."""
    ce = F.cross_entropy(logits, labels)
    weight = float(mcc_weight)
    if weight == 0.0:
        return ce, ce, logits.new_zeros(()), None
    mcc, stats = core.minimum_class_confusion(
        logits, temperature=float(mcc_config['temperature']),
        entropy_epsilon=float(mcc_config['entropy_epsilon']),
        row_epsilon=float(mcc_config['row_denominator_epsilon']))
    return ce + weight * mcc, ce, mcc, stats


def _adapt_teacher(name: str, config: dict, identity: dict, roots: dict,
                   models: dict, data: dict, fingerprint: str,
                   device: torch.device, stop_file: str | None,
                   max_epochs: int | None = None) -> tuple[torch.nn.Module | None, dict]:
    if name not in TEACHER_NAMES:
        raise ValueError(f'unknown adapted teacher: {name}')
    code = name
    final_path = roots['run_weights'] / f'teacher_{code}.pt'
    manifest_path = roots['run_result'] / f'teacher_{code}.json'
    resume_path = roots['run_result'] / f'resume_teacher_{code}.pt'
    method_fingerprint = _json_hash({'run_fingerprint': fingerprint, 'teacher': code,
                                     'epochs': int(max_epochs or 10),
                                     'objective': 'CE' if code == 'T_CE' else 'CE+MCC'})
    if final_path.exists() or manifest_path.exists():
        if not (final_path.is_file() and manifest_path.is_file()):
            raise RuntimeError(f'{code}: incomplete terminal checkpoint/manifest')
        manifest = _read_json(manifest_path)
        if (manifest.get('method_fingerprint') != method_fingerprint
                or manifest.get('checkpoint_sha256') != _sha256(final_path)):
            raise RuntimeError(f'{code}: terminal teacher identity/hash mismatch')
        model = copy.deepcopy(models['teacher_model']).to(device)
        _load_model_state(model, final_path, method_fingerprint)
        return model, manifest

    model = copy.deepcopy(models['teacher_model']).to(device)
    model.eval()
    initial_hash = _state_hash(model.state_dict())
    t0_hash = models['teacher_checkpoint_sha256']
    teacher_cfg = config['teacher_adaptation']
    epochs = int(max_epochs if max_epochs is not None else teacher_cfg['epochs'])
    batch_size = int(teacher_cfg['batch_size'])
    seed = int(config['randomness']['teacher_schedule_seed_base']) + 1009 * fold_from_identity(identity)
    schedule = reference._epoch_schedule(len(identity['uids']['train']), batch_size, 10, seed)
    train_x = data['teacher_x']['train']
    train_y = identity['labels']['train']
    optimizer = torch.optim.Adam(
        model.parameters(), lr=float(teacher_cfg['lr']),
        betas=tuple(float(x) for x in teacher_cfg['betas']),
        eps=float(teacher_cfg['eps']), weight_decay=float(teacher_cfg['weight_decay']))
    start_epoch, history, diagnostics = 0, [], []
    if resume_path.is_file():
        saved = torch.load(resolve_local_file(resume_path), map_location='cpu', weights_only=False)
        if saved.get('method_fingerprint') != method_fingerprint:
            raise RuntimeError(f'{code}: teacher resume fingerprint mismatch')
        model.load_state_dict(saved['model'], strict=True)
        optimizer.load_state_dict(saved['optimizer'])
        _restore_rng(saved['rng'], device)
        start_epoch = int(saved['epoch'])
        history = list(saved.get('history', []))
        diagnostics = list(saved.get('diagnostics', []))
        if len(history) != start_epoch:
            raise RuntimeError(f'{code}: resume cursor/history mismatch')
    if start_epoch == 0:
        diagnostics.extend(_teacher_diagnostic_rows(
            code, model, models['teacher_adapter'], data, identity, device, 0))
    mcc_cfg = teacher_cfg['mcc']
    schedule_hash = reference._schedule_sha(schedule[:epochs], identity['uids']['train'])
    for epoch_index in range(start_epoch, epochs):
        model.eval()  # Keep Dropout off and BN running statistics fixed.
        total_ce = total_mcc = total_loss = total_grad_norm = 0.0
        total_samples = batches = 0
        first_gradient_record = {}
        for batch_index, indices in enumerate(schedule[epoch_index]):
            xb = torch.as_tensor(train_x[indices], dtype=torch.float32, device=device)
            yb = torch.as_tensor(train_y[indices], dtype=torch.long, device=device)
            optimizer.zero_grad(set_to_none=True)
            _, logits = models['teacher_adapter'].forward(model, xb)
            active_mcc_weight = (float(mcc_cfg['weight']) if code == 'T_MCC' else 0.0)
            loss, ce, mcc, mcc_stats = _teacher_objective(
                logits, yb, active_mcc_weight, mcc_cfg)
            weighted_mcc = active_mcc_weight * mcc
            clamped_rows = (int(mcc_stats['clamped_rows'].detach().cpu())
                            if mcc_stats is not None else 0)
            if not torch.isfinite(loss):
                raise FloatingPointError(f'{code}: non-finite teacher loss at epoch {epoch_index+1}')
            if epoch_index == 0 and batch_index == 0:
                parameters = [parameter for parameter in model.parameters() if parameter.requires_grad]
                ce_grads = torch.autograd.grad(ce, parameters, retain_graph=True, allow_unused=True)
                if code == 'T_MCC':
                    mcc_grads = torch.autograd.grad(weighted_mcc, parameters,
                                                    retain_graph=True, allow_unused=True)
                    ce_parts = [g.detach().norm() for g in ce_grads if g is not None]
                    mcc_parts = [g.detach().norm() for g in mcc_grads if g is not None]
                    ce_norm = torch.linalg.vector_norm(torch.stack(ce_parts))
                    mcc_norm = torch.linalg.vector_norm(torch.stack(mcc_parts))
                    dot = sum((g.detach() * h.detach()).sum()
                              for g, h in zip(ce_grads, mcc_grads)
                              if g is not None and h is not None)
                    cosine = dot / (ce_norm * mcc_norm).clamp_min(1e-12)
                    first_gradient_record = {
                        'first_batch_ce_grad_norm': float(ce_norm.detach().cpu()),
                        'first_batch_weighted_mcc_grad_norm': float(mcc_norm.detach().cpu()),
                        'first_batch_mcc_to_ce_grad_ratio': float(
                            (mcc_norm / ce_norm.clamp_min(1e-12)).detach().cpu()),
                        'first_batch_ce_mcc_gradient_cosine': float(cosine.detach().cpu()),
                    }
            loss.backward()
            grad_norm = torch.nn.utils.clip_grad_norm_(
                model.parameters(), float(teacher_cfg['gradient_clip_norm']))
            if not torch.isfinite(grad_norm):
                raise FloatingPointError(f'{code}: non-finite teacher gradient norm')
            optimizer.step()
            n = len(indices)
            total_samples += n
            batches += 1
            total_ce += float(ce.detach().cpu()) * n
            total_mcc += float(mcc.detach().cpu()) * n
            total_loss += float(loss.detach().cpu()) * n
            total_grad_norm += float(grad_norm.detach().cpu())
        epoch = epoch_index + 1
        record = {
            'teacher': code, 'epoch': epoch,
            'train_ce': total_ce / total_samples,
            'train_mcc': total_mcc / total_samples,
            'train_weighted_mcc': (float(mcc_cfg['weight']) * total_mcc / total_samples
                                   if code == 'T_MCC' else 0.0),
            'train_total': total_loss / total_samples,
            'mean_preclip_gradient_norm': total_grad_norm / batches,
            'n_train': total_samples, 'schedule_sha256': schedule_hash,
            'clamped_class_rows_last_batch': clamped_rows,
            **first_gradient_record,
        }
        history.append(record)
        if epoch in (1, 5, 10) or (max_epochs is not None and epoch == epochs):
            diagnostics.extend(_teacher_diagnostic_rows(
                code, model, models['teacher_adapter'], data, identity, device, epoch))
        _atomic_csv(roots['run_result'] / f'teacher_{code}_history.csv', history)
        _atomic_csv(roots['run_result'] / f'teacher_{code}_source_diagnostics.csv', diagnostics)
        state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
        _atomic_torch(resume_path, {
            'method_fingerprint': method_fingerprint, 'epoch': epoch,
            'model': state, 'optimizer': reference._cpu_tree(optimizer.state_dict()),
            'rng': _capture_rng(device), 'history': history, 'diagnostics': diagnostics,
            'initial_model_sha256': initial_hash, 'schedule_sha256': schedule_hash,
        })
        if _stop_requested(stop_file):
            return None, {'status': 'paused', 'completed_epoch': epoch}
    if _state_hash(models['teacher_model'].state_dict()) != _state_hash(
            torch.load(resolve_local_file(models['teacher_checkpoint']),
                       map_location='cpu', weights_only=False)['model']):
        raise RuntimeError('frozen T0 changed during teacher adaptation')
    checkpoint_sha = _atomic_checkpoint(
        final_path, method_fingerprint, model,
        teacher=code, final_epoch=epochs, initial_T0_checkpoint_sha256=t0_hash,
        initial_model_sha256=initial_hash)
    manifest = {
        'status': 'complete', 'teacher': code, 'method_fingerprint': method_fingerprint,
        'run_fingerprint': fingerprint, 'checkpoint': str(final_path),
        'checkpoint_sha256': checkpoint_sha, 'model_sha256': _state_hash(model.state_dict()),
        'initial_T0_checkpoint_sha256': t0_hash, 'initial_T0_model_sha256': initial_hash,
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'epochs': epochs, 'schedule_sha256': schedule_hash,
        'objective': ('CE' if code == 'T_CE' else 'CE + 1.0 * Minimum Class Confusion'),
        'optimizer': {'name': 'Adam', 'lr': teacher_cfg['lr'], 'betas': teacher_cfg['betas'],
                      'eps': teacher_cfg['eps'], 'weight_decay': teacher_cfg['weight_decay'],
                      'scheduler': 'none'},
        'forward_mode': 'eval with autograd; dropout off; BN running stats fixed',
        'history': str(roots['run_result'] / f'teacher_{code}_history.csv'),
        'source_diagnostics': str(roots['run_result'] / f'teacher_{code}_source_diagnostics.csv'),
    }
    _atomic_json(manifest_path, manifest)
    resume_path.unlink(missing_ok=True)
    model.eval()
    return model, manifest


def _make_fixed_q(dimension: int, fold: int, config: dict,
                  roots: dict, fingerprint: str) -> tuple[np.ndarray, dict]:
    path = roots['run_cache'] / 'fixed_student_Q.npz'
    manifest_path = roots['run_cache'] / 'fixed_student_Q.json'
    seed = int(config['feature_space']['projection_seed_base']) + int(fold)
    if path.exists() or manifest_path.exists():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete fixed IFNet projection artifact')
        manifest = _read_json(manifest_path)
        if (manifest.get('run_fingerprint') != fingerprint
                or int(manifest.get('dimension', -1)) != dimension
                or int(manifest.get('seed', -1)) != seed
                or manifest.get('file_sha256') != _sha256(path)):
            raise RuntimeError('fixed IFNet projection identity/hash mismatch')
        with np.load(resolve_local_file(path), allow_pickle=False) as saved:
            projection = np.asarray(saved['Q'], dtype=np.float64)
        return projection, manifest
    projection = core.fixed_student_projection(dimension, seed)
    path = require_external_output(path)
    _atomic_npz(path, Q=projection, Q_float32=projection.astype(np.float32))
    manifest = {
        'run_fingerprint': fingerprint, 'file_sha256': _sha256(path),
        'seed': seed, 'dimension': dimension, 'native_student_dimension': 512,
        'generator': 'numpy.default_rng float64 standard normal then thin QR; positive diag(R)',
        'Q_sha256_float64': hashlib.sha256(
            np.ascontiguousarray(projection).tobytes()).hexdigest(),
        'Q_sha256_float32': hashlib.sha256(
            np.ascontiguousarray(projection.astype(np.float32)).tobytes()).hexdigest(),
        'orthonormal_max_abs_error': float(np.max(np.abs(
            projection @ projection.T - np.eye(dimension)))),
        'centered_projection_rank': int(np.linalg.matrix_rank(
            projection @ (np.eye(512) - np.ones((512, 512)) / 512.0), tol=1e-10)),
    }
    _atomic_json(manifest_path, manifest)
    return projection, manifest


def _load_or_make_teacher_targets(config: dict, identity: dict, roots: dict,
                                 models: dict, data: dict, pca: dict,
                                 fingerprint: str, teachers: dict,
                                 teacher_manifests: dict,
                                 device: torch.device) -> tuple[dict, dict]:
    path = roots['run_cache'] / 'teacher_features_train_feedback.npz'
    manifest_path = roots['run_cache'] / 'teacher_features_train_feedback.json'
    if path.exists() or manifest_path.exists():
        if not (path.is_file() and manifest_path.is_file()):
            raise RuntimeError('incomplete teacher feature cache')
        manifest = _read_json(manifest_path)
        if (manifest.get('run_fingerprint') != fingerprint
                or manifest.get('file_sha256') != _sha256(path)
                or manifest.get('train_uid_sha256') != identity['uid_sha256']['train']
                or manifest.get('feedback_uid_sha256') != identity['uid_sha256']['feedback']):
            raise RuntimeError('teacher feature cache does not match this fold identity')
        with np.load(resolve_local_file(path), allow_pickle=False) as source:
            return {key: np.asarray(source[key]) for key in source.files}, manifest

    arrays = {
        'train_uids': identity['uids']['train'],
        'train_y': identity['labels']['train'],
        'feedback_uids': identity['uids']['feedback'],
        'feedback_y': identity['labels']['feedback'],
    }
    for name, model in teachers.items():
        for part in ('train', 'feedback'):
            _, features, logits = _teacher_metrics(
                models['teacher_adapter'], model, data['teacher_x'][part],
                identity['labels'][part], device)
            whitened = apply_pca_whitener(features, pca)
            arrays[f'{name}_{part}_features_native'] = features
            arrays[f'{name}_{part}_features_whitened'] = whitened
            arrays[f'{name}_{part}_logits'] = logits
    pca_path = roots['run_cache'] / 'fixed_T0_pca.npz'
    arrays['pca_train_uid_sha256'] = np.asarray(identity['uid_sha256']['train'])
    _atomic_npz(path, **arrays)
    norm_audit = {}
    upper_bound = math.sqrt(512.0)
    for name in TEACHER_NAMES:
        norms = np.linalg.norm(arrays[f'{name}_train_features_whitened'], axis=1)
        norm_audit[name] = {
            'train_norm_quantiles_50_90_95_99_max': [
                float(x) for x in np.quantile(norms, [0.5, 0.9, 0.95, 0.99, 1.0])],
            'fraction_above_student_projection_bound_sqrt512': float(np.mean(norms > upper_bound)),
            'student_projection_norm_upper_bound': upper_bound,
        }
    manifest = {
        'run_fingerprint': fingerprint, 'file_sha256': _sha256(path),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'pca_file': str(pca_path), 'pca_file_sha256': _sha256(pca_path),
        'teacher_checkpoint_sha256': {
            name: teacher_manifests[name]['checkpoint_sha256'] for name in TEACHER_NAMES},
        'fit_partition': 'D_train; feedback features are diagnostic only',
        'feature_tap': 'MIRepNet pooled 256-D feature immediately before clshead',
        'whitened_dimension': int(pca['dimension']),
        'feature_source': {
            name: f'the corresponding adapted teacher at final epoch '
                  f"{teacher_manifests[name]['epochs']}"
            for name in TEACHER_NAMES
        },
        'teacher_feature_norm_audit_D_train_only': norm_audit,
    }
    _atomic_json(manifest_path, manifest)
    return arrays, manifest


def _prior_init_state(dimension: int, seed: int, config: dict) -> tuple[dict, str]:
    with torch.random.fork_rng(devices=[]):
        torch.manual_seed(int(seed))
        model = core.EpsilonPredictor(
            dimension,
            hidden_dimensions=int(config['diffusion']['hidden_dimensions']),
            time_dimensions=int(config['diffusion']['time_embedding_dimensions']))
        state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
    return state, _state_hash(state)


def _diagnostic_tables(features: dict[str, np.ndarray], config: dict,
                       fold: int, device: torch.device,
                       alpha: torch.Tensor, sigma: torch.Tensor) -> dict:
    rng = np.random.default_rng(909000 + int(fold))
    count_cfg = int(config['diffusion']['diagnostic_examples_per_noise_bin'])
    result = {}
    for part in ('train', 'feedback'):
        values = np.asarray(features[f'T_CE_{part}_features_whitened'])
        # A per-teacher table is created at the call site; this first pass only
        # determines fixed UID-independent indexes, times, and Gaussian noise.
        tables = []
        for low, high in NOISE_BINS:
            count = min(count_cfg, len(values))
            ids = np.asarray(rng.choice(len(values), size=count, replace=False), dtype=np.int64)
            timesteps = np.asarray(rng.integers(low, high + 1, size=count), dtype=np.int64)
            noise = rng.standard_normal((count, values.shape[1])).astype(np.float32)
            tables.append({'ids': ids, 'timesteps': timesteps, 'noise': noise})
        result[part] = tables
    return result


def _measure_prior(prior: core.EpsilonPredictor, features: np.ndarray,
                   tables: list[dict], alpha: torch.Tensor,
                   sigma: torch.Tensor, device: torch.device,
                   part: str) -> dict:
    prior.eval()
    rows = {}
    with torch.no_grad():
        for (low, high), table in zip(NOISE_BINS, tables):
            ids = table['ids']
            t_np, eps_np = table['timesteps'], table['noise']
            t = torch.as_tensor(t_np, dtype=torch.long, device=device)
            eps = torch.as_tensor(eps_np, dtype=torch.float32, device=device)
            clean = torch.as_tensor(features[ids], dtype=torch.float32, device=device)
            noisy = alpha[t, None] * clean + sigma[t, None] * eps
            prediction = prior(noisy, t)
            mse = F.mse_loss(prediction, eps)
            rows[f'{part}_mse_t{low}_{high}'] = float(mse.cpu())
            rows[f'{part}_zero_predictor_mse_t{low}_{high}'] = float(eps.square().mean().cpu())
    return rows


def _fit_denoiser(name: str, config: dict, identity: dict, roots: dict,
                  teacher_features: dict, fingerprint: str,
                  initial_state: dict, initial_hash: str,
                  device: torch.device, stop_file: str | None,
                  max_updates: int | None = None) -> tuple[core.EpsilonPredictor | None, dict]:
    if name not in TEACHER_NAMES:
        raise ValueError(f'unknown teacher prior: {name}')
    final_path = roots['run_weights'] / f'prior_{name}.pt'
    manifest_path = roots['run_result'] / f'prior_{name}.json'
    resume_path = roots['run_result'] / f'resume_prior_{name}.pt'
    diffusion_cfg = config['diffusion']
    total_updates = int(max_updates if max_updates is not None else diffusion_cfg['training_updates'])
    method_fingerprint = _json_hash({
        'run_fingerprint': fingerprint, 'teacher': name,
        'teacher_features_sha256': _sha256(roots['run_cache'] / 'teacher_features_train_feedback.npz'),
        'prior_initial_model_sha256': initial_hash, 'updates': total_updates,
    })
    if final_path.exists() or manifest_path.exists():
        if not (final_path.is_file() and manifest_path.is_file()):
            raise RuntimeError(f'{name}: incomplete terminal prior checkpoint/manifest')
        manifest = _read_json(manifest_path)
        if (manifest.get('method_fingerprint') != method_fingerprint
                or manifest.get('checkpoint_sha256') != _sha256(final_path)):
            raise RuntimeError(f'{name}: terminal prior identity/hash mismatch')
        prior = core.EpsilonPredictor(
            int(teacher_features[f'{name}_train_features_whitened'].shape[1]),
            hidden_dimensions=int(diffusion_cfg['hidden_dimensions']),
            time_dimensions=int(diffusion_cfg['time_embedding_dimensions'])).to(device)
        _load_model_state(prior, final_path, method_fingerprint)
        for parameter in prior.parameters():
            parameter.requires_grad_(False)
            parameter.grad = None
        return prior, manifest

    train_features_np = np.asarray(teacher_features[f'{name}_train_features_whitened'], dtype=np.float32)
    feedback_features_np = np.asarray(teacher_features[f'{name}_feedback_features_whitened'], dtype=np.float32)
    dimension = int(train_features_np.shape[1])
    prior = core.EpsilonPredictor(
        dimension, hidden_dimensions=int(diffusion_cfg['hidden_dimensions']),
        time_dimensions=int(diffusion_cfg['time_embedding_dimensions'])).to(device)
    prior.load_state_dict(initial_state, strict=True)
    prior.train()
    optimizer = torch.optim.AdamW(
        prior.parameters(), lr=float(diffusion_cfg['lr']),
        betas=tuple(float(x) for x in diffusion_cfg['betas']),
        eps=float(diffusion_cfg['eps']), weight_decay=float(diffusion_cfg['weight_decay']))
    betas, alpha, sigma = core.ddpm_schedule(
        int(diffusion_cfg['steps']), float(diffusion_cfg['beta_start']),
        float(diffusion_cfg['beta_end']), device)
    train_features = torch.as_tensor(train_features_np, dtype=torch.float32, device=device)
    rng_seed = int(config['randomness']['diffusion_seed_base']) + 101 * fold_from_identity(identity)
    generator = torch.Generator(device=device)
    generator.manual_seed(rng_seed)
    fixed_tables = {}
    # Generate one fixed diagnostic noise table per teacher; it is never used
    # for fitting and does not consume the training generator.
    for part, values in (('train', train_features_np), ('feedback', feedback_features_np)):
        rng = np.random.default_rng(909000 + 17 * fold_from_identity(identity)
                                    + (0 if part == 'train' else 1))
        fixed_tables[part] = []
        for low, high in NOISE_BINS:
            count = min(int(diffusion_cfg['diagnostic_examples_per_noise_bin']), len(values))
            ids = rng.choice(len(values), size=count, replace=False).astype(np.int64)
            times = rng.integers(low, high + 1, size=count, dtype=np.int64)
            eps = rng.standard_normal((count, dimension)).astype(np.float32)
            fixed_tables[part].append({'ids': ids, 'timesteps': times, 'noise': eps})
    start_update, history = 0, []
    if resume_path.is_file():
        saved = torch.load(resolve_local_file(resume_path), map_location='cpu', weights_only=False)
        if saved.get('method_fingerprint') != method_fingerprint:
            raise RuntimeError(f'{name}: prior resume fingerprint mismatch')
        prior.load_state_dict(saved['model'], strict=True)
        optimizer.load_state_dict(saved['optimizer'])
        generator.set_state(saved['generator_state'].cpu())
        _restore_rng(saved['rng'], device)
        start_update = int(saved['update'])
        history = list(saved.get('history', []))
        if history and int(history[-1]['update']) != start_update:
            raise RuntimeError(f'{name}: prior resume update/history cursor mismatch')
    if start_update == 0:
        diagnostic = {'teacher': name, 'update': 0}
        diagnostic.update(_measure_prior(prior, train_features_np, fixed_tables['train'],
                                         alpha, sigma, device, 'train'))
        diagnostic.update(_measure_prior(prior, feedback_features_np, fixed_tables['feedback'],
                                         alpha, sigma, device, 'feedback'))
        history.append(diagnostic)
    batch_size = min(int(diffusion_cfg['batch_size_cap']), len(train_features_np))
    grad_clip = float(diffusion_cfg['gradient_clip_norm'])
    for update in range(start_update + 1, total_updates + 1):
        ids = torch.randint(len(train_features), (batch_size,), device=device, generator=generator)
        t = torch.randint(1, int(diffusion_cfg['steps']) + 1, (batch_size,),
                          dtype=torch.long, device=device, generator=generator)
        eps = torch.randn((batch_size, dimension), dtype=torch.float32,
                          device=device, generator=generator)
        noisy = alpha[t, None] * train_features[ids] + sigma[t, None] * eps
        predicted = prior(noisy, t)
        loss = F.mse_loss(predicted, eps)
        if not torch.isfinite(loss):
            raise FloatingPointError(f'{name}: non-finite diffusion loss at update {update}')
        optimizer.zero_grad(set_to_none=True)
        loss.backward()
        grad_norm = torch.nn.utils.clip_grad_norm_(prior.parameters(), grad_clip)
        if not torch.isfinite(grad_norm):
            raise FloatingPointError(f'{name}: non-finite prior gradient at update {update}')
        optimizer.step()
        if update % int(diffusion_cfg['diagnostic_interval']) == 0 or update == total_updates:
            record = {'teacher': name, 'update': update,
                      'train_loss_batch': float(loss.detach().cpu()),
                      'mean_preclip_gradient_norm': float(grad_norm.detach().cpu())}
            record.update(_measure_prior(prior, train_features_np, fixed_tables['train'],
                                         alpha, sigma, device, 'train'))
            record.update(_measure_prior(prior, feedback_features_np, fixed_tables['feedback'],
                                         alpha, sigma, device, 'feedback'))
            history.append(record)
            _atomic_csv(roots['run_result'] / f'prior_{name}_history.csv', history)
            state = {key: value.detach().cpu().clone() for key, value in prior.state_dict().items()}
            _atomic_torch(resume_path, {
                'method_fingerprint': method_fingerprint, 'update': update,
                'model': state, 'optimizer': reference._cpu_tree(optimizer.state_dict()),
                'generator_state': generator.get_state().cpu(), 'rng': _capture_rng(device),
                'history': history, 'prior_initial_model_sha256': initial_hash,
                'teacher_features_sha256': _sha256(
                    roots['run_cache'] / 'teacher_features_train_feedback.npz'),
            })
            if _stop_requested(stop_file):
                return None, {'status': 'paused', 'completed_update': update}
    prior.eval()
    for parameter in prior.parameters():
        parameter.requires_grad_(False)
        parameter.grad = None
    checkpoint_sha = _atomic_checkpoint(
        final_path, method_fingerprint, prior, teacher=name, final_update=total_updates,
        prior_initial_model_sha256=initial_hash)
    manifest = {
        'status': 'complete', 'teacher': name, 'method_fingerprint': method_fingerprint,
        'run_fingerprint': fingerprint, 'checkpoint': str(final_path),
        'checkpoint_sha256': checkpoint_sha, 'model_sha256': _state_hash(prior.state_dict()),
        'initial_model_sha256': initial_hash, 'final_update': total_updates,
        'train_uid_sha256': identity['uid_sha256']['train'],
        'feedback_uid_sha256': identity['uid_sha256']['feedback'],
        'conditioning': 'unconditional; no labels/logits/student features',
        'optimizer': {'name': 'AdamW', 'lr': diffusion_cfg['lr'],
                      'betas': diffusion_cfg['betas'], 'eps': diffusion_cfg['eps'],
                      'weight_decay': diffusion_cfg['weight_decay']},
        'noise_schedule': {'steps': diffusion_cfg['steps'],
                           'beta_start': diffusion_cfg['beta_start'],
                           'beta_end': diffusion_cfg['beta_end']},
        'history': str(roots['run_result'] / f'prior_{name}_history.csv'),
    }
    _atomic_json(manifest_path, manifest)
    resume_path.unlink(missing_ok=True)
    return prior, manifest


def _student_feedback_diagnostic(condition: str, model: torch.nn.Module,
                                 models: dict, data: dict, identity: dict,
                                 teacher_features: dict, projection: torch.Tensor,
                                 prior: core.EpsilonPredictor | None,
                                 alpha: torch.Tensor, sigma: torch.Tensor,
                                 device: torch.device, epoch: int) -> dict:
    native, logits = reference._infer(
        models['student_adapter'], model, data['student_x']['feedback'], device,
        batch_size=64, return_features=True)
    metrics = reference._metric_values(identity['labels']['feedback'], logits)
    teacher_name = CONDITION_TEACHER[condition]
    target = teacher_features[f'{teacher_name}_feedback_features_whitened']
    native_t = torch.as_tensor(native, dtype=torch.float32, device=device)
    projected = core.student_features_to_vector(native_t, projection)
    target_t = torch.as_tensor(target, dtype=torch.float32, device=device)
    direct_loss = F.mse_loss(projected, target_t)
    if CONDITION_DIFFUSION[condition]:
        rng = np.random.default_rng(808000 + epoch)
        fixed_noise = torch.as_tensor(
            rng.standard_normal(tuple(projected.shape)).astype(np.float32), device=device)
        assert prior is not None
        with torch.no_grad():
            denoised = core.ddim_denoise(projected, fixed_noise, prior, alpha, sigma)
            transferred_loss = F.mse_loss(denoised, target_t)
        sensitivity = None
    else:
        transferred_loss = direct_loss
        sensitivity = 1.0
    return {
        'condition': condition, 'epoch': int(epoch), 'partition': 'feedback',
        'accuracy': metrics['accuracy'], 'balanced_accuracy': metrics['balanced_accuracy'],
        'cross_entropy': metrics['cross_entropy'], 'kappa': metrics['kappa'],
        'prediction_counts': json.dumps(metrics['prediction_counts']),
        'collapse': metrics['collapse'],
        'direct_feature_mse': float(direct_loss.detach().cpu()),
        'transferred_feature_mse': float(transferred_loss.detach().cpu()),
        'student_feature_effective_rank': float(reference._effective_rank(native)),
        'aux_feature_effective_rank': float(reference._effective_rank(
            projected.detach().cpu().numpy())),
        'fixed_map_sensitivity_identity_reference': sensitivity,
    }


def _student_feature_loss(auxiliary: torch.Tensor, targets: torch.Tensor,
                         feature_weight: float, use_diffusion: bool,
                         prior: core.EpsilonPredictor | None,
                         alpha: torch.Tensor | None, sigma: torch.Tensor | None,
                         generator: torch.Generator | None) -> torch.Tensor | None:
    """Shared loss dispatch; zero weight skips both DDIM and noise sampling."""
    if float(feature_weight) == 0.0:
        return None
    transferred = auxiliary
    if use_diffusion:
        if prior is None or alpha is None or sigma is None or generator is None:
            raise RuntimeError('positive DiffKD weight requires frozen prior, schedule, and RNG')
        noise = torch.randn(auxiliary.shape, dtype=auxiliary.dtype,
                            device=auxiliary.device, generator=generator)
        transferred = core.ddim_denoise(auxiliary, noise, prior, alpha, sigma)
    return F.mse_loss(transferred, targets)


def _fit_student(condition: str, config: dict, identity: dict, roots: dict,
                 models: dict, initial: dict, data: dict,
                 teacher_features: dict, feature_manifest: dict,
                 projection_np: np.ndarray, projection_manifest: dict,
                 prior: core.EpsilonPredictor | None, prior_manifest: dict | None,
                 fingerprint: str, device: torch.device, stop_file: str | None,
                 max_epochs: int | None = None) -> tuple[dict | None, dict]:
    if condition not in METHODS:
        raise ValueError(f'unknown student condition: {condition}')
    teacher_name = CONDITION_TEACHER[condition]
    use_diffusion = CONDITION_DIFFUSION[condition]
    if use_diffusion and (prior is None or prior_manifest is None):
        raise RuntimeError(f'{condition} requires its frozen teacher diffusion prior')
    if prior is not None and any(parameter.requires_grad or parameter.grad is not None
                                 for parameter in prior.parameters()):
        raise RuntimeError(f'{condition}: prior must be frozen and have no stale gradients')
    final_path = roots['run_weights'] / f'student_{condition}.pt'
    manifest_path = roots['run_result'] / f'student_{condition}.json'
    resume_path = roots['run_result'] / f'resume_student_{condition}.pt'
    student_cfg = config['student']
    epochs = int(max_epochs if max_epochs is not None else student_cfg['epochs'])
    method_fingerprint = _json_hash({
        'run_fingerprint': fingerprint, 'condition': condition,
        'teacher': teacher_name, 'teacher_features_sha256': feature_manifest['file_sha256'],
        'Q_sha256': projection_manifest['Q_sha256_float32'],
        'prior_sha256': prior_manifest.get('checkpoint_sha256') if prior_manifest else None,
        'initial_IFNet_sha256': initial['model_sha256'],
        'student_schedule_sha256': models['student_schedule_sha256'],
        'epochs': epochs, 'feature_weight': float(student_cfg['feature_weight']),
    })
    if final_path.exists() or manifest_path.exists():
        if not (final_path.is_file() and manifest_path.is_file()):
            raise RuntimeError(f'{condition}: incomplete terminal Student checkpoint/manifest')
        manifest = _read_json(manifest_path)
        if (manifest.get('method_fingerprint') != method_fingerprint
                or manifest.get('checkpoint_sha256') != _sha256(final_path)):
            raise RuntimeError(f'{condition}: terminal Student identity/hash mismatch')
        return None, manifest

    model, optimizer, scheduler, _ = reference._new_paired_student(
        models, initial, int(config['randomness']['student_initialization_seed']), device)
    projection = torch.as_tensor(projection_np.astype(np.float32), device=device)
    alpha = sigma = None
    if use_diffusion:
        _, alpha, sigma = core.ddpm_schedule(
            int(config['diffusion']['steps']), float(config['diffusion']['beta_start']),
            float(config['diffusion']['beta_end']), device)
    noise_generator = None
    if use_diffusion:
        noise_generator = torch.Generator(device=device)
        noise_generator.manual_seed(
            int(config['randomness']['student_feature_noise_seed_base'])
            + 1009 * fold_from_identity(identity))
    start_epoch, history, diagnostics = 0, [], []
    if resume_path.is_file():
        saved = torch.load(resolve_local_file(resume_path), map_location='cpu', weights_only=False)
        if saved.get('method_fingerprint') != method_fingerprint:
            raise RuntimeError(f'{condition}: student resume fingerprint mismatch')
        model.load_state_dict(saved['model'], strict=True)
        optimizer.load_state_dict(saved['optimizer'])
        scheduler.load_state_dict(saved['scheduler'])
        _restore_rng(saved['rng'], device)
        if noise_generator is not None:
            noise_generator.set_state(saved['noise_generator_state'].cpu())
        start_epoch = int(saved['epoch'])
        history = list(saved.get('history', []))
        diagnostics = list(saved.get('diagnostics', []))
        if len(history) != start_epoch:
            raise RuntimeError(f'{condition}: student resume cursor/history mismatch')
    train_x = data['student_x']['train']
    train_y = identity['labels']['train']
    target_features_np = np.asarray(
        teacher_features[f'{teacher_name}_train_features_whitened'], dtype=np.float32)
    if (len(train_x) != len(train_y) or len(train_y) != len(target_features_np)
            or not np.array_equal(identity['uids']['train'], teacher_features['train_uids'])):
        raise RuntimeError(f'{condition}: UID-aligned student/teacher features do not match')
    target_features = torch.as_tensor(target_features_np, dtype=torch.float32, device=device)
    schedule = models['student_schedule']
    feature_weight = float(student_cfg['feature_weight'])
    if epochs > len(schedule):
        raise RuntimeError('student schedule does not cover the configured epoch budget')
    if start_epoch == 0:
        if _state_hash(model.state_dict()) != initial['model_sha256']:
            raise RuntimeError(f'{condition}: fresh IFNet does not equal shared I_s')
    for epoch_index in range(start_epoch, epochs):
        model.train()
        seen = steps = 0
        total_ce = total_feature = total_loss = total_grad_norm = 0.0
        first_gradient_record = {}
        for batch_index, indices in enumerate(schedule[epoch_index]):
            xb = torch.as_tensor(train_x[indices], dtype=torch.float32, device=device)
            yb = torch.as_tensor(train_y[indices], dtype=torch.long, device=device)
            targets = target_features[torch.as_tensor(indices, dtype=torch.long, device=device)]
            optimizer.zero_grad(set_to_none=True)
            native, logits = models['student_adapter'].forward(model, xb)
            if native.shape[1] != 512:
                raise RuntimeError(f'IFNet feature tap changed: expected 512, got {native.shape}')
            auxiliary = core.student_features_to_vector(native, projection)
            ce = F.cross_entropy(logits, yb)
            feature_loss = _student_feature_loss(
                auxiliary, targets, feature_weight, use_diffusion,
                prior, alpha, sigma, noise_generator)
            if feature_loss is None:
                loss = ce
                feature_loss_value = 0.0
            else:
                loss = ce + feature_weight * feature_loss
                feature_loss_value = float(feature_loss.detach().cpu())
            if not torch.isfinite(loss):
                raise FloatingPointError(f'{condition}: non-finite student loss')
            if epoch_index == 0 and batch_index == 0:
                native.retain_grad()
            loss.backward()
            if epoch_index == 0 and batch_index == 0:
                native_grad = (float(native.grad.detach().norm().cpu())
                               if native.grad is not None else 0.0)
                stem_norms = [parameter.grad.detach().norm() for name, parameter
                              in model.named_parameters()
                              if name.startswith('stem.') and parameter.grad is not None]
                stem_norm = (float(torch.linalg.vector_norm(torch.stack(stem_norms)).cpu())
                             if stem_norms else 0.0)
                first_gradient_record = {
                    'first_batch_native_feature_grad_norm': native_grad,
                    'first_batch_stem_grad_norm': stem_norm,
                }
                if native_grad <= 0 or stem_norm <= 0:
                    raise RuntimeError(f'{condition}: feature loss did not reach IFNet stem')
            grad_parts = [parameter.grad.detach().norm() for parameter in model.parameters()
                          if parameter.grad is not None]
            grad_norm = torch.linalg.vector_norm(torch.stack(grad_parts))
            if not torch.isfinite(grad_norm):
                raise FloatingPointError(f'{condition}: non-finite IFNet gradient')
            optimizer.step()
            count = len(indices)
            seen += count
            steps += 1
            total_ce += float(ce.detach().cpu()) * count
            total_feature += feature_loss_value * count
            total_loss += float(loss.detach().cpu()) * count
            total_grad_norm += float(grad_norm.detach().cpu())
        epoch = epoch_index + 1
        lr_used = float(optimizer.param_groups[0]['lr'])
        scheduler.step()
        row = {
            'condition': condition, 'teacher': teacher_name,
            'student_loss_type': 'CE + 0.5 * DiffKD_feature_MSE'
            if use_diffusion else 'CE + 0.5 * paired_feature_MSE',
            'epoch': epoch, 'train_ce': total_ce / seen,
            'train_feature_mse': total_feature / seen,
            'weighted_feature_mse': feature_weight * total_feature / seen,
            'train_total': total_loss / seen,
            'mean_gradient_norm': total_grad_norm / steps,
            'lr_used': lr_used, 'lr_next': float(scheduler.get_last_lr()[0]),
            'n_train': seen, 'schedule_sha256': models['student_schedule_sha256'],
            **first_gradient_record,
        }
        history.append(row)
        if epoch in (1, 10, 50, 100) or (max_epochs is not None and epoch == epochs):
            diagnostics.append(_student_feedback_diagnostic(
                condition, model, models, data, identity, teacher_features,
                projection, prior, alpha, sigma, device, epoch))
        _atomic_csv(roots['run_result'] / f'student_{condition}_history.csv', history)
        if diagnostics:
            _atomic_csv(roots['run_result'] / f'student_{condition}_feedback_diagnostics.csv', diagnostics)
        _atomic_torch(resume_path, {
            'method_fingerprint': method_fingerprint, 'epoch': epoch,
            'model': reference._cpu_tree(model.state_dict()),
            'optimizer': reference._cpu_tree(optimizer.state_dict()),
            'scheduler': scheduler.state_dict(), 'rng': _capture_rng(device),
            'noise_generator_state': (noise_generator.get_state().cpu()
                                      if noise_generator is not None else None),
            'history': history, 'diagnostics': diagnostics,
            'initial_model_sha256': initial['model_sha256'],
            'student_schedule_sha256': models['student_schedule_sha256'],
        })
        if _stop_requested(stop_file):
            return {'status': 'paused', 'completed_epoch': epoch}, {}
    state_hash = _state_hash(model.state_dict())
    checkpoint_sha = _atomic_checkpoint(
        final_path, method_fingerprint, model, condition=condition,
        final_epoch=epochs, initial_model_sha256=initial['model_sha256'])
    manifest = {
        'status': 'complete', 'condition': condition, 'teacher': teacher_name,
        'uses_diffusion': use_diffusion, 'method_fingerprint': method_fingerprint,
        'run_fingerprint': fingerprint, 'checkpoint': str(final_path),
        'checkpoint_sha256': checkpoint_sha, 'model_sha256': state_hash,
        'initial_model_sha256': initial['model_sha256'],
        'initial_model_file': initial['path'],
        'final_epoch': epochs, 'train_uid_sha256': identity['uid_sha256']['train'],
        'student_schedule_sha256': models['student_schedule_sha256'],
        'teacher_features_file': str(roots['run_cache'] / 'teacher_features_train_feedback.npz'),
        'teacher_features_sha256': feature_manifest['file_sha256'],
        'fixed_Q_file': str(roots['run_cache'] / 'fixed_student_Q.npz'),
        'fixed_Q_sha256': projection_manifest['file_sha256'],
        'prior_checkpoint_sha256': prior_manifest.get('checkpoint_sha256') if prior_manifest else None,
        'loss': {'ce_weight': 1.0, 'feature_weight': feature_weight,
                 'feature': 'mean over batch and whitened dimensions', 'start_epoch': 1},
        'student_config': models['student_cfg'],
        'optimizer': {'name': 'AdamW', 'lr': models['student_cfg']['lr'],
                      'weight_decay': models['student_cfg']['weight_decay'],
                      'defaults': models['optimizer_defaults']['student'],
                      'scheduler': 'CosineAnnealingLR, T_max=100, epoch-end'},
        'history': str(roots['run_result'] / f'student_{condition}_history.csv'),
        'feedback_diagnostics': str(
            roots['run_result'] / f'student_{condition}_feedback_diagnostics.csv'),
    }
    _atomic_json(manifest_path, manifest)
    resume_path.unlink(missing_ok=True)
    return None, manifest


def _train_fold(config_path: Path, run_id: str, fold: int, config: dict,
                device: torch.device, stop_file: str | None,
                smoke: bool = False) -> dict:
    seed = int(config['seed'])
    roots = _run_roots(run_id, seed, fold)
    resolved, cfg = reference._load_resolved_reference(config)
    identity = reference._load_fold_identity(
        config, fold, roots, cfg['resolved_config_sha256'], include_test_labels=False)
    models = reference._load_reference_checkpoints(
        config, resolved, cfg, identity, roots, device)
    initial = _load_initial_student(config, identity, roots, models)
    fingerprint = _fold_fingerprint(config_path, config, cfg, identity, roots,
                                    models, initial, run_id)
    _save_resolved(roots, config_path, config, cfg, identity, models, initial, fingerprint)
    complete_path = roots['run_result'] / 'fold_complete.json'
    if complete_path.is_file():
        old = _read_json(complete_path)
        if old.get('run_fingerprint') != fingerprint or old.get('status') != 'complete':
            raise RuntimeError(f'fold {fold+1}: completion record belongs to another run')
        return {'status': 'complete', 'target_subject_id': fold + 1,
                'run_fingerprint': fingerprint, 'resumed_existing': True}
    data = reference._load_fold_data(identity, roots, resolved['student_config'], device)
    smoke_audits = {}
    if smoke:
        smoke_audits['mcc_equation'] = _audit_mcc_formula(device)
        smoke_audits['zero_weight_mcc_trajectory'] = _audit_zero_mcc_trajectory(
            config, identity, models, data, device)
    t0_train, t0_train_logits = _teacher_metrics(
        models['teacher_adapter'], models['teacher_model'], data['teacher_x']['train'],
        identity['labels']['train'], device)[1:]
    t0_feedback, t0_feedback_logits = _teacher_metrics(
        models['teacher_adapter'], models['teacher_model'], data['teacher_x']['feedback'],
        identity['labels']['feedback'], device)[1:]
    if t0_train.shape != (len(identity['uids']['train']), 256):
        raise RuntimeError(f'fold {fold+1}: MIRepNet feature tap is not the audited [N,256] vector')
    pca_record = _fit_t0_pca(t0_train, roots, config, fingerprint)
    pca = pca_record['pca']
    npz_path = roots['run_cache'] / 'T0_features_train_feedback.npz'
    npz_manifest = roots['run_cache'] / 'T0_features_train_feedback.json'
    t0_logits_cache = npz_path.is_file() and npz_manifest.is_file()
    if t0_logits_cache:
        saved = _read_json(npz_manifest)
        if (saved.get('run_fingerprint') != fingerprint
                or saved.get('file_sha256') != _sha256(npz_path)):
            raise RuntimeError('T0 feature cache provenance/hash mismatch')
        with np.load(resolve_local_file(npz_path), allow_pickle=False) as values:
            if (not np.array_equal(values['train_uids'], identity['uids']['train'])
                    or not np.array_equal(values['feedback_uids'], identity['uids']['feedback'])):
                raise RuntimeError('T0 feature cache UID mismatch')
    else:
        _atomic_npz(npz_path,
                    train_uids=identity['uids']['train'],
                    train_y=identity['labels']['train'],
                    train_features=t0_train, train_logits=t0_train_logits,
                    feedback_uids=identity['uids']['feedback'],
                    feedback_y=identity['labels']['feedback'],
                    feedback_features=t0_feedback, feedback_logits=t0_feedback_logits)
        _atomic_json(npz_manifest, {
            'run_fingerprint': fingerprint, 'file_sha256': _sha256(npz_path),
            'train_uid_sha256': identity['uid_sha256']['train'],
            'feedback_uid_sha256': identity['uid_sha256']['feedback'],
            'feature_tap': 'MIRepNet pooled feature before clshead, 256-D',
            'fit_partition': 'D_train only; feedback is diagnostics only',
            'pca_file_sha256': pca_record['manifest']['file_sha256'],
        })
    t0_diagnostics = _teacher_diagnostic_rows(
        'T0', models['teacher_model'], models['teacher_adapter'], data, identity, device, 0)
    _atomic_csv(roots['run_result'] / 'teacher_T0_source_diagnostics.csv', t0_diagnostics)
    if _stop_requested(stop_file):
        return {'status': 'paused', 'target_subject_id': fold + 1,
                'stage': 'after_T0_feature_audit'}

    epochs = 2 if smoke else None
    teachers = {}
    teacher_manifests = {}
    for name in TEACHER_NAMES:
        model, manifest = _adapt_teacher(
            name, config, identity, roots, models, data, fingerprint,
            device, stop_file, max_epochs=epochs)
        if model is None:
            return {'status': 'paused', 'target_subject_id': fold + 1,
                    'stage': f'teacher_{name}', **manifest}
        teachers[name] = model
        teacher_manifests[name] = manifest

    teacher_features, feature_manifest = _load_or_make_teacher_targets(
        config, identity, roots, models, data, pca, fingerprint,
        teachers, teacher_manifests, device)
    dimension = int(pca['dimension'])
    projection_np, projection_manifest = _make_fixed_q(
        dimension, fold, config, roots, fingerprint)
    prior_seed = int(config['randomness']['diffusion_seed_base']) + 31 * fold
    prior_init, prior_init_hash = _prior_init_state(dimension, prior_seed, config)
    prior_init_path = roots['run_weights'] / 'prior_initial_state.pt'
    if prior_init_path.is_file():
        prior_saved = torch.load(resolve_local_file(prior_init_path), map_location='cpu',
                                 weights_only=False)
        if (prior_saved.get('run_fingerprint') != fingerprint
                or prior_saved.get('model_sha256') != prior_init_hash
                or _state_hash(prior_saved.get('model', {})) != prior_init_hash):
            raise RuntimeError('saved common diffusion prior initialization differs')
    else:
        _atomic_torch(prior_init_path, {'run_fingerprint': fingerprint,
                                        'model_sha256': prior_init_hash,
                                        'model': prior_init})
    prior_updates = 4 if smoke else None
    priors, prior_manifests = {}, {}
    for name in TEACHER_NAMES:
        prior, manifest = _fit_denoiser(
            name, config, identity, roots, teacher_features, fingerprint,
            prior_init, prior_init_hash, device, stop_file,
            max_updates=prior_updates)
        if prior is None:
            return {'status': 'paused', 'target_subject_id': fold + 1,
                    'stage': f'prior_{name}', **manifest}
        priors[name] = prior
        prior_manifests[name] = manifest
    if smoke:
        _, alpha_audit, sigma_audit = core.ddpm_schedule(
            int(config['diffusion']['steps']), float(config['diffusion']['beta_start']),
            float(config['diffusion']['beta_end']), device)
        smoke_audits['ddim_gradient'] = _audit_ddim_gradient(
            priors['T_CE'], alpha_audit, sigma_audit, dimension, device)
        smoke_audits['zero_feature_weight_trajectory'] = _audit_zero_feature_trajectory(
            config, identity, models, data, initial, teacher_features,
            projection_np, priors, device)
    if _state_hash(priors['T_CE'].state_dict()) == _state_hash(priors['T_MCC'].state_dict()):
        raise RuntimeError('teacher-specific diffusion priors unexpectedly have identical terminal states')
    student_epochs = 2 if smoke else None
    student_manifests = {}
    for condition in METHODS:
        teacher_name = CONDITION_TEACHER[condition]
        prior = priors[teacher_name] if CONDITION_DIFFUSION[condition] else None
        prior_manifest = prior_manifests[teacher_name] if CONDITION_DIFFUSION[condition] else None
        status, manifest = _fit_student(
            condition, config, identity, roots, models, initial, data,
            teacher_features, feature_manifest, projection_np, projection_manifest,
            prior, prior_manifest, fingerprint, device, stop_file,
            max_epochs=student_epochs)
        if status is not None:
            return {'status': 'paused', 'target_subject_id': fold + 1,
                    'stage': f'student_{condition}', **status}
        student_manifests[condition] = manifest

    if smoke:
        smoke_record = {
            'status': 'passed', 'run_fingerprint': fingerprint,
            'target_subject_id': fold + 1,
            'teacher_adaptations': {key: value['epochs'] for key, value in teacher_manifests.items()},
            'prior_updates': {key: value['final_update'] for key, value in prior_manifests.items()},
            'student_epochs': {key: value['final_epoch'] for key, value in student_manifests.items()},
            'pca_dimension': dimension,
            'student_feature_gradient_reached_stem': {
                key: bool(_read_json(roots['run_result'] / f'student_{key}.json').get(
                    'status') == 'complete') for key in METHODS},
            'implementation_audits': smoke_audits,
            'target_evaluated': False,
        }
        _atomic_json(roots['run_result'] / 'smoke_result.json', smoke_record)
        return {'status': 'complete', 'target_subject_id': fold + 1,
                'run_fingerprint': fingerprint, 'smoke': True}

    completion = {
        'status': 'complete', 'experiment_id': config['experiment_id'],
        'run_fingerprint': fingerprint, 'seed': seed,
        'target_subject_id': fold + 1,
        'feedback_subject_id': int(identity['manifest']['feedback_subject_id']),
        'train_uid_sha256': identity['uid_sha256']['train'],
        'teacher_manifests': {key: value['checkpoint_sha256']
                              for key, value in teacher_manifests.items()},
        'prior_manifests': {key: value['checkpoint_sha256']
                            for key, value in prior_manifests.items()},
        'student_manifests': {key: value['checkpoint_sha256']
                              for key, value in student_manifests.items()},
        'pca_sha256': pca_record['manifest']['file_sha256'],
        'Q_sha256': projection_manifest['file_sha256'],
        'teacher_features_sha256': feature_manifest['file_sha256'],
        'completed_at_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
    }
    _atomic_json(complete_path, completion)
    _atomic_json(roots['run_result'] / 'resolved_run.json', {
        **_read_json(roots['run_result'] / 'resolved_run.json'),
        'status': 'complete', 'run_fingerprint': fingerprint,
        'fold_complete_sha256': _sha256(complete_path),
    })
    return {'status': 'complete', 'target_subject_id': fold + 1,
            'run_fingerprint': fingerprint}


def _strict_kd_raw_reference(config: dict, identity: dict, roots: dict,
                             models: dict, initial: dict, fold: int) -> tuple[dict | None, str]:
    run_id = 'student_aware_diffusion_prior_loso_v1_seed666'
    old = reference._run_roots(run_id, int(config['seed']), fold)
    manifest_path = old['run_result'] / 'student_B.json'
    checkpoint_path = old['run_weights'] / 'student_B.pt'
    if not manifest_path.is_file() or not checkpoint_path.is_file():
        return None, 'raw T0 logits-KD checkpoint/manifest is missing'
    manifest = _read_json(manifest_path)
    expected_loss = {'ce_weight': 1.0, 'kd_start_epoch': 1,
                     'kd_weight': 0.5, 'temperature': 2.0}
    if (manifest.get('status') != 'complete' or manifest.get('method') != 'KD_RAW'
            or manifest.get('method_code') != 'B' or int(manifest.get('final_epoch', -1)) != 100
            or manifest.get('initial_model_sha256') != initial['model_sha256']
            or manifest.get('train_uid_sha256') != identity['uid_sha256']['train']
            or manifest.get('student_schedule_sha256') != models['student_schedule_sha256']
            or manifest.get('student_config') != models['student_cfg']
            or manifest.get('loss') != expected_loss
            or manifest.get('teacher_target', {}).get('method') != 'T0'
            or manifest.get('teacher_target', {}).get('checkpoint_sha256')
            != models['teacher_checkpoint_sha256']
            or manifest.get('checkpoint_sha256') != _sha256(checkpoint_path)):
        return None, 'raw T0 logits-KD artifacts do not match this run protocol exactly'
    return {'manifest': manifest, 'checkpoint': str(checkpoint_path),
            'checkpoint_sha256': _sha256(checkpoint_path)}, 'strictly matched historical reference'


def _audit_mcc_formula(device: torch.device) -> dict:
    cuda_devices = [device] if device.type == 'cuda' else []
    with torch.random.fork_rng(devices=cuda_devices):
        torch.manual_seed(20261030)
        logits = torch.randn(13, 2, device=device, requires_grad=True)
        value, stats = core.minimum_class_confusion(logits, temperature=2.0)
        probabilities = torch.softmax(logits / 2.0, dim=1)
        entropy = -(probabilities * torch.log(probabilities + 1e-5)).sum(dim=1)
        weights = len(logits) * (1.0 + torch.exp(-entropy.detach()))
        weights = weights / weights.sum()
        confusion = probabilities.transpose(0, 1) @ (weights[:, None] * probabilities)
        manual = confusion / confusion.sum(dim=1, keepdim=True).clamp_min(1e-8)
        expected = (manual.sum() - manual.diagonal().sum()) / 2
        if not torch.allclose(value, expected, rtol=1e-7, atol=1e-7):
            raise RuntimeError('MCC implementation differs from the preregistered equation')
        value.backward()
        if logits.grad is None or not torch.isfinite(logits.grad).all():
            raise RuntimeError('MCC gradient audit failed')
        if int(stats['clamped_rows']) != 0:
            raise RuntimeError('MCC synthetic audit unexpectedly hit a row clamp')
        return {'scalar_abs_error': float(torch.abs(value - expected).detach().cpu()),
                'finite_gradient_norm': float(logits.grad.norm().detach().cpu()),
                'clamped_rows': int(stats['clamped_rows'])}


def _audit_zero_mcc_trajectory(config: dict, identity: dict, models: dict,
                               data: dict, device: torch.device) -> dict:
    teacher_cfg = config['teacher_adaptation']
    seed = int(config['randomness']['teacher_schedule_seed_base']) + 1009 * fold_from_identity(identity)
    schedule = reference._epoch_schedule(len(identity['uids']['train']),
                                         int(teacher_cfg['batch_size']), 2, seed)
    final_hashes = []
    mcc_skipped = []
    for mode, weight in (('CE_ONLY', 0.0), ('MCC_CONFIGURED_AT_ZERO', 0.0)):
        model = copy.deepcopy(models['teacher_model']).to(device)
        model.eval()
        optimizer = torch.optim.Adam(
            model.parameters(), lr=float(teacher_cfg['lr']),
            betas=tuple(teacher_cfg['betas']), eps=float(teacher_cfg['eps']),
            weight_decay=float(teacher_cfg['weight_decay']))
        for epoch in range(2):
            for indices in schedule[epoch]:
                xb = torch.as_tensor(data['teacher_x']['train'][indices],
                                     dtype=torch.float32, device=device)
                yb = torch.as_tensor(identity['labels']['train'][indices],
                                     dtype=torch.long, device=device)
                optimizer.zero_grad(set_to_none=True)
                _, logits = models['teacher_adapter'].forward(model, xb)
                loss, _, _, stats = _teacher_objective(
                    logits, yb, weight, teacher_cfg['mcc'])
                mcc_skipped.append(stats is None)
                loss.backward()
                torch.nn.utils.clip_grad_norm_(model.parameters(),
                                               float(teacher_cfg['gradient_clip_norm']))
                optimizer.step()
        final_hashes.append(_state_hash(model.state_dict()))
    if final_hashes[0] != final_hashes[1]:
        raise RuntimeError('zero-weight MCC did not reproduce the CE teacher trajectory')
    if not all(mcc_skipped):
        raise RuntimeError('zero-weight MCC unexpectedly evaluated the MCC graph')
    return {'CE_only_hash': final_hashes[0], 'CE_plus_zero_MCC_hash': final_hashes[1],
            'exact_match': True, 'mcc_skipped': True, 'epochs': 2}


def _audit_zero_feature_trajectory(config: dict, identity: dict, models: dict,
                                   data: dict, initial: dict, teacher_features: dict,
                                   projection_np: np.ndarray,
                                   priors: dict, device: torch.device) -> dict:
    projection = torch.as_tensor(projection_np.astype(np.float32), device=device)
    hashes = {}
    for condition in METHODS:
        teacher_name = CONDITION_TEACHER[condition]
        use_diffusion = CONDITION_DIFFUSION[condition]
        model, optimizer, scheduler, _ = reference._new_paired_student(
            models, initial, int(config['randomness']['student_initialization_seed']), device)
        generator = None
        if use_diffusion:
            generator = torch.Generator(device=device)
            generator.manual_seed(
                int(config['randomness']['student_feature_noise_seed_base'])
                + 1009 * fold_from_identity(identity))
            noise_before = generator.get_state().clone()
        targets = torch.as_tensor(
            teacher_features[f'{teacher_name}_train_features_whitened'],
            dtype=torch.float32, device=device)
        for epoch in range(2):
            model.train()
            for indices in models['student_schedule'][epoch]:
                xb = torch.as_tensor(data['student_x']['train'][indices],
                                     dtype=torch.float32, device=device)
                yb = torch.as_tensor(identity['labels']['train'][indices],
                                     dtype=torch.long, device=device)
                batch_ids = torch.as_tensor(indices, dtype=torch.long, device=device)
                optimizer.zero_grad(set_to_none=True)
                native, logits = models['student_adapter'].forward(model, xb)
                auxiliary = core.student_features_to_vector(native, projection)
                feature_loss = _student_feature_loss(
                    auxiliary, targets[batch_ids], 0.0, use_diffusion,
                    priors[teacher_name] if use_diffusion else None,
                    None, None, generator)
                if feature_loss is not None:
                    raise RuntimeError('zero feature weight unexpectedly produced a loss')
                ce = F.cross_entropy(logits, yb)
                ce.backward()
                optimizer.step()
            scheduler.step()
        hashes[condition] = _state_hash(model.state_dict())
        if use_diffusion and not torch.equal(generator.get_state(), noise_before):
            raise RuntimeError(f'{condition}: zero feature weight consumed the noise RNG')
    if len(set(hashes.values())) != 1:
        raise RuntimeError('zero feature weight did not collapse all four arms to the same CE trajectory')
    return {'condition_hashes': hashes, 'exact_match': True,
            'denoiser_skipped': True, 'noise_rng_unchanged': True}


def _audit_ddim_gradient(prior: core.EpsilonPredictor, alpha: torch.Tensor,
                         sigma: torch.Tensor, dimension: int,
                         device: torch.device) -> dict:
    calls = {'count': 0}
    hook = prior.network.register_forward_hook(
        lambda _module, _inputs, _output: calls.__setitem__('count', calls['count'] + 1))
    try:
        cuda_devices = [device] if device.type == 'cuda' else []
        with torch.random.fork_rng(devices=cuda_devices):
            torch.manual_seed(554433)
            features = torch.randn(3, dimension, device=device, requires_grad=True)
            noise = torch.randn(3, dimension, device=device)
            output = core.ddim_denoise(features, noise, prior, alpha, sigma)
            output.square().mean().backward()
    finally:
        hook.remove()
    if calls['count'] != 5 or features.grad is None or not torch.isfinite(features.grad).all():
        raise RuntimeError('five-step DDIM call-count/input-gradient audit failed')
    if float(features.grad.norm().detach().cpu()) <= 0:
        raise RuntimeError('DDIM gradient to the student feature is zero')
    if any(parameter.grad is not None for parameter in prior.parameters()):
        raise RuntimeError('frozen DDIM prior received gradients')
    return {'epsilon_calls': calls['count'],
            'input_gradient_norm': float(features.grad.norm().detach().cpu()),
            'prior_parameters_have_no_grad': True}


def _load_new_checkpoint(adapter, checkpoint_path: Path,
                         manifest: dict, device: torch.device) -> torch.nn.Module:
    rng = _capture_rng(device)
    try:
        model = adapter.build(2).to(device)
    finally:
        _restore_rng(rng, device)
    _load_model_state(model, checkpoint_path, manifest['method_fingerprint'])
    return model


def _evaluate_fold(config_path: Path, config: dict, run_id: str, fold: int,
                   device: torch.device) -> tuple[list[dict], dict]:
    seed = int(config['seed'])
    roots = _run_roots(run_id, seed, fold)
    resolved, cfg = reference._load_resolved_reference(config)
    identity = reference._load_fold_identity(
        config, fold, roots, cfg['resolved_config_sha256'], include_test_labels=True)
    models = reference._load_reference_checkpoints(
        config, resolved, cfg, identity, roots, device)
    initial = _load_initial_student(config, identity, roots, models)
    data = reference._load_fold_data(identity, roots, resolved['student_config'], device)
    completion_path = roots['run_result'] / 'fold_complete.json'
    completion = _read_json(completion_path)
    if completion.get('status') != 'complete':
        raise RuntimeError(f'target evaluation blocked: fold {fold+1} has no completion record')
    if completion.get('run_fingerprint') != _fold_fingerprint(
            config_path, config, cfg, identity, roots, models, initial, run_id):
        # The path is carried by the caller in config; all training inputs are
        # separately verified above, so this additionally rejects stale folds.
        raise RuntimeError(f'target evaluation blocked: fold {fold+1} identity changed')
    rows = {}
    predictions = {'uids': identity['uids']['test'], 'y': identity['labels']['test']}
    teacher_x = np.asarray(data['teacher_input_mmap'][data['teacher_rows']['test']], dtype=np.float32)
    raw_student = np.asarray(data['student_raw_mmap'][data['student_rows']['test']], dtype=np.float32)
    student_x = reference._preprocess_student(data['student_adapter'], raw_student)
    models_to_eval = {
        'TEACHER_T0': (models['teacher_adapter'], models['teacher_model'], 'teacher'),
        'IFNET_CE_REFERENCE': (models['student_adapter'], models['student_model'], 'student'),
    }
    adapted_teachers = {}
    for teacher_name in TEACHER_NAMES:
        manifest = _read_json(roots['run_result'] / f'teacher_{teacher_name}.json')
        if (manifest.get('status') != 'complete'
                or manifest.get('checkpoint_sha256') != completion['teacher_manifests'].get(teacher_name)):
            raise RuntimeError(f'{teacher_name}: final checkpoint is incomplete or changed')
        model = _load_new_checkpoint(
            models['teacher_adapter'], roots['run_weights'] / f'teacher_{teacher_name}.pt',
            manifest, device)
        adapted_teachers[teacher_name] = (model, manifest)
        models_to_eval[f'TEACHER_{teacher_name}'] = (
            models['teacher_adapter'], model, 'teacher')
    kdraw, kdraw_status = _strict_kd_raw_reference(config, identity, roots, models, initial, fold)
    if kdraw is not None:
        kdraw_adapter = models['student_adapter']
        kdraw_model = _load_new_checkpoint(
            kdraw_adapter, Path(kdraw['checkpoint']), kdraw['manifest'], device)
        models_to_eval['IFNET_KD_RAW_REFERENCE'] = (kdraw_adapter, kdraw_model, 'student')
    for condition in METHODS:
        manifest = _read_json(roots['run_result'] / f'student_{condition}.json')
        if (manifest.get('status') != 'complete'
                or manifest.get('checkpoint_sha256') != completion['student_manifests'].get(condition)):
            raise RuntimeError(f'{condition}: final IFNet checkpoint is incomplete or changed')
        student_model = _load_new_checkpoint(
            models['student_adapter'], roots['run_weights'] / f'student_{condition}.pt',
            manifest, device)
        models_to_eval[condition] = (models['student_adapter'], student_model, 'student')

    subject_rows = []
    for name, (adapter, model, role) in models_to_eval.items():
        values = teacher_x if role == 'teacher' else student_x
        _, logits = reference._infer(adapter, model, values, device,
                                     batch_size=32 if role == 'teacher' else 64,
                                     return_features=False)
        metrics = reference._metric_values(identity['labels']['test'], logits)
        predictions[f'{name}_logits'] = logits
        code = name if name in METHODS else name
        row = {
            'condition': code, 'model_role': role,
            'model': 'MIRepNet' if role == 'teacher' else 'IFNet',
            'seed': seed, 'target_subject_id': fold + 1,
            'n_trials': metrics['n'],
            'balanced_accuracy': metrics['balanced_accuracy'],
            'accuracy': metrics['accuracy'], 'kappa': metrics['kappa'],
            'cross_entropy': metrics['cross_entropy'],
            'prediction_count_0': metrics['prediction_counts'][0],
            'prediction_count_1': metrics['prediction_counts'][1],
            'collapse': metrics['collapse'],
            'test_uid_sha256': identity['uid_sha256']['test'],
        }
        subject_rows.append(row)
        rows[name] = metrics
    prediction_path = roots['run_result'] / 'test_predictions.npz'
    _atomic_npz(prediction_path, **predictions)
    _atomic_json(roots['run_result'] / 'test_evaluation.json', {
        'status': 'complete', 'seed': seed, 'target_subject_id': fold + 1,
        'run_fingerprint': completion['run_fingerprint'],
        'test_uid_sha256': identity['uid_sha256']['test'],
        'prediction_file': str(prediction_path),
        'prediction_file_sha256': _sha256(prediction_path),
        'strict_kd_raw_reference': kdraw_status,
        'prediction_models': list(rows),
    })
    return subject_rows, {'fold': fold + 1, 'metrics': rows,
                          'strict_kd_raw_reference': kdraw_status,
                          'kd_raw_reference_available': kdraw is not None}


def _load_saved_eval(config_path: Path, config: dict, run_id: str,
                     fold: int) -> tuple[list[dict], dict] | None:
    """Read verified predictions after model weights have been pruned."""
    roots = _run_roots(run_id, int(config['seed']), fold)
    prediction_path = roots['run_result'] / 'test_predictions.npz'
    evaluation_path = roots['run_result'] / 'test_evaluation.json'
    completion_path = roots['run_result'] / 'fold_complete.json'
    if not (prediction_path.is_file() and evaluation_path.is_file()
            and completion_path.is_file()):
        return None
    evaluation = _read_json(evaluation_path)
    completion = _read_json(completion_path)
    resolved_run_path = roots['run_result'] / 'resolved_run.json'
    if not resolved_run_path.is_file():
        raise RuntimeError(f'fold {fold+1}: resolved run provenance is missing')
    resolved_run = _read_json(resolved_run_path)
    if (evaluation.get('status') != 'complete'
            or evaluation.get('prediction_file_sha256') != _sha256(prediction_path)
            or evaluation.get('run_fingerprint') != completion.get('run_fingerprint')
            or resolved_run.get('status') != 'complete'
            or resolved_run.get('run_fingerprint') != completion.get('run_fingerprint')
            or resolved_run.get('experiment_id') != config['experiment_id']
            or int(resolved_run.get('seed', -1)) != int(config['seed'])
            or int(resolved_run.get('target_subject_id', -1)) != fold + 1
            or resolved_run.get('config_sha256') != _sha256(config_path)
            or resolved_run.get('fold_complete_sha256') != _sha256(completion_path)):
        raise RuntimeError(f'fold {fold+1}: saved target predictions fail provenance/hash checks')
    resolved, cfg = reference._load_resolved_reference(config)
    identity = reference._load_fold_identity(
        config, fold, roots, cfg['resolved_config_sha256'], include_test_labels=True)
    if (evaluation.get('test_uid_sha256') != identity['uid_sha256']['test']
            or completion.get('train_uid_sha256') != identity['uid_sha256']['train']
            or resolved_run.get('train_uid_sha256') != identity['uid_sha256']['train']
            or resolved_run.get('feedback_uid_sha256') != identity['uid_sha256']['feedback']
            or resolved_run.get('test_uid_sha256') != identity['uid_sha256']['test']
            or resolved_run.get('split_manifest_sha256') != identity['sha256']
            or resolved_run.get('source_manifest_sha256') != identity['source_manifest_sha256']
            or resolved_run.get('source_trial_manifest_sha256')
            != identity['source_trial_manifest_sha256']):
        raise RuntimeError(f'fold {fold+1}: saved target prediction UIDs differ from the manifest')
    with np.load(resolve_local_file(prediction_path), allow_pickle=False) as source:
        if (not np.array_equal(source['uids'], identity['uids']['test'])
                or not np.array_equal(source['y'], identity['labels']['test'])):
            raise RuntimeError(f'fold {fold+1}: saved prediction labels/UIDs mismatch')
        arrays = {key: np.asarray(source[key]) for key in source.files}
    model_names = evaluation.get('prediction_models', [])
    required = ['TEACHER_T0', 'TEACHER_T_CE', 'TEACHER_T_MCC',
                'IFNET_CE_REFERENCE', *METHODS]
    if not set(required).issubset(model_names):
        raise RuntimeError(f'fold {fold+1}: saved predictions are incomplete')
    subject_rows, by_name = [], {}
    for name in model_names:
        key = f'{name}_logits'
        if key not in arrays or arrays[key].shape != (len(identity['labels']['test']), 2):
            raise RuntimeError(f'fold {fold+1}: missing or malformed saved logits for {name}')
        role = 'teacher' if name.startswith('TEACHER_') else 'student'
        metrics = reference._metric_values(identity['labels']['test'], arrays[key])
        subject_rows.append({
            'condition': name, 'model_role': role,
            'model': 'MIRepNet' if role == 'teacher' else 'IFNet',
            'seed': int(config['seed']), 'target_subject_id': fold + 1,
            'n_trials': metrics['n'], 'balanced_accuracy': metrics['balanced_accuracy'],
            'accuracy': metrics['accuracy'], 'kappa': metrics['kappa'],
            'cross_entropy': metrics['cross_entropy'],
            'prediction_count_0': metrics['prediction_counts'][0],
            'prediction_count_1': metrics['prediction_counts'][1],
            'collapse': metrics['collapse'], 'test_uid_sha256': identity['uid_sha256']['test'],
        })
        by_name[name] = metrics
    status = evaluation.get('strict_kd_raw_reference', 'not recorded')
    return subject_rows, {
        'fold': fold + 1, 'metrics': by_name,
        'strict_kd_raw_reference': status,
        'kd_raw_reference_available': 'IFNET_KD_RAW_REFERENCE' in model_names,
        'loaded_from_saved_predictions': True,
    }


def _prune_run_weights(run_id: str, seed: int) -> dict:
    """Drop this pilot's temporary trained weights after reports/predictions verify."""
    report_dir = REPORTS_ROOT / 'distill' / run_id
    required_reports = [report_dir / 'summary.json', report_dir / 'summary.csv',
                        report_dir / 'summary.xlsx', report_dir / 'subject_metrics.csv',
                        report_dir / 'paired_comparisons.csv']
    if not all(path.is_file() and path.stat().st_size > 0 for path in required_reports):
        raise RuntimeError('refusing to prune model weights before all user reports are complete')
    weight_root = WEIGHTS_ROOT / run_id / f'seed_{seed}'
    if not weight_root.exists():
        return {'status': 'already_absent', 'path': str(weight_root)}
    external_results = RESULTS_ROOT / 'distill' / run_id / f'seed_{seed}'
    fold_dirs = sorted((external_results / 'folds').glob('target_*'))
    if len(fold_dirs) != 9:
        raise RuntimeError(f'refusing to prune weights; expected nine completed folds, found {len(fold_dirs)}')
    for fold_dir in fold_dirs:
        evaluation_path = fold_dir / 'test_evaluation.json'
        prediction_path = fold_dir / 'test_predictions.npz'
        completion_path = fold_dir / 'fold_complete.json'
        resolved_run_path = fold_dir / 'resolved_run.json'
        if not (evaluation_path.is_file() and prediction_path.is_file()
                and completion_path.is_file() and resolved_run_path.is_file()):
            raise RuntimeError(f'refusing to prune weights; incomplete fold {fold_dir.name}')
        evaluation = _read_json(evaluation_path)
        completion = _read_json(completion_path)
        resolved_run = _read_json(resolved_run_path)
        if (evaluation.get('prediction_file_sha256') != _sha256(prediction_path)
                or evaluation.get('run_fingerprint') != completion.get('run_fingerprint')
                or resolved_run.get('status') != 'complete'
                or resolved_run.get('fold_complete_sha256') != _sha256(completion_path)):
            raise RuntimeError(f'refusing to prune weights; invalid prediction provenance in {fold_dir}')
    shutil.rmtree(weight_root)
    try:
        (WEIGHTS_ROOT / run_id).rmdir()
    except OSError:
        pass
    return {'status': 'pruned_after_verified_evaluation', 'path': str(weight_root)}


def _aggregate_rows(subject_rows: list[dict], condition: str) -> dict:
    selected = [row for row in subject_rows if row['condition'] == condition]
    if len(selected) != 9 or len({int(row['target_subject_id']) for row in selected}) != 9:
        raise RuntimeError(f'{condition}: expected exactly one result for each of nine subjects')
    out = {'condition': condition, 'model_role': selected[0]['model_role'],
           'model': selected[0]['model'], 'seed': 666, 'n_seeds': 1, 'n_subjects': 9}
    for metric in ('balanced_accuracy', 'accuracy', 'kappa', 'cross_entropy'):
        values = np.asarray([float(row[metric]) for row in selected], dtype=np.float64)
        scale = 100.0 if metric in ('balanced_accuracy', 'accuracy') else 1.0
        out[f'{metric}_mean'] = float(values.mean() * scale)
        out[f'{metric}_std'] = float(values.std(ddof=1) * scale)
        if scale == 100.0:
            out[f'{metric}_mean_pm_std'] = (
                f"{out[f'{metric}_mean']:.2f} ± {out[f'{metric}_std']:.2f}")
    return out


def _paired_row(label: str, left: str, right: str,
                rows_by_condition: dict, draws: np.ndarray) -> dict:
    left_rows = sorted(rows_by_condition[left], key=lambda row: int(row['target_subject_id']))
    right_rows = sorted(rows_by_condition[right], key=lambda row: int(row['target_subject_id']))
    if ([row['target_subject_id'] for row in left_rows]
            != [row['target_subject_id'] for row in right_rows]):
        raise RuntimeError(f'{label}: paired subject IDs do not match')
    left_values = np.asarray([row['balanced_accuracy'] for row in left_rows]) * 100.0
    right_values = np.asarray([row['balanced_accuracy'] for row in right_rows]) * 100.0
    differences = left_values - right_values
    bootstrap_means = differences[draws].mean(axis=1)
    return {
        'comparison': label, 'method': left, 'baseline': right,
        'seed': 666, 'n_seeds': 1, 'n_subjects': len(differences),
        'mean_balanced_accuracy_difference_pp': float(differences.mean()),
        'std_difference_pp': float(differences.std(ddof=1)),
        'paired_bootstrap_95ci_low_pp': float(np.percentile(bootstrap_means, 2.5)),
        'paired_bootstrap_95ci_high_pp': float(np.percentile(bootstrap_means, 97.5)),
        'bootstrap_draws': len(draws), 'bootstrap_seed': 20261030,
        'wins': int(np.sum(differences > 1e-6)),
        'ties': int(np.sum(np.abs(differences) <= 1e-6)),
        'losses': int(np.sum(differences < -1e-6)),
        'per_subject_differences_pp': json.dumps(differences.tolist()),
    }


def _comparison_fill(value: float, teacher_value: float, student_value: float) -> str | None:
    greater_teacher = value > teacher_value
    greater_student = value > student_value
    if greater_teacher and greater_student:
        return 'FFC7CE'
    if greater_teacher:
        return 'C6EFCE'
    if greater_student:
        return 'FFF2CC'
    return None


def _write_workbook(path: Path, summary_rows: list[dict],
                    paired_rows: list[dict], teacher_rows: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Summary'
    sheet.sheet_view.showGridLines = False
    headers = ['student', 'method', 'n_seed', 'n_subject',
               'Balanced Accuracy mean (%)', 'Balanced Accuracy SD (pp)',
               'Accuracy mean (%)', 'Accuracy SD (pp)', 'Kappa mean', 'Cross Entropy mean']
    sheet.append(headers)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ordered = [
        ('TEACHER_T0', 'MIRepNet', 'T0 reference'),
        ('TEACHER_T_CE', 'MIRepNet', 'CE continue 10 ep'),
        ('TEACHER_T_MCC', 'MIRepNet', 'CE + MCC 10 ep'),
        ('IFNET_CE_REFERENCE', 'IFNet', 'none (supervised CE)'),
        ('IFNET_KD_RAW_REFERENCE', 'IFNet', 'logits KD (strict historical reference)'),
        ('CE_FEATURE', 'IFNet', 'CE_FEATURE'),
        ('MCC_FEATURE', 'IFNet', 'MCC_FEATURE'),
        ('CE_DIFFKD', 'IFNet', 'CE_DIFFKD'),
        ('MCC_DIFFKD', 'IFNet', 'MCC_DIFFKD'),
    ]
    by_name = {row['condition']: row for row in summary_rows}
    ordered = [item for item in ordered if item[0] in by_name]
    start_row = 2
    row_index = {}
    for condition, student, method in ordered:
        row = by_name[condition]
        sheet.append([
            student, method, 1, 9,
            row['balanced_accuracy_mean'], row['balanced_accuracy_std'],
            row['accuracy_mean'], row['accuracy_std'],
            row['kappa_mean'], row['cross_entropy_mean'],
        ])
        row_index[condition] = sheet.max_row
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for col in (5, 6, 7, 8):
            sheet.cell(sheet.max_row, col).number_format = '0.00'
        for col in (9, 10):
            sheet.cell(sheet.max_row, col).number_format = '0.0000'
    if 'TEACHER_T0' in row_index:
        for cell in sheet[row_index['TEACHER_T0']]:
            cell.fill = PatternFill('solid', fgColor='FAC090')
    if 'IFNET_CE_REFERENCE' in row_index:
        for cell in sheet[row_index['IFNET_CE_REFERENCE']]:
            cell.fill = PatternFill('solid', fgColor='93CDDD')
    teacher_ref = by_name.get('TEACHER_T0', {}).get('balanced_accuracy_mean')
    none_ref = by_name.get('IFNET_CE_REFERENCE', {}).get('balanced_accuracy_mean')
    teacher_acc = by_name.get('TEACHER_T0', {}).get('accuracy_mean')
    none_acc = by_name.get('IFNET_CE_REFERENCE', {}).get('accuracy_mean')
    if teacher_ref is not None and none_ref is not None:
        for condition in METHODS:
            if condition not in row_index:
                continue
            row = by_name[condition]
            ba_color = _comparison_fill(row['balanced_accuracy_mean'], teacher_ref, none_ref)
            acc_color = _comparison_fill(row['accuracy_mean'], teacher_acc, none_acc)
            if ba_color:
                sheet.cell(row_index[condition], 5).fill = PatternFill('solid', fgColor=ba_color)
            if acc_color:
                sheet.cell(row_index[condition], 7).fill = PatternFill('solid', fgColor=acc_color)
    legend_row = sheet.max_row + 2
    sheet.cell(legend_row, 1, 'Legend')
    legend = [
        ('大于小模型', 'FFF2CC'), ('大于大模型', 'C6EFCE'),
        ('大于二者', 'FFC7CE'), ('教师参考', 'FAC090'), ('IFNet CE 基线', '93CDDD'),
    ]
    for offset, (label, color) in enumerate(legend, start=1):
        cell = sheet.cell(legend_row + offset, 1, label)
        cell.fill = PatternFill('solid', fgColor=color)
        cell.alignment = Alignment(horizontal='left')
    note_row = legend_row + len(legend) + 2
    notes = [
        '任务：BNCI2014004 session_3；指标分别报告 Balanced Accuracy 与 Accuracy。',
        '汇总：seed 666；九个 LOSO 目标被试等权平均，SD 为被试间样本标准差，不是跨 seed SD。',
        '四个新条件：从 T0 做 CE 或 CE+MCC 教师适配，再训练新 IFNet；FEATURE 是直接配对特征 MSE，DIFFKD 是五步 DDIM 后配对特征 MSE。',
        '颜色按对应指标与 MIRepNet T0、同一 IFNet 的 none CE 基线做严格大于比较；相等不计为大于。',
        'KD logits 仅在全部九折初始化、数据顺序、配置、epoch 与 T0 provenance 均严格匹配时列为历史参考。',
    ]
    for offset, note in enumerate(notes):
        row = note_row + offset
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cell = sheet.cell(row, 1, note)
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    widths = [16, 36, 10, 12, 25, 24, 18, 18, 14, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'A2'

    paired = workbook.create_sheet('Paired comparisons')
    paired.sheet_view.showGridLines = False
    if paired_rows:
        paired.append(list(paired_rows[0].keys()))
        for cell in paired[1]:
            cell.fill = header_fill
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for record in paired_rows:
            paired.append([record[key] for key in paired_rows[0]])
        for row in paired.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        paired.freeze_panes = 'A2'
        for column in paired.columns:
            letter = get_column_letter(column[0].column)
            paired.column_dimensions[letter].width = min(
                42, max(14, max(len(str(cell.value or '')) for cell in column) + 2))

    teacher = workbook.create_sheet('Teacher results')
    teacher.sheet_view.showGridLines = False
    if teacher_rows:
        teacher.append(list(teacher_rows[0].keys()))
        for cell in teacher[1]:
            cell.fill = header_fill
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for record in teacher_rows:
            teacher.append([record[key] for key in teacher_rows[0]])
        teacher.freeze_panes = 'A2'
        for column in teacher.columns:
            letter = get_column_letter(column[0].column)
            teacher.column_dimensions[letter].width = min(
                40, max(14, max(len(str(cell.value or '')) for cell in column) + 2))
    path = require_report_output(path)
    workbook.save(path)


def _write_reports(run_id: str, config: dict, subject_rows: list[dict],
                   fold_records: list[dict]) -> dict:
    report_dir = require_report_output(REPORTS_ROOT / 'distill' / run_id / 'summary.csv').parent
    report_dir.mkdir(parents=True, exist_ok=True)
    all_kd_raw = all(row.get('kd_raw_reference_available', False) for row in fold_records)
    if not all_kd_raw:
        subject_rows = [row for row in subject_rows
                        if row['condition'] != 'IFNET_KD_RAW_REFERENCE']
    conditions = ['TEACHER_T0', 'TEACHER_T_CE', 'TEACHER_T_MCC', 'IFNET_CE_REFERENCE']
    if all_kd_raw:
        conditions.append('IFNET_KD_RAW_REFERENCE')
    conditions.extend(METHODS)
    summary_rows = [_aggregate_rows(subject_rows, name) for name in conditions]
    by_condition = {name: [row for row in subject_rows if row['condition'] == name]
                    for name in conditions}
    rng = np.random.default_rng(int(config['report']['bootstrap_seed']))
    draws = rng.integers(0, 9, size=(int(config['report']['bootstrap_draws']), 9))
    comparisons = [
        ('PRIMARY: MCC_DIFFKD - MCC_FEATURE', 'MCC_DIFFKD', 'MCC_FEATURE'),
        ('MCC_DIFFKD - CE_DIFFKD', 'MCC_DIFFKD', 'CE_DIFFKD'),
        ('MCC_FEATURE - CE_FEATURE', 'MCC_FEATURE', 'CE_FEATURE'),
        ('CE_DIFFKD - CE_FEATURE', 'CE_DIFFKD', 'CE_FEATURE'),
        ('MCC_DIFFKD - CE_FEATURE', 'MCC_DIFFKD', 'CE_FEATURE'),
        ('MCC_DIFFKD - IFNET_CE_REFERENCE', 'MCC_DIFFKD', 'IFNET_CE_REFERENCE'),
    ]
    if all_kd_raw:
        comparisons.append(('MCC_DIFFKD - IFNET_KD_RAW_REFERENCE',
                            'MCC_DIFFKD', 'IFNET_KD_RAW_REFERENCE'))
    paired_rows = [_paired_row(label, left, right, by_condition, draws)
                   for label, left, right in comparisons]
    teacher_comparisons = []
    for left, right in (('TEACHER_T_MCC', 'TEACHER_T_CE'),
                        ('TEACHER_T_CE', 'TEACHER_T0'),
                        ('TEACHER_T_MCC', 'TEACHER_T0')):
        left_rows = sorted(by_condition[left], key=lambda row: row['target_subject_id'])
        right_rows = sorted(by_condition[right], key=lambda row: row['target_subject_id'])
        difference = (np.asarray([row['balanced_accuracy'] for row in left_rows])
                      - np.asarray([row['balanced_accuracy'] for row in right_rows])) * 100.0
        teacher_comparisons.append({
            'comparison': f'{left} - {right}', 'metric': 'Balanced Accuracy (pp)',
            'seed': 666, 'n_subjects': 9, 'mean_difference_pp': float(difference.mean()),
            'std_difference_pp': float(difference.std(ddof=1)),
            'wins': int(np.sum(difference > 1e-6)),
            'ties': int(np.sum(np.abs(difference) <= 1e-6)),
            'losses': int(np.sum(difference < -1e-6)),
            'per_subject_differences_pp': json.dumps(difference.tolist()),
        })

    reference_status = [
        {'target_subject_id': row['fold'], 'available': row['kd_raw_reference_available'],
         'status': row['strict_kd_raw_reference']} for row in fold_records]
    _atomic_csv(report_dir / 'summary.csv', summary_rows)
    _atomic_csv(report_dir / 'subject_metrics.csv', subject_rows)
    _atomic_csv(report_dir / 'paired_comparisons.csv', paired_rows)
    _atomic_csv(report_dir / 'teacher_comparisons.csv', teacher_comparisons)
    reference._atomic_json(report_dir / 'reference_artifact_audit.json', {
        'strict_KD_RAW_reference': reference_status,
        'included_in_summary': all_kd_raw,
        'method': 'previous T0 logits KD; no retraining',
    })
    _write_workbook(report_dir / 'summary.xlsx', summary_rows, paired_rows,
                    [row for row in summary_rows if row['model_role'] == 'teacher'])
    primary = paired_rows[0]
    mcc_teacher = teacher_comparisons[0]
    conclusions = [
        (f"主要比较 MCC_DIFFKD−MCC_FEATURE：ΔBA={primary['mean_balanced_accuracy_difference_pp']:+.2f} pp，"
         f"配对被试 bootstrap 95% CI [{primary['paired_bootstrap_95ci_low_pp']:+.2f}, "
         f"{primary['paired_bootstrap_95ci_high_pp']:+.2f}]，W/T/L="
         f"{primary['wins']}/{primary['ties']}/{primary['losses']}。"),
        (f"教师 MCC 适配相对 CE 适配：ΔBA={mcc_teacher['mean_difference_pp']:+.2f} pp，"
         f"W/T/L={mcc_teacher['wins']}/{mcc_teacher['ties']}/{mcc_teacher['losses']}。"),
        ('结果是 seed 666 的 9 被试探索性 pilot；LOSO 训练折重叠，区间不代表跨 seed 或跨数据集稳定性。'),
    ]
    summary = {
        'experiment_id': config['experiment_id'], 'run_id': run_id,
        'dataset': config['dataset'], 'session': config['selected_session'],
        'source_variant': config['source_variant'], 'seed': 666, 'n_seeds': 1,
        'subjects': list(range(1, 10)), 'primary_metric': 'Balanced Accuracy (%)',
        'aggregation': 'equal-weight mean and sample SD over nine LOSO target subjects; seed 666 only',
        'bootstrap_unit': 'target subject', 'bootstrap_draws': int(config['report']['bootstrap_draws']),
        'bootstrap_seed': int(config['report']['bootstrap_seed']),
        'primary_comparison': primary, 'student_summary': summary_rows,
        'student_paired_comparisons': paired_rows,
        'teacher_paired_comparisons': teacher_comparisons,
        'strict_KD_RAW_reference_included': all_kd_raw,
        'strict_KD_RAW_reference_audit': reference_status,
        'summary_csv': str(report_dir / 'summary.csv'),
        'subject_metrics_csv': str(report_dir / 'subject_metrics.csv'),
        'paired_comparisons_csv': str(report_dir / 'paired_comparisons.csv'),
        'teacher_comparisons_csv': str(report_dir / 'teacher_comparisons.csv'),
        'workbook': str(report_dir / 'summary.xlsx'), 'conclusions': conclusions,
    }
    reference._atomic_json(report_dir / 'summary.json', summary)
    md = [
        '# MCC_TEACHER_ADAPT_DIFFKD_LOSO_V1', '',
        '数据：BNCI2014004 session_3；seed 666，9 个 LOSO 目标被试。',
        '主指标：逐被试 Balanced Accuracy；先对九被试等权汇总，报告均值和被试间样本 SD。',
        '四个正式条件均为新初始化 IFNet，CE + 配对特征损失；教师为 CE 或 CE+MCC 适配。',
        '', '## 汇总', '',
        '| 模型 | 条件 | Balanced Accuracy (%) 均值±SD | Accuracy (%) 均值±SD | Kappa | CE |',
        '|---|---|---:|---:|---:|---:|',
    ]
    for row in summary_rows:
        md.append(f"| {row['model']} | {row['condition']} | "
                  f"{row['balanced_accuracy_mean_pm_std']} | {row['accuracy_mean_pm_std']} | "
                  f"{row['kappa_mean']:.4f} | {row['cross_entropy_mean']:.4f} |")
    md.extend(['', '## 配对比较', '',
               '| 比较 | ΔBA (pp) | 被试配对 bootstrap 95% CI | W/T/L |',
               '|---|---:|---:|---:|'])
    for row in paired_rows:
        md.append(f"| {row['comparison']} | {row['mean_balanced_accuracy_difference_pp']:+.2f} | "
                  f"[{row['paired_bootstrap_95ci_low_pp']:+.2f}, "
                  f"{row['paired_bootstrap_95ci_high_pp']:+.2f}] | "
                  f"{row['wins']}/{row['ties']}/{row['losses']} |")
    md.extend(['', '## 结论', ''] + [f'{i}. {line}' for i, line in enumerate(conclusions, 1)])
    md.append('')
    md.append('95% CI 是按被试重抽样的探索性区间；没有报告显著性检验。')
    md_path = require_report_output(report_dir / 'conclusions.md')
    md_path.write_text('\n'.join(md) + '\n', encoding='utf-8')
    return summary


def evaluate_run(config_path: Path, config: dict, run_id: str,
                 gpu: int, progress_path: str | None = None) -> dict:
    device = _device(gpu)
    subject_rows, fold_records = [], []
    for fold in range(9):
        saved = _load_saved_eval(config_path, config, run_id, fold)
        if saved is None:
            rows, record = _evaluate_fold(config_path, config, run_id, fold, device)
        else:
            rows, record = saved
        subject_rows.extend(rows)
        fold_records.append(record)
        source = 'saved predictions' if record.get('loaded_from_saved_predictions') else 'inference'
        print(f"[target-eval] fold={fold+1}/9 subject={fold+1} rows={len(rows)} source={source}",
              flush=True)
    summary = _write_reports(run_id, config, subject_rows, fold_records)
    cleanup = _prune_run_weights(run_id, int(config['seed']))
    summary['model_artifact_cleanup'] = cleanup
    _atomic_json(REPORTS_ROOT / 'distill' / run_id / 'summary.json', summary)
    summary_path = RESULTS_ROOT / 'distill' / run_id / 'final_summary.json'
    _atomic_json(summary_path, summary)
    if progress_path:
        progress = _read_json(progress_path)
        _progress(progress_path, run_id, progress.get('records', []), 9,
                  status='completed', message='all nine folds trained and target evaluation completed')
    print(f"[report] {REPORTS_ROOT / 'distill' / run_id / 'summary.xlsx'}", flush=True)
    return summary


def enqueue_run(config_path: Path, config: dict, smoke: bool = False,
                queue_root: str | None = None) -> str:
    from experiments.distill.queue import DEFAULT_ROOT, enqueue_snapshot
    config_path = Path(config_path).resolve()
    formal_run_id = str(config['storage']['external_run_id'])
    run_id = formal_run_id + '_smoke' if smoke else formal_run_id
    folds = [0] if smoke else list(range(9))
    snapshot = {
        'kind': 'mcc_diffkd_loso_v1', 'config': str(config_path),
        'config_sha256': _sha256(config_path), 'run_id': run_id,
        'stage': 'smoke' if smoke else 'train_evaluate', 'folds': folds,
        'args': {
            'dataset': config['dataset'], 'teacher': 'mirepnet', 'student': 'ifnet',
            'seeds': [int(config['seed'])], 'methods': list(METHODS),
            'conditions': list(METHODS), 'protocol': 'loso_7_1_1',
            'keys': folds, 'epochs': 2 if smoke else int(config['student']['epochs']),
            'teacher_epochs': 2 if smoke else int(config['teacher_adaptation']['epochs']),
            'diffusion_updates': 4 if smoke else int(config['diffusion']['training_updates']),
        },
    }
    job_id = enqueue_snapshot(
        snapshot, root=queue_root or DEFAULT_ROOT,
        argv=['mcc-diffkd', '--smoke' if smoke else '--full'],
        label=(f'MCC-DiffKD LOSO smoke seed={config["seed"]}' if smoke else
               f'MCC-DiffKD LOSO seed={config["seed"]}, 9 folds sequential'))
    return job_id


def _run_training(config_path: Path, config: dict, run_id: str, folds: list[int],
                  gpu: int, progress_path: str | None,
                  stop_file: str | None, smoke: bool) -> int:
    device = _device(gpu)
    records = []
    for index, fold in enumerate(folds, start=1):
        print(f"[train] fold={index}/{len(folds)} target_subject={fold+1}", flush=True)
        record = _train_fold(config_path, run_id, fold, config, device, stop_file, smoke=smoke)
        if record.get('status') == 'paused':
            records.append({'fold': fold + 1, 'status': 'paused', 'stage': record.get('stage')})
            _progress(progress_path, run_id, records, len(folds), status='paused',
                      message=f"paused at target subject {fold+1}")
            return 75
        records.append({'fold': fold + 1, 'status': 'complete',
                        'run_fingerprint': record.get('run_fingerprint')})
        _progress(progress_path, run_id, records, len(folds),
                  message=f"completed {len(records)}/{len(folds)} requested folds")
        if _stop_requested(stop_file):
            _progress(progress_path, run_id, records, len(folds), status='paused',
                      message='stop requested after fold boundary')
            return 75
    if smoke:
        smoke_result = roots['run_result'] / 'smoke_result.json'
        if not smoke_result.is_file() or _read_json(smoke_result).get('status') != 'passed':
            raise RuntimeError('smoke completed without a verified smoke_result.json')
        _progress(progress_path, run_id, records, len(folds), status='completed',
                  message='train-only MCC-DiffKD smoke passed; no target evaluation')
        print('[smoke] passed; no target evaluation was run', flush=True)
        seed_root = f"seed_{int(config['seed'])}"
        for base in (WEIGHTS_ROOT, DATA_CACHE_ROOT):
            path = require_external_output(base / run_id / seed_root)
            if path.exists():
                shutil.rmtree(path)
        cleanup = {'status': 'passed-smoke-weights-and-cache-removed',
                   'run_id': run_id, 'seed': int(config['seed']),
                   'retained': str(RESULTS_ROOT / 'distill' / run_id / seed_root)}
        _atomic_json(roots['run_result'] / 'smoke_cleanup.json', cleanup)
        return 0
    if folds != list(range(9)):
        _progress(progress_path, run_id, records, len(folds), status='completed',
                  message='selected training folds complete; target evaluation not run')
        return 0
    evaluate_run(config_path, config, run_id, gpu, progress_path)
    return 0


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, default=DEFAULT_CONFIG)
    parser.add_argument('--run-id', default=None)
    parser.add_argument('--folds', nargs='+', type=int, default=list(range(9)))
    parser.add_argument('--gpu', type=int, default=0)
    parser.add_argument('--progress', default=None)
    parser.add_argument('--stop-file', default=None)
    parser.add_argument('--expected-config-sha256', default=None)
    parser.add_argument('--queue-root', default=None)
    parser.add_argument('--enqueue', action='store_true')
    parser.add_argument('--smoke', action='store_true')
    parser.add_argument('--evaluate', action='store_true')
    args = parser.parse_args(argv)
    if any(fold < 0 or fold > 8 for fold in args.folds) or len(set(args.folds)) != len(args.folds):
        parser.error('--folds must be unique zero-based subject indices in [0,8]')
    if args.expected_config_sha256 and _sha256(args.config) != args.expected_config_sha256:
        parser.error('config hash differs from queued immutable snapshot')
    return args


def main(argv=None) -> int:
    args = parse_args(argv)
    try:
        config_path = args.config.resolve()
        config = _read_yaml(config_path)
        _validate_config(config)
        run_id = str(args.run_id or config['storage']['external_run_id'])
        if args.enqueue:
            job = enqueue_run(config_path, config, smoke=args.smoke, queue_root=args.queue_root)
            print(f'[queue] job={job} run_id={run_id}', flush=True)
            return 0
        if args.evaluate:
            evaluate_run(config_path, config, run_id, args.gpu, args.progress)
            return 0
        folds = args.folds
        if args.smoke:
            folds = [folds[0]]
            if run_id == config['storage']['external_run_id']:
                run_id = run_id + '_smoke'
        return _run_training(config_path, config, run_id, folds, args.gpu,
                             args.progress, args.stop_file, args.smoke)
    except Exception as exc:
        traceback.print_exc()
        print(f'[mcc-diffkd] error: {type(exc).__name__}: {exc}',
              file=sys.stderr, flush=True)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
